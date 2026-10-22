"""V4 manuscript from actual frozen results; preserves all earlier drafts."""
from pathlib import Path
import argparse,json,re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1];D=ROOT/'artifacts/v4'
NAMES={'constant':'恒定加速＋门控','eta':'ETA排序＋门控','search':'125动作局部搜索','bc':'行为克隆＋门控','raw_rl':'无门控RL',
       'no_memory':'丢失指令记忆','individual':'独立许可拼接','joint_rl':'联合门控RL'}


def build(out):
    data=json.loads((D/'summary.json').read_text());audit=json.loads((D/'priority_audit/summary.json').read_text())
    timing=json.loads((D/'timing/results.json').read_text());doc=Document();sec=doc.sections[0]
    for border in doc.styles.element.xpath('.//w:pBdr'):border.getparent().remove(border)
    sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=sec.bottom_margin=Cm(2.1);sec.left_margin=sec.right_margin=Cm(2.3)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Caption'):
        st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.line_spacing=1.2;doc.styles['Normal'].paragraph_format.space_after=Pt(5)
    doc.styles['Title'].font.size=Pt(18);doc.styles['Heading 1'].font.size=Pt(14);doc.styles['Heading 2'].font.size=Pt(12)
    doc.styles['Subtitle'].font.size=Pt(11);doc.styles['Subtitle'].font.italic=False
    doc.core_properties.author='李东旭';doc.core_properties.title='联合可恢复制动证书下的多智能体合流策略切换'
    def p(text,style=None):
        q=doc.add_paragraph(text,style)
        if style is None:q.paragraph_format.first_line_indent=Cm(.74)
        return q
    def h(text):doc.add_heading(text,1)
    def eq(text):
        q=doc.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;m=OxmlElement('m:oMath')
        def run(parent,value):
            r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=value;t.set(qn('xml:space'),'preserve');r.append(t);parent.append(r)
        pos=0
        for token in re.finditer(r'([A-Za-z]+)_(?:\{([^}]+)\}|([A-Za-z0-9]+))',text):
            if token.start()>pos:run(m,text[pos:token.start()])
            sub=OxmlElement('m:sSub');base=OxmlElement('m:e');index=OxmlElement('m:sub');run(base,token.group(1));run(index,token.group(2) or token.group(3));sub.append(base);sub.append(index);m.append(sub);pos=token.end()
        if pos<len(text):run(m,text[pos:])
        q._p.append(m)
    def table(headers,rows,widths):
        t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
        for col,w in zip(t.columns,widths):col.width=Cm(w)
        for c,w,s in zip(t.rows[0].cells,widths,headers):c.width=Cm(w);c.text=s
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            for c,w,s in zip(t.add_row().cells,widths,row):c.width=Cm(w);c.text=str(s)
        for i,row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for c in row.cells:
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7ECF2' if i==0 else ('F3F5F7' if i%2==0 else 'FFFFFF'));c._tc.get_or_add_tcPr().append(shade)
                for q in c.paragraphs:
                    q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.space_after=q.paragraph_format.space_before=Pt(4);q.paragraph_format.line_spacing=1
                    if len(rows)<=4:q.paragraph_format.keep_with_next=i<len(t.rows)-1
                    for r in q.runs:r.font.size=Pt(9);r.bold=i==0
        q=doc.add_paragraph();q.paragraph_format.line_spacing=Pt(4);q.paragraph_format.space_after=Pt(4)
    def avg(split,m,k):return sum(r[k] for r in data['summary'][split][m])/3
    p('联合可恢复制动证书下的\n多智能体合流策略切换','Title');p('李东旭','Subtitle')
    h('摘要')
    p('针对学习控制与传统备用切换中的联合可执行性问题，本文研究三辆受控车从两条固定路线汇入一条道路的仿真系统。'
      '在路线进度、速度及上一加速度指令构成的状态空间中，建立联合制动备用轨迹，并对单个仿真子步内的冲突区占用时间区间进行精确检查。'
      '给出可恢复集合下的条件安全证明，以及独立许可不能直接拼接、安全停车不蕴含任务完成的反例。'
      '三个独立训练种子、八种方法共完成960次固定协议SUMO评估。联合门控RL在两类场景中全部完成且无合同违规，但其平均到达时间未优于恒定加速加同一门控。'
      '结果表明安全过滤本身可以承担主要协调作用，不能将闭环安全与完成率直接归因于多智能体学习。本文结论限于全受控、固定路线及精确离散动力学。')
    p('关键词  多智能体强化学习  合流协调  模型预测屏障  可恢复性  控制器切换')
    h('1 引言与相关研究')
    p('以置信度决定是否接受学习动作时，分数只有与具体约束及可执行备用联系起来才具有明确含义。本文将接受条件定义为备用轨迹存在性，'
      '而非把神经网络输出或sigmoid分数视作安全概率。前期单车道跟驰实验中，恒定加速加安全门控已经超过学习策略的平均进展，促使本文转向存在次序冲突的双路合流任务。')
    p('模型预测安全认证通过连接安全终端集合的轨迹保护学习控制[1,2]；MAMPS将学习与备用切换用于多智能体系统[3]。'
      'DMPS进一步指出屏障可能妨碍任务推进，并利用局部恢复规划结合短期进展与长期收益[4]。自动驾驶中的安全MARL也已有合流实验和公开实现[5]。'
      '因此，安全切换、收益导向屏障及多智能体合流本身均不是本文首次提出的思想。')
    p('本阶段提供可复现的路线占用检查、带指令记忆的联合证书、针对机制的消融和明确的负面学习结果。'
      '核心研究问题是：在相同执行接口和安全证书下，多智能体学习是否提供超过简单加速或传统排序的额外收益？')
    h('2 系统模型与路线合同')
    p('系统含三辆完全受控车，来路标签为(0,1,0)，同来路ego0始终在ego2之前。路线进度s以汇入口为零点，状态还包含速度v和上一指令u⁻。'
      '决策周期Δ=0.5秒，子步h=0.1秒，速度上限20 m/s，加速度指令范围−5至2.5 m/s²，相邻指令差不超过2.5 m/s²。')
    eq('v_i⁺ = clip(v_i + h u_i, 0, 20),     s_i⁺ = s_i + h v_i⁺')
    p('同来路车辆保持至少6.8米车头间距，其中车长4.8米、间距余量2米。不同来路车辆不得同时占据冲突区Z=[−25,6.8]；'
      '当二者均进入s≥−25的范围后，还要求其路线坐标差绝对值至少6.8米，防止一车刚离区、后车已贴近的边界情况。')
    p('每子步内路线坐标仿射变化。解不等式sᵢ(λ)∈Z得到占区时间区间Iᵢ，要求不同来路的Iᵢ与Iⱼ不相交；'
      '两车进入保护范围后的最小相对距离则在区间端点或相对坐标过零点取得。该检查覆盖子步内部，而非仅检查两个采样端点。')
    eq('s_i(λ) = s_{i,0} + λ(s_{i,1} − s_{i,0}),     0 ≤ λ ≤ 1')
    p('真实SUMO使用两条入边与一条出边，通过TraCI逐子步setSpeed执行声明模型，并用累计行驶距离读取路线进度。关闭隐式防撞及换道干预，'
      '不使用内部路口连接；空间绘制可能在接点几何跳接。因此证明仅覆盖路线合同，不能直接作为任意车辆外形下的连续空间防碰撞证明。仿真另行记录SUMO碰撞。')
    h('3 联合备用证书及理论性质')
    eq('B_i(u_i⁻) = max(−5, u_i⁻ − 5Δ)')
    p('联合备用令三车遵守上述指令递推直至停止。终端集合T要求所有车速度为零、上一指令非正且路线合同成立。'
      '没有外部交通时，继续执行B保持T。可恢复集合K由执行B时始终满足合同并有限时间到达T的状态组成。'
      '候选首步后接备用尾部，整条轨迹通过才放行；每车有提案和备用两个模式，共枚举八种组合，选择平方干预最小者。')
    p('命题1：精确动力学、固定路线和全车服从联合指令的条件下，K在B下不变；从K开始的联合证书切换保持路线合同。'
      '证明：备用保留上一指令，因此执行首步后的递推等于原证书尾部；尾部仍安全且到达T。通过认证的候选同样提供安全首步和可恢复后继。'
      '全备用保证可行模式非空，逐步归纳即可。实现检查24个备用周期，覆盖所声明速度及指令界内的停止过程。')
    p('命题2：可恢复性不蕴含有限时间完成。证明：令一车停在冲突区，其他车停在区外，均满足合同且上一指令非正；该状态属于T，'
      '持续备用仍安全，却不会继续通过。因此必须将安全与任务推进分别评估，不能仅以零碰撞判断控制器质量。')
    p('独立许可反例：三车(s,v)为(−40,10)、(−40,10)、(−90,10)，上一指令均为0。'
      '分别检查“本车加速2.5、其他车制动−2.5”可获许可，但三个许可拼接为同时加速后无法通过联合证书。'
      '两个来路的备用占区区间重叠，联合选择则令其中一辆先让行。该反例针对独立拼接消融，不是对已有联合屏障的反驳。')
    p('不同来路的冲突关系随合流推进变化，不能直接继承固定单车道链式动态规划的线性复杂度结论。当前三车实现采用完整模式枚举；'
      '自由换道、背景车辆、模型误差与更大规模的处理仍需另行建立合同。')
    h('4 学习与实验协议')
    p('采用共享Actor和集中式Critic，输入分别为12维与36维，两层64单元。动作集合为−5、−2.5、0、1.25、2.5 m/s²。'
      '三个训练种子127、139、151各使用20回合恒定加速门控示范、15轮行为克隆，再进行6000步门控内PPO。'
      '概率比对应原始提案，执行替换属于环境转移。奖励由截断进展、完成、未完成时间成本、指令jerk和干预组成。')
    p('协议与源文件哈希在训练前冻结，最终检查点直接测试。仅常规场景用于训练；高密度场景提高初速并使到达时刻更接近。'
      '每方法、每种子、每分布20回合，八种方法共960次执行。所有方法共享场景，不把重复的确定性传统基线作为额外独立样本。')
    p('对照包含恒定加速、ETA排序、125个联合动作的局部搜索、行为克隆、无门控RL、丢失备用指令记忆、独立许可拼接和联合门控RL。'
      '局部搜索使用与RL相同五动作，对每个联合动作认证，随后以2秒恒定指令预测的进展减动作变化惩罚选择；它不是全局最优MPC或已发表算法的完整复现。')
    p('主指标为到达s=60的逐车平均时间，未完成车辆按25秒计，碰撞回合全部车辆按25秒计。完成车辆统一转入巡航提案生命周期，'
      '避免停在目标线后阻碍后车，提案仍经过同一证书。同时报告回合完成率、合同违规、碰撞及指令jerk。配对bootstrap只作描述性分析，不签发总体风险保证。')
    h('5 结果与机制分析')
    for split,label in [('nominal','常规场景'),('dense','高密度场景')]:
        doc.add_heading(label,2)
        table(['方法','种子127','种子139','种子151','到达时间 s'],
            [[name]+[f'{r["completed"]}/{r["collisions"]}/{r["violating_episodes"]}' for r in data['summary'][split][m]]+[f'{avg(split,m,"mean_time"):.3f}'] for m,name in NAMES.items()], [4.4,2.6,2.6,2.6,4.2])
        p('计数列为完成／碰撞／合同违规回合数，各项分母20；到达时间取三个种子均值，越低越好。')
    p('联合门控RL的120次执行均完成且无碰撞或合同违规，但种子127和139的逐场景到达时间与恒定加速门控一致，种子151略慢。'
      '常规三种子平均11.408秒，对照11.308秒；高密度10.489秒，对照10.467秒。因此没有观察到MARL相对于这一简单强基线的效率提升。')
    p('ETA排序全部完成，但平均用时更长。2秒局部搜索只有10/20常规及5/20高密度场景完全完成，尽管没有合同违规，'
      '与安全不蕴含推进的命题一致。此结果揭示具体启发式的短视，不能解释为RL胜过成熟的预测控制。')
    p('丢失指令记忆在常规和高密度均出现大量合同违规，高密度每个种子发生一次碰撞；独立许可拼接也产生违规。'
      '这些消融支持保留指令记忆与最终联合检查的重要性。行为克隆的种子151几乎不能完成，说明初始化本身存在明显种子敏感性。')
    doc.add_picture(str(D/'results.png'),width=Cm(16.3));doc.paragraphs[-1].paragraph_format.keep_with_next=True
    p('图1 三训练种子的到达时间与回合完成数，黑点表示逐种子结果','Caption')
    h('6 开发诊断与工程测量')
    p('主实验后另取20个开发场景，比较原始选择及三个静态优先次序，共80次执行；优先级仅打破相同平方干预代价下的平局，不改变安全条件。'
      f'逐场景事后挑选最好的结果，仅在{audit["improved_cases"]}/20个场景改善，平均节省{audit["mean_seconds_saved"]:.2f}秒。'
      '这个后验选择不是可部署控制器，也不是学到的性能，只表明当前优先级变体的可利用空间有限。')
    p('对每种子、每分布首条保存轨迹的非初始状态审计，策略提出最大加速的比例约90.8%至96.7%。这与执行表现接近恒定加速一致，'
      '但该样本不是全部测试决策，不能据此断言策略处处相同。')
    table(['方法','均值 ms','P95 ms','最大 ms'],[[NAMES[m],f'{r["mean_ms"]:.2f}',f'{r["p95_ms"]:.2f}',f'{r["max_ms"]:.2f}'] for m,r in timing.items()],[6,3.4,3.4,3.6])
    p('每方法五个预定场景，主实验结束后单进程测量，包含策略前向、门控与SUMO步进。所有测量均低于500毫秒决策周期；'
      '样本P95与最大值不是最坏执行时间保证。完整源代码、协议、权重、逐回合数据与重放脚本公开于CoopGate仓库。')
    h('7 讨论与结论')
    p('本阶段把联合停止证书从单车道接入具有路线冲突的合流任务，验证了条件安全、独立许可失效和安全停车可能阻碍任务推进。'
      '主要限制为全受控、固定路线、精确速度适配器、保守冲突区及较短训练预算，尚无实车或混合交通验证。')
    p('实验没有支持学习优于恒定加速门控。仅增加训练量或改变算法名称不足以建立贡献；在继续扩大MARL前，应先以充分时域的小规模规划参考，'
      '检验场景是否存在显著的长期次序收益，再研究能在可恢复性约束下利用该收益的决策变量。模型误差下的证书与对等已有方法比较仍是必要工作。')
    h('参考文献')
    for ref in [
      '[1] Wabersich K P, Zeilinger M N. Linear model predictive safety certification for learning-based control. 2018. https://arxiv.org/abs/1803.08552',
      '[2] Wabersich K P, Zeilinger M N. A predictive safety filter for learning-based control of constrained nonlinear dynamical systems. Automatica, 2021, 129:109597.',
      '[3] Zhang W, Bastani O, Kumar V. MAMPS: Safe Multi-Agent Reinforcement Learning via Model Predictive Shielding. 2019. https://arxiv.org/abs/1910.12639',
      '[4] Banerjee A, Rahmani K, Biswas J, Dillig I. Dynamic Model Predictive Shielding for Provably Safe Reinforcement Learning. NeurIPS, 2024. https://arxiv.org/abs/2405.13863',
      '[5] Safe Multi-Agent Reinforcement Learning with Bilevel Optimization in Autonomous Driving. Author implementation: https://github.com/SAIL-Research-Lab/Safe-MARL-in-Autonomous-Driving',
      '[6] CoopGate source and data. https://github.com/Leo-7GOAT/CoopGate']:
        q=p(ref);q.paragraph_format.first_line_indent=Cm(0);q.paragraph_format.space_after=Pt(3)
        for r in q.runs:r.font.size=Pt(9)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER;f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
    doc.save(out);print(out)


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--out',required=True);build(p.parse_args().out)
