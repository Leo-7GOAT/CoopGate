"""Substantive V10 manuscript using the retained document's native design."""
import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
from zipfile import ZipFile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table

ROOT=Path(__file__).resolve().parents[1]


def build(reference,out):
    before_hash=hashlib.sha256(reference.read_bytes()).hexdigest();doc=Document(reference)
    template=deepcopy(doc.tables[0]._tbl)
    r=json.loads((ROOT/'artifacts/v10/summary.json').read_text());main=r['overall']['rollout_learned'];gain=r['contrasts_perturbed']['all']['v9_learned_pair']
    title='有界执行误差下的合流控制\n联合安全证书与滚动策略改进'
    doc.paragraphs[0].text=title;doc.core_properties.title=title.replace('\n',' ');doc.core_properties.version='10'
    cut=doc.paragraphs[3]._p;remove=False
    for e in list(doc.element.body):
        if e is cut:remove=True
        if remove and e.tag!=qn('w:sectPr'):doc.element.body.remove(e)

    def p(text,style=None):
        a=doc.add_paragraph(text,style)
        if style is None:a.paragraph_format.first_line_indent=Cm(.74)
        return a
    def h(text):doc.add_heading(text,1)
    def page(text):doc.add_page_break();h(text)
    def eq(text):
        a=doc.add_paragraph();a.alignment=WD_ALIGN_PARAGRAPH.CENTER;m=OxmlElement('m:oMath')
        def run(parent,value):
            z=OxmlElement('m:r');t=OxmlElement('m:t');t.text=value;t.set(qn('xml:space'),'preserve');z.append(t);parent.append(z)
        pos=0
        for token in re.finditer(r'([A-Za-zΣ0-9]+)(?:_\{([^}]+)\}(?:\^\{([^}]+)\})?|\^\{([^}]+)\})',text):
            run(m,text[pos:token.start()]);sub=token.group(2);sup=token.group(3) or token.group(4)
            z=OxmlElement('m:sSubSup' if sub and sup else 'm:sSub' if sub else 'm:sSup')
            base=OxmlElement('m:e');run(base,token.group(1));z.append(base)
            if sub:
                index=OxmlElement('m:sub');run(index,sub);z.append(index)
            if sup:
                index=OxmlElement('m:sup');run(index,sup);z.append(index)
            m.append(z);pos=token.end()
        run(m,text[pos:]);a._p.append(m)
    def table(headers,rows,widths=None):
        element=deepcopy(template);doc.element.body.insert(len(doc.element.body)-1,element);t=Table(element,doc._body)
        while len(t.rows)>1:t._tbl.remove(t.rows[-1]._tr)
        while len(t.columns)<len(headers):t.add_column(Cm(2.5))
        t.autofit=False;widths=widths or [8.2,4.1,4.1]
        for cell,value in zip(t.rows[0].cells,headers):cell.text=value
        header_fill=deepcopy(t.rows[0].cells[0]._tc.get_or_add_tcPr().find(qn('w:shd')))
        if header_fill is not None:
            for cell in t.rows[0].cells:
                props=cell._tc.get_or_add_tcPr();old_fill=props.find(qn('w:shd'))
                if old_fill is not None:props.remove(old_fill)
                props.append(deepcopy(header_fill))
        for values in rows:
            for cell,value in zip(t.add_row().cells,values):cell.text=str(value)
        for col,width in zip(t.columns,widths):col.width=Cm(width)
        for i,row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell,width in zip(row.cells,widths):
                cell.width=Cm(width)
                for a in cell.paragraphs:
                    a.alignment=WD_ALIGN_PARAGRAPH.CENTER;a.paragraph_format.space_before=a.paragraph_format.space_after=Pt(5)
                    a.paragraph_format.line_spacing=1.05;a.paragraph_format.keep_with_next=i<len(t.rows)-1
                    for run in a.runs:run.font.size=Pt(10);run.bold=i==0
        doc.add_paragraph().paragraph_format.space_after=Pt(2)

    p('针对学习控制与传统规划切换中安全约束引起的通行效率损失，本文研究有界执行误差下的三车固定路线合流。'
      '在联合鲁棒备用证书内，逐周期比较传统与学习基策略的名义余程代价，并通过逐车坐标前瞻选择温和让行动作。'
      '每个替换指令必须同时满足联合安全证书和严格代价下降条件。构造独立最早到达下界，提前结束无法改进当前方案的尾部评价。')
    p(f"在200个全新初态、四类执行误差和八种方法共6400次SUMO执行中，新控制器完成{main['completed']}/800回合，"
      f"记录碰撞{main['collision_episodes']}回合、合同违规{main['contract_violation_episodes']}回合、中止{main['aborted']}回合。"
      f"非零误差下相对上一版学习双方案，平均截断到达代价改善{gain['gain_percent']:.2f}%，平均节省{gain['mean_saved_s']:.3f} s，"
      f"95%整簇配对区间为[{gain['paired_cluster_ci95_s'][0]:.3f}, {gain['paired_cluster_ci95_s'][1]:.3f}] s。"
      f"暖机控制周期P95为{main['warm_controller_cycle_p95_ms']:.2f} ms。名义性能改进与条件鲁棒安全分别论证，扰动下的性能由实验评价。")
    p('关键词  合流控制  联合安全证书  有界执行误差  坐标前瞻  基策略改进')
    h('1 引言与研究定位')
    p('安全过滤能够阻止不满足约束的学习动作，但满足约束并不意味着通行高效。上一版在起点比较传统策略与学习提案，'
      '此后主要在提案加速与制动备用之间作反馈切换。遇到冲突时，较强减速及后续加速可能增加任务代价；'
      '只优化起点的策略选择，无法充分利用执行过程中重新测得的状态。')
    p('已有预测安全过滤器通过备用控制约束学习策略[1]；多智能体坐标rollout已降低联合动作枚举量并建立基策略改进性质[2,3]。'
      '本文沿用这些理论思想，研究如何在带指令记忆、非可分冲突约束与有界执行误差的合流任务中组合并核验。'
      '本文不把坐标前瞻或鲁棒安全过滤本身当作首次提出，也不将整个收益归因于强化学习。')

    page('2 模型与任务定义')
    p('三车路线标签为(0,1,0)，同路线车辆0始终在车辆2前方。状态由路线坐标s、速度v、上一指令及全局决策时刻k组成。'
      '控制周期Δ=0.5 s，子步h=0.1 s，周期内保持指令u不变；每个子步有独立或相关的有界误差w。')
    eq('v_{i,t+1} = clip(v_{i,t} + h(u_{i,k} + w_{i,t}), 0, v_{max})')
    eq('s_{i,t+1} = s_{i,t} + h v_{i,t+1},     |w_{i,t}| ≤ ε')
    table(['参数','符号或含义','取值'],[
        ['最大速度','v上限','20 m/s'],['指令范围','制动与加速','[−5, 2.5] m/s²'],
        ['指令变化率上限','每周期限幅','5 m/s³'],['已知执行误差界','ε','0.5 m/s²'],
        ['最小路线间距','d安全','6.8 m'],['冲突区','Z','[−25, 6.8] m'],
        ['任务终点与截止时间','路线坐标及时间','60 m / 25 s']])
    p('同路线车辆保持至少6.8 m间距；不同路线车辆不能同时占据冲突区，且双方进入s≥−25后需保持路线间距。'
      '每个子步采用端点间仿射位置插值，对整段检查合同。这里的保证针对声明的路线合同，SUMO碰撞事件另行记录。')
    p('任务代价为三车首次在决策时刻到达终点的平均时间，未到达者按25 s截断。速度非负使到达标志可由位置恢复。'
      '已到达车辆继续接受基控制器指令；滚动优化不再修改其指令分量。控制器中止按整回合25 s计入代价，不删除失败样本。')
    p('仿真采用Eclipse SUMO及TraCI速度接口，三车均为合作车辆，固定路线且不自由换道。测量仍精确。'
      '执行误差界预先给定，不是由测试数据估计出的置信概率；指令变化率限制也不等同于实际含噪加速度的变化率限制。')

    page('3 鲁棒联合备用证书')
    p('沿用上一版区间证书，在当前精确测量处建立退化状态区间。动力学单调性允许以u−ε和u+ε传播速度与位置端点：')
    eq('v_{L}^{+} = clip(v_{L} + h(u − ε), 0, v_{max})')
    eq('v_{U}^{+} = clip(v_{U} + h(u + ε), 0, v_{max})')
    eq('s_{L}^{+} = s_{L} + h v_{L}^{+},     s_{U}^{+} = s_{U} + h v_{U}^{+}')
    p('对任意满足界的误差序列，逐步归纳可得真实状态始终位于区间内。同路线检查前车位置下界与后车位置上界；'
      '不同路线以最早可能进入、最晚可能离开时间形成冲突占用外包区间。外包不相交且进路间距成立，是实际路线合同成立的充分条件。')
    p('候选动作执行一个周期后，追加联合制动备用律：')
    eq('B_{i}(u^{-}) = max(−b, u_{i}^{-} − j_{max} Δ)')
    p('证书要求候选动作及最多24个备用周期的区间检查全部通过，最终三车速度上界精确为0且末指令不大于−ε。'
      '在该终端集合中，任何允许误差都不能使速度脱离0，故停止状态不变。区间外包可能保守，本研究没有把它当作精确可达集。')
    p('命题1  条件鲁棒安全。若初始状态属于此可恢复集合，则新控制器执行的每个周期都满足路线合同，且下一实测状态仍有可行备用。')
    p('证明：基控制器的初始动作已获证书，每次坐标替换也核验整个联合指令，因此最终动作始终获证书。'
      '真实下一状态包含于该动作的一周期包络；由该实测点重新传播的同一制动后缀包含于旧证书的后缀包络内。'
      '合同及停止终端条件继续成立，归纳完成。该结论依赖模型及准确测量；有限仿真通过不等于证明所有实际交通安全。')

    page('4 每周期基策略比较与动作前瞻')
    p('冻结学习器在初始状态上提出一个反馈策略，与传统策略形成集合P={0,pθ}；相同提案去重。'
      '每个策略确定协作优先级及初期滑行阶段，但实际指令均接受第3节证书。学习网络为此前训练的集中式上下文bandit，'
      '本轮没有新训练分布式MARL，也不读取未来误差带。')
    p('以每周期未到达车辆数作为整数阶段代价，完整尾部代价为：')
    eq('ℓ(x) = Σ_{i} 1[s_{i} < 60],     V_{k}^{p}(z) = Σ_{t=k}^{N−1} ℓ(x_{t}^{p})')
    p('将该整数乘Δ/3即可得到相应的剩余截断平均到达时间。每周期重新比较两条完整基策略尾部，平局保留传统策略。'
      '滑行阶段使用全局k，前瞻时不重置时间；否则有限时域价值函数与实际策略将不一致。')
    p('取当前较优基策略p*的动作作为incumbent。每车允许区间由物理限幅与上一指令变化率共同给出：')
    eq('L_{i} = max(−b, u_{i}^{−} − j_{max} Δ),     U_{i} = min(a_{max}, u_{i}^{−} + j_{max} Δ)')
    p('每个区间取五个等分值，按基策略权重降序逐车扫描。同权重按车号排序。每次只改变一个分量，但必须检查完整三车指令；'
      '不同车辆各自可行的单独修改不能直接拼接。候选通过安全证书后，执行名义一步并沿p*评价剩余过程：')
    eq('Q_{k}^{p*}(z,a) = ℓ(x) + V_{k+1}^{p*}(F_{0}(z,a))')
    p('只在整数Q严格下降时更新incumbent，平局保持已有动作。一次扫描最多试探15个动作，跳过未变化分量及已到达车辆。'
      '这是局部坐标改进，不承诺得到125个联合网格动作中的全局最优；每个试探的尾部评价和联合证书仍有计算成本。')

    page('5 名义性能改进与精确剪枝')
    p('命题2  名义闭环不退化。令W为同一时刻所有基策略尾部代价的最小值，ρ为每周期重复上述操作的控制器。'
      '在零误差、相同截止时间及精确评价条件下，ρ的闭环代价不高于W，从起点看不差于上一版学习双方案。')
    eq('W_{k}(z) = min_{p∈P} V_{k}^{p}(z)')
    eq('ℓ(x) + V_{k+1}^{p*}(F_{0}(z,ρ(z))) ≤ W_{k}(z)')
    p('证明：初始incumbent等于较优基策略的首动作，其完整Q恰为W；所有替换严格降低Q。'
      '截止时刻闭环与基策略剩余代价均为0。若下一时刻闭环代价不高于W，而W不高于所选p*的尾部代价，'
      '代入上式即得到当前时刻闭环代价不高于W，反向归纳成立。')
    p('该结论不能直接推广到非零误差：真实含扰动后继与名义后继不同，名义收益可能不能沿真实轨迹保持。'
      '因此论文分别陈述条件鲁棒安全和名义性能不退化，不把二者合并为扰动下的全性能保证。')
    p('命题3  剪枝保持选择。忽略车间冲突，使每辆车独立采用符合指令变化率的最大加速序列，得到剩余最早到达周期总数L。'
      'L是任意基策略尾部代价的下界；若累计代价g与L之和已不小于incumbent，候选不可能严格改进，可以停止评价。')
    eq('L_{k}(z) ≤ V_{k}^{p}(z),     g + L_{k}(z) ≥ Q_{best} ⇒ reject')
    p('证明：最大加速包络逐步支配任意符合限幅及变化率的指令。饱和速度与位置更新单调，独立车辆不会比受冲突限制时更晚到达。'
      '使用同一截止时间求和仍保持下界。由于只接收严格改进，达到incumbent的平局也可剪枝。整数代价避免浮点收益比较误差。')
    p('理论结构继承坐标rollout的基策略改进思想[2,3]。本项目的具体工作是将其与含指令记忆的联合合流证书结合，'
      '保留有限时域时钟、可行incumbent及可复核下界，并通过差分审计检查实现是否保持这些条件。')

    page('6 冻结实验与统计方法')
    p('开发阶段仅使用已公开的旧场景：先16个初态320回合，再32个初态896回合；这些结果不进入最终证据。'
      '双轮坐标扫描未显示稳定额外收益而增加计算，故选单轮、五级动作。代码哈希、验收标准和新测试种子在正式实验前冻结[4]。')
    p('最终采用200个全新初态，常规与高密度各100个。常规初速度8至14 m/s，高密度为12至17 m/s；'
      '生成器将主车插入位置分别采样在距汇入口70至110 m和70至85 m范围，同路线后车位于前车之后30至45 m。'
      'SUMO插入步后的实际初态另存原始数据；全部初态均保留，不按方法表现筛除。')
    table(['执行误差','规则','每种方法回合数'],[
        ['零误差','每个子步为0',200],['随机误差','独立均匀于声明界内',200],
        ['持续偏差','每车固定为正或负极值',200],['间歇脉冲','1秒极值脉冲与1秒空隙',200]])
    p('八种方法均使用同一鲁棒证书：V9传统、V9学习双方案、V9完整搜索，以及无学习、固定提案、随机提案、学习二值动作和学习五级动作的滚动优化。'
      '固定索引来自此前训练集；随机提案使用冻结随机规则。相同初态与误差模式共享外生误差带，方法执行顺序随机化。')
    p('预先要求新控制器：无碰撞、合同违规或中止；完成率至少99%；非零误差下相对V9学习双方案改善至少3%，主要配对区间下界大于0；'
      '相对V9完整搜索改善至少2%；高密度场景相对V9学习双方案改善至少2%；零误差更差初态为0；'
      '暖机控制周期P95不超过20 ms，暖机首响应P95不超过25 ms，共九项标准。')
    p('三个非零误差条件共享同一初态。先对每个初态取三个模式的配对差均值，再按200个初态整簇bootstrap一万次，报告95%区间。'
      '主要比较预先指定为新方法对V9学习双方案，其他区间为描述性结果，不宣称多重比较显著性。')

    page('7 主要结果与不利场景')
    names=[('v9_constant','V9传统'),('v9_learned_pair','V9学习双方案'),('v9_full','V9完整搜索'),
           ('rollout_constant','无学习滚动'),('rollout_fixed_pair','固定提案滚动'),('rollout_random_pair','随机提案滚动'),
           ('rollout_learned_binary','学习二值动作'),('rollout_learned','V10学习滚动')]
    p('表1  各误差模式的平均截断到达代价  单位为秒','Caption')
    table(['方法','零误差','随机','持续','脉冲'],[[name,*[f"{r['summary_by_error'][mode][m]['mean_cost_s']:.3f}" for mode in ('clean','uniform','bias','pulse')]] for m,name in names],[5.2,2.8,2.8,2.8,2.8])
    p(f"新控制器完成{main['completed']}/800回合，中止{main['aborted']}次、合同违规{main['contract_violation_episodes']}回合、"
      f"碰撞{main['collision_episodes']}回合、超时{main['timeouts']}回合。故障取各类事件并集，避免将同一回合重复相加。")
    gfull=r['contrasts_perturbed']['all']['v9_full'];gbase=r['contrasts_perturbed']['all']['v9_constant']
    p(f"非零误差下，相对V9学习双方案平均改善{gain['gain_percent']:.2f}%，相对V9完整搜索改善{gfull['gain_percent']:.2f}%，"
      f"相对V9传统基线改善{gbase['gain_percent']:.2f}%。这些方法均有相同鲁棒证书，收益不能解释为放松间距或执行误差界。")
    counts=gain['better_tied_worse_clusters']
    p(f"对V9学习双方案，按三种扰动的初态平均差，更好、持平、更差分别为{counts[0]}、{counts[1]}、{counts[2]}个。"
      f"常规与高密度场景分别改善{r['contrasts_perturbed']['nominal']['v9_learned_pair']['gain_percent']:.2f}%和"
      f"{r['contrasts_perturbed']['dense']['v9_learned_pair']['gain_percent']:.2f}%。零误差更差初态为{r['clean_worse_than_v9_learned']}/200。"
      '不利扰动场景仍保留，不能把总体均值改进写成逐场景支配。')

    page('8 消融与工程计算成本')
    p('表2  新方法相对各滚动优化消融的平均节省','Caption')
    nm=dict(names);ablations=['rollout_constant','rollout_fixed_pair','rollout_random_pair','rollout_learned_binary']
    table(['对照','平均节省 秒','95%整簇区间 秒'],[[nm[m],f"{r['contrasts_perturbed']['all'][m]['mean_saved_s']:.3f}",
        f"[{r['contrasts_perturbed']['all'][m]['paired_cluster_ci95_s'][0]:.3f}, {r['contrasts_perturbed']['all'][m]['paired_cluster_ci95_s'][1]:.3f}]"] for m in ablations])
    p('这些对照分离学习提案与通用滚动优化的贡献。无学习版本仍有前瞻和安全证书，固定及随机提案具有同样两个基策略预算。'
      '二值动作对照保留前瞻，只移除允许区间内的中间采样值。')
    learning=r['contrasts_perturbed']['all']['rollout_constant']
    p(f"相对无学习滚动优化，学习提案平均额外节省{learning['mean_saved_s']:.3f} s。"+
      ('该区间包含0，当前实验未可靠建立学习模块的额外通行收益。' if learning['paired_cluster_ci95_s'][0]<=0<=learning['paired_cluster_ci95_s'][1]
       else '其描述性区间见表2；该比较不构成对所有场景或更多车辆规模的保证。')+
      '主要提升应首先归于每周期前瞻和具体动作优化。')
    p('表3  暖机计算分位数  单位为毫秒','Caption')
    table(['方法','首响应P95','控制周期P95'],[[name,f"{r['overall'][m]['warm_first_response_p95_ms']:.2f}",
        f"{r['overall'][m]['warm_controller_cycle_p95_ms']:.2f}"] for m,name in names])
    p('首响应包含初始策略选择与第一次TraCI执行，控制周期计时只统计控制算法；两者均已预热。')
    audit=json.loads((ROOT/'artifacts/v10/pruning_audit.json').read_text())
    saved=100*(1-audit['pruned_tail_steps']/audit['unpruned_tail_steps'])
    p(f"在预指定首种子轨迹每隔一个决策时刻抽取的{audit['states']}个相关诊断状态上，关闭与启用剪枝的动作和值不一致数量均为0；"
      f"剪枝减少{saved:.2f}%的尾部反馈步数。该审计用于验证选择不变，不计作新的独立场景证据。")

    page('9 讨论 结论与研究边界')
    jerk_old=r['perturbed']['v9_learned_pair']['mean_command_jerk'];jerk_new=r['perturbed']['rollout_learned']['mean_command_jerk']
    p(f"滚动前瞻把优化从起点的策略筛选推进到每个控制周期的具体动作，并保留可恢复的联合备用。"
      f"非零误差下平均指令变化率由{jerk_old:.3f}升至{jerk_new:.3f} m/s³，增加{100*(jerk_new/jerk_old-1):.2f}%；这是指令层指标，不是含执行误差的实际jerk保证。"
      f"本轮九项预先声明标准全部通过：{'是' if r['all_criteria_passed'] else '否'}。")
    p('工程代价是增加周期内的前瞻计算。仿真在计算期间暂停，记录的暖机分位数不是最坏执行时间，也没有验证计算延迟下的闭环。'
      '本轮未加入通信延迟统计。三辆合作车辆、固定路线、已知误差界及精确测量仍是适用条件。'
      '更多参与者、非合作行为、感知误差、未知误差边界和实车验证尚未覆盖。')
    p('算法由已有坐标rollout与鲁棒备用过滤思想构成，当前增量是其在带记忆的联合合流合同内的可核验组合、动作细化、'
      '名义闭环改进证明和冻结实验。学习提案的独立贡献应以消融结果为准；不能将这种实现称为新的MARL理论。'
      '投稿前仍需在统一观测、动作和计算预算下复现已发表的同类方法，并扩展任务规模。')
    h('参考文献')
    references=[
        '[1] Wabersich K P, Zeilinger M N. A predictive safety filter for learning-based control of constrained nonlinear dynamical systems. Automatica, 2021, 129:109597. https://arxiv.org/abs/1812.05506',
        '[2] Bertsekas D. Multiagent Rollout Algorithms and Reinforcement Learning. arXiv:1910.00120, revised 2020. https://arxiv.org/abs/1910.00120',
        '[3] Bhattacharya S, Kailas S, Badyal S, Gil S, Bertsekas D. Multiagent Rollout and Policy Iteration for POMDP with Application to Multi-Robot Repair Problems. CoRL 2020, PMLR 155:1814–1828, 2021. https://proceedings.mlr.press/v155/bhattacharya21a.html',
        '[4] CoopGate. Frozen V10 source hashes, paired protocol and raw SUMO executions. https://github.com/Leo-7GOAT/CoopGate/tree/main/artifacts/v10']
    for text in references:
        a=p(text);a.paragraph_format.first_line_indent=Cm(0);a.paragraph_format.space_after=Pt(4)
        for run in a.runs:run.font.size=Pt(9)
    doc.save(out)
    assert hashlib.sha256(reference.read_bytes()).hexdigest()==before_hash
    with ZipFile(reference) as old,ZipFile(out) as new:
        changed=[name for name in old.namelist() if old.read(name)!=new.read(name)]
        assert set(changed)<={'word/document.xml','docProps/core.xml','word/_rels/document.xml.rels','[Content_Types].xml'},changed
    print(json.dumps({'output':str(out),'reference_sha256':before_hash,'changed_parts':changed}))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--reference',type=Path,required=True);parser.add_argument('--out',type=Path,required=True)
    args=parser.parse_args();build(args.reference,args.out)
