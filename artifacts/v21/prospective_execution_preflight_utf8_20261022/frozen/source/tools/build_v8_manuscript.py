"""Rewrite the retained manuscript around exact acceleration and its evidence."""
import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
from zipfile import ZipFile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table

ROOT = Path(__file__).resolve().parents[1]
D = ROOT/'artifacts/v8'


def build(reference, out):
    original_hash = hashlib.sha256(reference.read_bytes()).hexdigest()
    doc = Document(reference); template = deepcopy(doc.tables[0]._tbl)
    result = json.loads((D/'summary.json').read_text()); summary = result['summary']; timing = result['timing']; mechanism = result['mechanism']
    nominal, dense = summary['nominal']['bb_natural'], summary['dense']['bb_natural']
    title = '带可恢复性证书的合流规划\n精确剪枝与学习排序'
    doc.paragraphs[0].text = title; doc.core_properties.title = title.replace('\n', ' '); doc.core_properties.version = '8'
    doc.paragraphs[3].text = ('针对学习控制与传统规控切换中性能评分不可靠、在线核验开销较大的问题，本文研究三辆完全受控车的固定路线合流。'
        '将传统基线保留在37个联合反馈候选中，以带指令记忆的制动证书约束执行；通过独立最早到达时间构造代价下界，'
        '在不改变有限候选库最优选择的前提下提前停止无改进可能的候选。学习模型仅对搜索顺序排序，不能替代下界或安全证书。')
    doc.paragraphs[4].text = (f"在200个全新场景、七种方法共1400次SUMO执行中，全部候选代价与原模型一致。自然顺序剪枝首次响应P95为"
        f"{timing['bb_natural']['first_response_p95_ms']:.2f} ms，原完整搜索为{timing['reference_full']['first_response_p95_ms']:.2f} ms；"
        f"相对传统基线的常规与高密度平均到达时间分别改善{nominal['gain_percent']:.2f}%和{dense['gain_percent']:.2f}%。"
        f"学习排序进一步减少{mechanism['learned_reduction_vs_natural']:.1%}证书调用，但推断开销使其P95略高。"
        '结果限于精确模型与声明候选库，不代表实车或未知交通条件下的安全与最优保证。')
    doc.paragraphs[5].text = '关键词  合流规划  可恢复性证书  代价下界  精确剪枝  学习排序'
    doc.paragraphs[7].text = ('学习策略的动作可行性、任务性能和计算成本需要分别验证。此前实验中，单侧性能校准在高密度分布下出现较多错误接受；'
        '更严格的选择性风险校验又因证据不足而完全回退。完整比较学习提案与传统基线虽能在精确模型中排除性能退化，'
        '但原实现首次响应P95为139.74 ms。本文进一步研究如何保留核验逻辑，同时降低开销并恢复完整候选库的任务收益。')
    doc.paragraphs[8].text = ('多智能体rollout已有相对基策略的改进性质[1]，动态模型预测屏障研究短期进展与长期安全的协调[2]。'
        'Learn then Test通过多重检验控制预测系统风险[3]，学习引导分支搜索也已有研究[4,5]。本文不把编译、分支定界或学习排序本身视为首次创新，'
        '而在统一反馈候选库、物理证书与代价目标下核查它们各自的收益与限制。')
    doc.paragraphs[9].text = ('本文给出与冻结执行模型一致的候选评价内核，构造保留指令变化率约束的乐观到达下界，并证明剪枝不改变有限候选库的最优选择。'
        '实验分别对照原Python实现、相同编译后端的完整搜索、自然排序及学习排序，以免将实现效率误认为学习贡献。')
    cut = next(p._p for p in doc.paragraphs if p.text.startswith('5 V5'))
    removing = False
    for element in list(doc.element.body):
        if element is cut:
            removing = True
        if removing and element.tag != qn('w:sectPr'):
            doc.element.body.remove(element)

    def p(text, style=None):
        paragraph = doc.add_paragraph(text, style)
        if style is None:
            paragraph.paragraph_format.first_line_indent = Cm(.74)
        return paragraph

    def h(text):
        doc.add_heading(text, 1)

    def eq(text):
        paragraph = doc.add_paragraph(); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        math = OxmlElement('m:oMath')
        def run(parent, text):
            r = OxmlElement('m:r'); t = OxmlElement('m:t'); t.text = text
            t.set(qn('xml:space'), 'preserve'); r.append(t); parent.append(r)
        pos = 0
        for token in re.finditer(r'([A-Za-zΣ0-9]+)(?:_\{([^}]+)\}|\^\{([^}]+)\})', text):
            run(math, text[pos:token.start()]); sub = token.group(2) is not None
            script = OxmlElement('m:sSub' if sub else 'm:sSup'); base = OxmlElement('m:e'); run(base, token.group(1)); script.append(base)
            index = OxmlElement('m:sub' if sub else 'm:sup'); run(index, token.group(2) or token.group(3)); script.append(index)
            math.append(script); pos = token.end()
        run(math, text[pos:]); paragraph._p.append(math)

    def table(headers, rows):
        element = deepcopy(template); doc.element.body.insert(len(doc.element.body)-1, element); t = Table(element, doc._body)
        while len(t.rows) > 1:
            t._tbl.remove(t.rows[-1]._tr)
        for cell, value in zip(t.rows[0].cells, headers):
            cell.text = value
        for values in rows:
            for cell, value in zip(t.add_row().cells, values):
                cell.text = str(value)
        for i, row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(5)
                    paragraph.paragraph_format.line_spacing = 1.05
                    paragraph.paragraph_format.keep_with_next = i < len(t.rows)-1
                    for run in paragraph.runs:
                        run.font.size = Pt(10); run.bold = i == 0
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    h('5 冻结学习模型与性能核验')
    p('学习模型以三车位置、速度、相对位置和预计到达时间为12维输入，输出37项方案的概率及相对代价。'
      '三个两层128单元网络在768个常规初态训练，以128个独立验证初态选检查点。标签来自完整模型枚举，'
      '训练属于完整信息离线上下文bandit，不是分布式MARL；本轮没有重训。')
    p('双方案核验仅比较学习提案与传统基线。完整搜索则保留全部37个候选；学习排序剪枝按冻结集成的概率顺序访问候选，'
      '始终先评估传统基线。学习只影响较优方案何时被发现，不能授权动作或判断一个候选可被剪枝。')

    doc.add_page_break(); h('6 乐观到达下界与精确剪枝')
    p('考虑某个候选已执行至第k个决策步的状态。对已经到达的车辆，保留实际到达时间；对未到达车辆，'
      '暂时去掉车辆间冲突和间隔约束，并从当前指令记忆出发，在每个后续决策步施加最大的允许指令。')
    eq('u_{i,k} = min(a_{max}, u_{i,k−1} + j_{max} Δ)')
    p('在相同子步动力学、速度上限、决策时刻到达判据和25秒截断目标下，逐车计算该无冲突运动的到达时间τᵢ，'
      '以三车平均值构造下界Lₖ。这里“无冲突”指移除车辆间约束，不表示该乐观轨迹可以实际执行。')
    eq('L_{k} = (τ₀ + τ₁ + τ₂) / 3')
    p('性质3：任意满足相同执行器约束的候选后续轨迹，其完整回合代价J均满足Lₖ≤J。'
      '证明：从同一指令记忆出发，任一可行第一步指令不超过最大允许指令。按变化率约束归纳，后续指令也不超过乐观指令。'
      '饱和速度与位置递推对指令、速度单调，故乐观车辆在每个子步的位置不落后。相同首次到达与截断判据下，其到达时间不会更晚。已到达车辆的代价已知，取平均即得结论。')
    p('先完整求解传统基线得到已验证的当前最好代价J★。对后续候选逐步展开；只有当下界严格大于J★并留出10⁻¹⁰ s数值裕量时，才提前停止该候选。')
    eq('L_{k} > J^{*} + 10^{−10}  ⇒  stop candidate')
    p('性质4：剪枝返回的代价与最终方案索引均与37方案完整枚举相同。证明：由性质3，被停止候选的完整代价不可能优于已有方案。'
      '其余候选完整评价并更新当前最好解。等价代价不触发剪枝，最终仍按原始候选索引破除同分，故保留完整枚举的最优值及索引。'
      '这只保证声明候选库内最优，不能外推为所有连续控制轨迹的全局最优。')
    p('学习排序越早找到较好方案，后续剪枝阈值就可能越紧。但任意错误排序都不改变正确性；最坏情况下只增加推断开销。'
      '访问顺序并不参与性质3的下界推导，因此该机制允许在相同证明下独立评价学习排序的计算价值。')

    doc.add_page_break(); h('7 实现与启动条件')
    p('实验在AMD Ryzen 5 9600X、64 GB内存的Windows工作站上运行，使用Python 3.11.9、NumPy 2.4.6和Numba 0.65.1。'
      '七种方法均采用单工作进程串行评测，文中计时对应这台工作站，不以GPU并行掩盖不同方法的计算开销。')
    p('编译内核采用float64、单线程和严格浮点运算，不开启fastmath。原每个子步的区间检查、执行器约束和八个联合模式均保留；'
      '将反复创建的小数组改为标量局部量，在已发现证书违规时停止该证书检查。编译循环加速属于实现优化[6]，不能当作强化学习的新算法。')
    p('每个实际执行周期独立检查备用可恢复性、计划状态匹配与拟执行指令的证书。计划失配后永久废弃缓存，调用原传统反馈控制器；'
      '若实测状态已在声明可恢复集合之外，则报错退出而不输出宣称安全的动作。失配后原完整回合性能最优性不再适用。')
    cold = result['cold_start']
    p(f"启动测量使用独立Python进程和全新编译缓存：包导入{cold['imports_ms']:.1f} ms，首次内核调用及编译{cold['first_kernel_call_ms']:.1f} ms，"
      f"随后同一开发初态调用{cold['second_kernel_call_ms']:.3f} ms。系统应在接管控制前完成预热；暖机毫秒值不能描述冷启动。")
    h('8 独立实验协议与改善标准')
    p('开发阶段用此前公开的360个初态逐项核对全部37项代价，这些场景不计入新增独立证据。随后冻结算法、模型、源文件哈希及改善标准，'
      '另取常规与高密度各100个全新初态。每个场景运行七种方法，共1400次真实SUMO执行，方法顺序按预先规定的随机种子打乱。')
    p('七种方法为传统基线、原Python双方案核验、编译双方案核验、编译完整搜索、自然顺序剪枝、学习排序剪枝和原Python完整搜索。'
      '每个新场景额外独立核对全部37项编译代价；评测标签不进入学习决策。原完整搜索在其计时窗口内重新规划，不能复用评测标签冒充在线计算。')
    p('冻结标准要求：编译双方案与原双方案代价逐场景相同；完整与剪枝方法均与原完整搜索选择及代价相同；无记录的安全合同违规；'
      '编译双方案及学习排序剪枝的首次响应P95不超过原双方案的20%；学习排序剪枝较同内核完整搜索平均减少至少30%的证书调用。')
    p('首响应包括高层选择与首个SUMO执行周期，所有方法均预热后串行运行。以场景为配对统计单元，报告描述性bootstrap区间。'
      '1400个方法回合不能视为1400个独立初态；单机计时也不是最坏执行时间或含计算延迟的闭环验证。')

    doc.add_page_break(); h('9 到达性能与选择一致性')
    names = [('constant', '传统基线'), ('reference_pair', '原双方案核验'), ('fast_pair', '编译双方案核验'),
             ('fast_full', '编译完整搜索'), ('bb_natural', '自然顺序剪枝'), ('bb_learned', '学习排序剪枝'), ('reference_full', '原完整搜索')]
    p('表1  相同新场景中的平均截断到达时间  单位为秒', 'Caption')
    table(['方法', '常规 100场景', '高密度 100场景'], [[name, f"{summary['nominal'][m]['mean_arrival_s']:.3f}", f"{summary['dense'][m]['mean_arrival_s']:.3f}"] for m, name in names])
    for regime, label in [('nominal', '常规'), ('dense', '高密度')]:
        r = summary[regime]['bb_natural']; v = result['contrasts'][regime]['full_vs_pair']; ci = v['ci95_s']
        p(f"{label}分布中，完整候选库相对传统基线改善{r['gain_percent']:.2f}%，相对双方案核验再节省{v['mean_saved_s']:.3f} s，"
          f"描述性95%配对区间为[{ci[0]:.3f}, {ci[1]:.3f}] s；比双方案更优/持平/更差分别为{v['better_tied_worse'][0]}/{v['better_tied_worse'][1]}/{v['better_tied_worse'][2]}。")
    p('编译完整搜索和两种剪枝均逐场景复现原完整搜索的方案索引及代价；编译双方案也逐场景复现原双方案。'
      '新增通行收益来自对更多候选的精确核验，与排序学习本身无关。学习排序与自然顺序的最终结果按理论应相同，本轮验证与此一致。')
    p('1400次执行全部完成，未记录碰撞、路线合同违规、不可认证动作、离开备用集合或计划失效回退。'
      '最大单步模型误差为1.42×10⁻¹³，符合当前仿真接口下的精确模型假设。零碰撞样本不等于对未知道路条件的零风险保证。')

    doc.add_page_break(); h('10 计算开销与贡献分离')
    p('表2  暖机运行时间P95  单位为毫秒', 'Caption')
    table(['方法', '高层选择', '首次响应'], [[name, f"{timing[m]['selection_p95_ms']:.3f}", f"{timing[m]['first_response_p95_ms']:.3f}"] for m, name in names])
    ratio = result['latency_ratios']['fast_pair_vs_reference_pair']['p95_ratio']
    p(f"编译双方案保留原选择的同时，首次响应P95降低{100*(1-ratio):.2f}%。自然顺序剪枝保留完整候选库最优结果，"
      f"P95为{timing['bb_natural']['first_response_p95_ms']:.3f} ms；相同编译后端不剪枝时为{timing['fast_full']['first_response_p95_ms']:.3f} ms。"
      '对比后两项可以分离剪枝贡献，避免只与较慢的Python实现比较。')
    p('表3  相同编译后端中的搜索量', 'Caption')
    table(['方法', '平均证书调用', '平均剪枝方案数'], [[name, f"{mechanism[m]['mean_certificate_calls']:.1f}", f"{mechanism[m]['mean_pruned_profiles']:.1f}"]
          for m, name in [('fast_full', '完整搜索'), ('bb_natural', '自然顺序剪枝'), ('bb_learned', '学习排序剪枝')]])
    p(f"学习排序较完整编译搜索减少{mechanism['learned_reduction_vs_full']:.1%}证书调用；较自然顺序额外减少{mechanism['learned_reduction_vs_natural']:.1%}。"
      f"但学习排序首次响应P95为{timing['bb_learned']['first_response_p95_ms']:.3f} ms，略高于自然顺序。"
      '网络推断开销抵消了部分搜索节省，不能据此宣称学习方法在总时延上胜过自然排序。预先声明的六项标准全部通过。')

    doc.add_page_break(); h('11 研究结论与适用边界')
    p('本轮获得了明确的工程进展：同一物理证书下，双方案核验降低响应开销，完整候选库剪枝同时恢复双方案遗漏的通行收益。'
      '理论上，独立最早到达时间提供合法下界，剪枝保留声明候选库的最优选择；学习排序可减小搜索量，但当前规模下尚未转化为更低总时延。')
    p('此前风险校准失败仍是未解决的问题。在另一批1024个校准初态中，13个候选阈值没有一个能同时满足两类交通的选择性性能风险证书；'
      '全回退策略没有通行收益。V8用模型核验替代对评分可信度的依赖，并未证明评分本身可靠，也没有获得未知模型下的性能置信保证。')
    p('适用范围仍是三车固定路线、车辆完全服从指令、精确动力学和有限反馈方案库。仿真在计算期间暂停，尚无含执行延迟的闭环测试、'
      '非合作交通参与者、感知误差、模型扰动或实车验证。若模型失配导致计划失效，原回合性能保证不再成立。')
    p('目前的新增证据支持继续研究与工程展示，但不能单凭编译提速或模块组合宣称达到顶刊顶会的新颖性要求。'
      '后续应比较更大候选库及更复杂交互任务中的搜索尺度效应，并与同类已发表方法统一观测、动作和计算预算。'
      '高层模型仍为集中式离线bandit，不应重新命名为分布式多智能体强化学习。')
    h('参考文献')
    references = [
        '[1] Bertsekas D. Multiagent Rollout Algorithms and Reinforcement Learning. https://arxiv.org/abs/1910.00120',
        '[2] Banerjee A, et al. Dynamic Model Predictive Shielding for Provably Safe Reinforcement Learning. NeurIPS, 2024. https://arxiv.org/abs/2405.13863',
        '[3] Angelopoulos A N, et al. Learn then Test: Calibrating Predictive Algorithms to Achieve Risk Control. https://arxiv.org/abs/2110.01052',
        '[4] Balcan M F, Dick T, Sandholm T, Vitercik E. Learning to Branch. ICML, 2018. https://proceedings.mlr.press/v80/balcan18a.html',
        '[5] Gasse M, Chételat D, Ferroni N, Charlin L, Lodi A. Exact Combinatorial Optimization with Graph Convolutional Neural Networks. NeurIPS, 2019. https://arxiv.org/abs/1906.01629',
        '[6] Numba. Performance Tips. https://numba.readthedocs.io/en/stable/user/performance-tips.html']
    for text in references:
        paragraph = p(text); paragraph.paragraph_format.first_line_indent = Cm(0); paragraph.paragraph_format.space_after = Pt(3)
        for run in paragraph.runs:
            run.font.size = Pt(9)
    doc.save(out)
    assert hashlib.sha256(reference.read_bytes()).hexdigest() == original_hash
    with ZipFile(reference) as before, ZipFile(out) as after:
        changed = [n for n in before.namelist() if before.read(n) != after.read(n)]
        assert set(changed) <= {'word/document.xml', 'word/_rels/document.xml.rels', 'docProps/core.xml', '[Content_Types].xml'}, changed
    print(json.dumps({'output': str(out), 'changed_parts': changed}))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--reference', type=Path, required=True); parser.add_argument('--out', type=Path, required=True)
    args = parser.parse_args(); build(args.reference, args.out)
