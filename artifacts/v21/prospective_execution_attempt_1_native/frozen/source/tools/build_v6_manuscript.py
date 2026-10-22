"""Extend the retained fifth draft with actual learned-selection evidence."""
import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
D = ROOT/'artifacts/v6'


def build(reference, out):
    source_hash = hashlib.sha256(reference.read_bytes()).hexdigest()
    doc = Document(reference); result = json.loads((D/'summary.json').read_text())
    summary, timing = result['summary'], result['timing']
    a, b = summary['nominal']['calibrated'], summary['dense']['calibrated']
    table_template = deepcopy(doc.tables[0]._tbl)
    doc.paragraphs[0].text = '性能置信门控下的\n学习合流方案选择与安全执行'
    doc.core_properties.title = '性能置信门控下的学习合流方案选择与安全执行'
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith('针对强化学习与传统规控切换中'):
            paragraph.text = ('针对学习控制与传统规控切换中可执行性和性能可信度混淆的问题，本文研究三辆完全受控车的双路合流。'
                              '低层保留带指令记忆的联合制动证书，高层通过离线完整信息bandit选择通行偏好和提前让行方案，'
                              '并以独立校准的相对基线代价上界决定接受或回退。证明分别对应精确模型下的路线合同与可交换条件下的边际性能退化界。')
        elif paragraph.text.startswith('对80个独立新场景的五种方法'):
            paragraph.text = (f"前阶段完整规划在80个场景中获得约7.29%的均值收益；本阶段另以120个独立场景、六种方法完成720次SUMO执行。"
                              f"校准门控在常规与高密度分布分别取得{a['mean_saved_s']:.3f} s和{b['mean_saved_s']:.3f} s的平均时间节省，"
                              f"保留完整规划教师{a['teacher_gain_retained_percent']:.1f}%和{b['teacher_gain_retained_percent']:.1f}%的均值收益。"
                              '所有结果限定于三车固定路线和声明动力学；性能校准不等同于碰撞风险保证，高密度偏移不受常规校准理论覆盖。')
        elif paragraph.text.startswith('本文提供可运行的有限候选规划器'):
            paragraph.text = ('本文首先以完整规划识别可实现的改进空间，再学习高层方案选择并校准相对传统基线的性能差。'
                              '学习目标是每回合选择一个联合反馈方案；这是集中式高层学习，不是新的分布式MARL。全文分别报告规划收益、学习收益和门控的收益损失。')
        elif paragraph.text.startswith('关键词'):
            paragraph.text = '关键词  合流协调  上下文bandit  性能校准  可恢复性  控制器切换'
        elif paragraph.text.startswith('5 冻结协议'):
            paragraph.text = '5 V5规划参照与消融协议'
        elif paragraph.text.startswith('6 实验结果'):
            paragraph.text = '6 V5规划结果与收益来源'
        elif paragraph.text.startswith('7 计算成本'):
            paragraph.text = '7 V5完整规划的计算成本与边界'
    # Retain the established physical and V5 evidence sections, replace the old
    # prospective discussion and references with the completed V6 study.
    cut = next(p._p for p in doc.paragraphs if p.text.startswith('8 讨论与后续'))
    removing = False
    for element in list(doc.element.body):
        if element is cut:
            removing = True
        if removing and element.tag != qn('w:sectPr'):
            doc.element.body.remove(element)

    def p(text, style=None):
        paragraph = doc.add_paragraph(text, style)
        if style is None:
            paragraph.paragraph_format.first_line_indent = Cm(.74)
        return paragraph

    def h(text):
        doc.add_heading(text, 1)

    def eq(text):
        paragraph = doc.add_paragraph(); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        math = OxmlElement('m:oMath')
        def run(parent, value):
            r = OxmlElement('m:r'); t = OxmlElement('m:t'); t.text = value
            t.set(qn('xml:space'), 'preserve'); r.append(t); parent.append(r)
        pos = 0
        pattern = r'([A-Za-zΣ]+)(?:_(?:\{([^}]+)\}|([A-Za-z0-9]+)))?(?:\^(?:\{([^}]+)\}|([A-Za-z0-9*]+)))?'
        for token in re.finditer(pattern, text):
            run(math, text[pos:token.start()])
            low, high = token.group(2) or token.group(3), token.group(4) or token.group(5)
            if low is None and high is None:
                run(math, token.group(1))
            else:
                script = OxmlElement('m:sSubSup' if low and high else ('m:sSub' if low else 'm:sSup'))
                base = OxmlElement('m:e'); run(base, token.group(1)); script.append(base)
                for tag, value in [('m:sub', low), ('m:sup', high)]:
                    if value:
                        index = OxmlElement(tag); run(index, value); script.append(index)
                math.append(script)
            pos = token.end()
        run(math, text[pos:]); paragraph._p.append(math)

    def table(headers, rows):
        from docx.table import Table
        element = deepcopy(table_template); doc.element.body.insert(len(doc.element.body)-1, element)
        table = Table(element, doc._body)
        while len(table.rows) > 1:
            table._tbl.remove(table.rows[-1]._tr)
        for cell, value in zip(table.rows[0].cells, headers):
            cell.text = value
        for values in rows:
            for cell, value in zip(table.add_row().cells, values):
                cell.text = str(value)
        for i, row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(5)
                    paragraph.paragraph_format.line_spacing = 1.05
                    paragraph.paragraph_format.keep_with_next = i < len(table.rows)-1
                    for run in paragraph.runs:
                        run.font.size = Pt(10); run.bold = i == 0
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    h('8 高层方案学习')
    p('将V5的37项方案作为每回合的离散动作，学习器输入三车位置、速度、位置差及到冲突区的预计时间，共12维。'
      '共享编码器为两层128单元tanh网络，分别输出37维动作概率和37维相对基线代价。其集中式结构控制三辆车，三个训练模型是集成成员而不是车辆智能体。')
    p('训练使用768个独立常规初态的完整方案代价；128个验证初态用于选检查点，256个校准初态用于定阈值。'
      '前20轮为成本敏感软教师蒸馏，后续至180轮优化所有动作的期望相对代价与熵项，并回归代价差。'
      '三个随机种子分别训练，每10轮按验证集贪心代价选择最优检查点，同分保留最早者；测试不参与选择。')
    p('该训练利用模型生成的完整动作回报，是离线完整信息上下文bandit，不是在线试错PPO。保留第20轮蒸馏模型作为消融，'
      '最终策略将三个模型的动作概率与代价预测分别平均。在线只根据初态选方案，之后低层每步反馈执行并认证。')

    doc.add_page_break()
    h('9 相对基线的性能校准')
    p('记方案j到达代价为Jⱼ，传统基线为J₀，差值dⱼ=Jⱼ−J₀。对已经固定的学习提案a(z)，用m(z)表示代价头预测的相对基线差。'
      '第i个校准场景所提方案的真实差记为dᵢ，预测差为mᵢ，计算两者的单侧残差。模型、验证选择和集成规则均在校准之前固定。')
    eq('r_i = d_i − m_i,     k = ⌈(n+1)(1−α)⌉')
    p('设n=256，α=0.1，q为第k小残差；若k大于样本量则取正无穷。此有限样本秩修正不能简单替换成90%经验分位数。'
      f"本次q为{result['calibration_quantile_s']:.6f} s。以预测差加q构成上界，只有上界不大于零才采用学习提案，否则选择传统基线。")
    eq('U(z) = m(z) + q,     accept ⇔ U(z) ≤ 0')
    p('−U可以解释为性能置信裕量，单位为秒，阈值为零。它不是逐场景安全概率；低层的联合制动证书继续独立检查每个动作的可执行性。')
    p('性质3：若校准与新常规初态可交换，且执行代价等于校准所用精确模型代价，则最终切换方案比传统基线慢的边际概率不超过α。'
      '证明：split conformal秩论证给出P(dₐ>U)≤α。回退的代价差为零；若最终变慢，则必须接受且dₐ>0。接受时U≤0，因此该事件包含于dₐ>U。')
    p('该概率涉及校准集与新初态共同的随机性，不是固定某份校准数据后每个场景的保证，也不是接受子集内的条件保证。'
      '高密度分布改变初态分布，本项目不对其主张同一界；每回合反复重选或存在未知模型误差时也需重新论证。该推论使用标准conformal理论[4]。')
    h('10 独立测试与对照设计')
    p('新测试包含常规与高密度各60个初态，不复用V5测试场景。六种方法为传统基线、训练集最优固定方案、纯蒸馏集成、策略优化集成、'
      '校准门控集成和完整规划教师，共720次真实SUMO执行。原始37方案标签只用于配对核验与评价，不进入学习器在线决策。')
    p('同时报告平均截断到达代价、相对教师的收益保留、胜平负、非基线方案采用率、总体变慢率与接受后的条件变慢率。'
      '纯蒸馏对照检验后续策略优化是否有额外价值；训练集最优固定方案检验学习器是否只记住一个普遍有效的模板。')

    doc.add_page_break()
    h('11 学习收益与门控代价')
    names = [('constant', '传统基线'), ('fixed', '训练集最优固定方案'), ('imitation', '纯蒸馏选择器'),
             ('learned', '策略优化选择器'), ('calibrated', '校准门控选择器'), ('teacher', '完整规划教师')]
    p('表2  V6平均截断到达时间  单位为秒', 'Caption')
    table(['方法', '常规 60场景', '高密度 60场景'],
          [[name, f"{summary['nominal'][method]['mean_arrival_s']:.3f}", f"{summary['dense'][method]['mean_arrival_s']:.3f}"] for method, name in names])
    for split, label in [('nominal', '常规'), ('dense', '高密度')]:
        r = summary[split]['calibrated']; raw = summary[split]['learned']
        p(f"{label}分布中，未加性能门控的学习器保留教师{raw['teacher_gain_retained_percent']:.1f}%的均值收益；"
          f"校准后保留{r['teacher_gain_retained_percent']:.1f}%，平均节省{r['mean_saved_s']:.3f} s。"
          f"门控采用非基线方案{r['nonbaseline_adoptions']}/60次，最终变慢{r['wins_ties_losses'][2]}/60次，提案上界覆盖率{r['proposal_coverage']:.1%}。")
    picture = doc.add_paragraph(); picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.add_run().add_picture(str(D/'results.png'), width=Cm(16.3))
    p('图2  到达代价与相对完整规划教师的收益保留比例', 'Caption')
    for split, label in [('nominal', '常规'), ('dense', '高密度')]:
        r = summary[split]['calibrated']; ci = r['saved_ci95_s']
        conditional = '无采用样本' if r['conditional_regression_rate'] is None else f"{r['conditional_regression_rate']:.1%}"
        p(f"{label}校准门控平均节省时间的描述性95%配对区间为[{ci[0]:.3f}, {ci[1]:.3f}] s；"
          f"在采用非基线方案的子集内，条件变慢率为{conditional}。整体回退率会影响这两个指标，不能只报告总体退化率。")

    doc.add_page_break()
    h('12 工程验证与研究结论')
    p('表3将高层选择计入首次响应，完整规划教师的模型展开也计入其首次响应。周期用时包含当前证书、SUMO步进和指标更新；'
      '单机观测用时不等于WCET，也没有模拟计算期间车辆继续前行。离线1152个初态的标签生成成本与在线推断分开。')
    p('表3  V6运行时间P95  单位为毫秒', 'Caption')
    table(['方法', '首次响应', '执行周期'], [[name, f"{timing[m]['first_response_p95_ms']:.2f}", f"{timing[m]['cycle_p95_ms']:.2f}"]
                                             for m, name in [('learned', '策略优化选择器'), ('calibrated', '校准门控选择器'), ('teacher', '完整规划教师')]])
    total = sum(r['completed'] for split in summary.values() for r in split.values())
    collisions = sum(r['collisions'] for split in summary.values() for r in split.values())
    violations = sum(r['violating_episodes'] for split in summary.values() for r in split.values())
    p(f"720次执行完成{total}次，碰撞回合{collisions}次，路线合同违规回合{violations}次。实际选中方案的代价逐回合与独立模型标签一致。"
      '所有训练数据、模型、校准分位数及测试记录均保存，代码验证数据分离、验证检查点选择、校准秩和预先指定轨迹重放。')
    p('本轮实现并检验了性能置信切换：学习器选择有任务意义的高层方案，性能上界决定是否采用，物理证书继续约束低层。'
      '常规测试的上界覆盖率83.3%低于名义90%，不能宣称经验可靠度已达标。高密度上界覆盖率70%，门控仅保留28.4%教师收益。'
      '尚未实现分布式MARL、未知动力学下的鲁棒证书或实车验证。')
    p('Conformal策略切换已有研究[5,6]，因此模块组合与边际界本身不能作为首次创新。投稿仍需在更丰富的交互任务中证明稳定学习贡献，'
      '并与已发表方法做相同观测、动作和计算预算下的对照。')
    h('参考文献')
    references = [
        '[1] Bertsekas D. Multiagent Rollout Algorithms and Reinforcement Learning. arXiv:1910.00120. https://arxiv.org/abs/1910.00120',
        '[2] Banerjee A, et al. Dynamic Model Predictive Shielding for Provably Safe Reinforcement Learning. NeurIPS, 2024. https://arxiv.org/abs/2405.13863',
        '[3] Li Y, et al. On-Line Policy Iteration with Trajectory-Driven Policy Generation. arXiv:2604.15004, 2026. https://arxiv.org/abs/2604.15004',
        '[4] Angelopoulos A N, Bates S. A Gentle Introduction to Conformal Prediction and Distribution-Free Uncertainty Quantification. https://arxiv.org/abs/2107.07511',
        '[5] Conformal Policy Learning for Sensorimotor Control Under Distribution Shifts. https://arxiv.org/abs/2311.01457',
        '[6] Huriot S, Tabbara I, Sibai H. Safe Control using Learned Safety Filters and Adaptive Conformal Inference. L4DC, 2026. https://proceedings.mlr.press/v331/huriot26a.html']
    for reference_text in references:
        paragraph = p(reference_text); paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_after = Pt(3)
        for run in paragraph.runs:
            run.font.size = Pt(9)
    doc.save(out)
    assert hashlib.sha256(reference.read_bytes()).hexdigest() == source_hash
    allowed = {'word/document.xml', 'word/_rels/document.xml.rels', 'docProps/core.xml', '[Content_Types].xml'}
    with ZipFile(reference) as before, ZipFile(out) as after:
        changed = [name for name in before.namelist() if before.read(name) != after.read(name)]
        assert set(changed) <= allowed, changed
    print(json.dumps({'output': str(out), 'changed_existing_parts': changed}))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--reference', type=Path, required=True); parser.add_argument('--out', type=Path, required=True)
    args = parser.parse_args(); build(args.reference, args.out)
