"""Locally update the retained Word manuscript with audited final evidence."""
import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table
from lxml.etree import tostring

ROOT = Path(__file__).resolve().parents[1]
REFERENCE_SHA = '9bca1a2bc2711c26161e6d79c9c36ffee38a250967e68b8a3b8f03c799e868c0'
SEEDS = ('1802', '1812', '1822')


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def build(reference, out, evidence_out):
    if reference.resolve() == out.resolve() or digest(reference) != REFERENCE_SHA:
        raise ValueError('Use the unchanged retained fourteenth manuscript')
    root = ROOT / 'artifacts/v19'
    manifest = json.loads((root / 'final_report/sha256.json').read_text())
    if digest(root / 'final_report/report.json') != manifest['report.json']:
        raise ValueError('Final reporting identity changed')
    report = json.loads((root / 'final_report/report.json').read_text())
    if not report['original_report_values_reproduced'] or not report['real_whole_profile_replays_verified']:
        raise ValueError('Completed independent reporting and actual deployments are required')
    for relative, expected in report['input_sha256'].items():
        if digest(root / relative) != expected:
            raise ValueError('Final evidence changed: ' + relative)
    results = report['results']
    assert results['portfolio0']['passing_training_seeds'] == ['1812']
    assert not report['strict_gate_all_replicates_meet_original_targets']
    replay = json.loads((root / 'comparison_replay/audit.json').read_text())
    assert replay['actual_deployment_replays'] == 156 and replay['actual_shadow_sumo_runs'] == 10
    doc = Document(reference)
    original_section = doc.sections[0]._sectPr.xml
    original_math = [tostring(e, method='c14n', exclusive=True) for e in doc.element.body.iter(qn('m:oMath'))]
    prototype = deepcopy(doc.tables[0]._tbl)
    doc.core_properties.version = '15'

    def replace(prefix, content):
        matches = [p for p in doc.paragraphs if p.text.startswith(prefix)]
        if len(matches) != 1:
            raise ValueError('Ambiguous edit slot: ' + prefix)
        matches[0].text = content
        return matches[0]

    replace('在256个新初态、7,168次SUMO运行的冻结对照中，',
        '单车前置验证与三车模型式协同训练完成后，本文进一步冻结1,024个初态的40,960次风险校准和512个新初态的20,480次最终测试。'
        '严格联合门控在三个模型上保留0.67%–0.72%的团队成本改善及20.9%–22.3%的学习采用率，收益区间均为正；'
        '仅模型1812同时通过两项5%条件风险认证与预定收益、采用率要求。匹配选择器的156次实际部署回放验证了执行一致性。'
        '6,400次直接迁移揭示密集交通、紧密跟车收益及强惯性失败，完整调用仍不满足最坏情况实时要求。')
    replace('本文考察完整联合组合对实际采用、逐车性能和计算代价的影响，',
        '本文考察完整联合组合对实际采用、逐车性能和计算代价的影响，并检验有限扰动条件下的保证及动力学迁移边界。'
        '风险校准采用已有Learn then Test工具，性能分解属于已有rollout思想。最终确认显示严格门控保留实际收益，'
        '但只有一份训练模型满足全部预定条件，因此跨训练复现的风险稳定性仍是待解决的问题。')
    replace('共形预测与Learn then Test为不确定性和风险检验提供了工具',
        '共形预测与Learn then Test为不确定性和风险检验提供了工具[6,7]；MA-COPP处理多智能体策略改变后的联合轨迹分布[8]。'
        '阈值型安全策略改进与深度模型式改进已有相关研究[9,10]。本文的配对仿真观察候选与基线的完整结果。'
        '第12节给出同候选、同观测、同样本预算的共形、条件风险与均值改进适配及最终结果；这些适配的保证对象不同。')
    replace('校准从名义与密集分布的等权混合中抽取1,024个IID初态，',
        '校准从名义与密集分布的等权混合中抽取1,024个IID初态，完成40,960次SUMO；最终确认使用另512个IID初态，完成20,480次SUMO。'
        '全部控制器、阈值、样本量与选择规则事先冻结。联合要求为两项风险认证、至少20%的实际采用率、至少0.5%的团队成本改善，'
        '以及正的配对收益区间下界。三个训练模型分别报告，不能用最佳模型代替全部复现。')
    replace('已有结果支持联合检查的必要性及其在本任务中的实际收益。',
        '独立最终结果支持联合检查在本任务中的实际收益。名义门控取得较高平均收益，但采用后逐车退化仍较多；'
        '严格扰动组合将主统计中的经验退化比例降至0.88%–2.80%，保留0.67%–0.72%的平均成本改善。'
        '这些测试频率不同于校准风险上界，只有模型1812满足全部预定条件，不能宣称三个训练复现均成功。')
    replace('理论保证分为三个层次：',
        '可恢复物理安全依赖误差与初始集合假设；逐车性能不退化依赖精确模型和有限扰动序列；采用后的条件风险依赖独立统计校准及声明的总体分布。'
        '本次校准仅认证一份严格模型，零采用不能形成学习条件风险证书。迁移实验中的强惯性超出物理假设，'
        '因此正确回退也不能消除基线自身的失配失败。')
    replace('后续结论应以冻结校准和最终确认结果为准。',
        '保守共形适配在最终测试全部回退基线；LTT仅在模型1812部署严格门控，其余模型回退。均值改进选择器保留收益，却不提供本研究的采用后风险保证。'
        '这些结果比较具体保证对象和部署规则，不证明严格门控普遍优于共形方法，也不能把使用LTT认证写成对LTT的替代。')
    replace('本文将统计校准、策略改进恒等式和模型式多智能体训练作为已有工具。',
        '本文将统计校准、策略改进恒等式和模型式多智能体训练作为已有工具。研究证据集中于联合组合的采用机制、逐车相对性能约束及其风险与收益权衡。'
        '后续时变路径和动作调整策略使用独立开发数据；其代码与协议在查看本次最终结果前已冻结，仍须新的风险校准与最终确认。'
        '当前未获得跨模型一致认证或硬实时保证。')

    chapter13 = next(p for p in doc.paragraphs if p.text == '13 跨场景与动力学直接迁移')
    anchor = chapter13._p.getprevious()  # Preserve its original page break.

    def place(element):
        anchor.addprevious(element)

    def p(text):
        paragraph = doc.add_paragraph(text)
        paragraph.paragraph_format.first_line_indent = Cm(.74)
        place(paragraph._p)
        return paragraph

    def page(title):
        place(doc.add_page_break()._p)
        place(doc.add_heading(title, 2)._p)

    def table(headers, rows, widths):
        element = deepcopy(prototype)
        place(element)
        t = Table(element, doc._body)
        while len(t.rows) > 1:
            t._tbl.remove(t.rows[-1]._tr)
        while len(t.columns) < len(headers):
            t.add_column(Cm(2.5))
        assert len(t.columns) == len(headers) and abs(sum(widths) - 16.4) < 1e-9
        t.autofit = False
        for cell, value in zip(t.rows[0].cells, headers):
            cell.text = value
        fill = deepcopy(t.rows[0].cells[0]._tc.get_or_add_tcPr().find(qn('w:shd')))
        for cell in t.rows[0].cells:
            props = cell._tc.get_or_add_tcPr()
            old = props.find(qn('w:shd'))
            if old is not None:
                props.remove(old)
            if fill is not None:
                props.append(deepcopy(fill))
        for values in rows:
            for cell, value in zip(t.add_row().cells, values):
                cell.text = str(value)
        for col, width in zip(t.columns, widths):
            col.width = Cm(width)
        for index, row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for para in cell.paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    para.paragraph_format.space_before = para.paragraph_format.space_after = Pt(5)
                    para.paragraph_format.line_spacing = 1.05
                    para.paragraph_format.keep_with_next = index < len(t.rows) - 1
                    for run in para.runs:
                        run.font.size = Pt(10)
                        run.bold = index == 0
        blank = doc.add_paragraph()
        blank.paragraph_format.space_after = Pt(2)
        place(blank._p)

    page('独立风险校准与最终确认结果')
    p('全部校准记录完成后才计算风险上界，并由独立程序重建18项检验。下表报告三个严格模型的单侧上界，'
      '对应总体错误认证预算5%在九候选与两个目标之间的固定分配。两项上界均不超过5%才通过。')
    calibration_rows = []
    for seed in SEEDS:
        record = report['calibration_reconstruction'][f'portfolio0_{seed}']
        actual, bundle = record['actual_random_mode'], record['worst_mode_bundle']
        calibration_rows.append([seed, f"{100 * actual['threshold_upper']:.4f}%", f"{100 * bundle['threshold_upper']:.4f}%",
                                 '通过' if actual['certified'] and bundle['certified'] else '未通过'])
    table(['模型', '实际条件风险上界', '组合条件风险上界', '两项同时认证'], calibration_rows, [2., 4.8, 4.8, 4.8])
    p('模型1802的5.0340%严格大于5%，不能四舍五入判为通过。模型1822仅通过实际模式指标；模型1812同时通过两项。'
      '未认证不等于已经证明真实风险超标，但也不能据此给出风险保证。认证不限制单次退化幅度，且不覆盖任意噪声或分布迁移。')
    p('独立最终测试按预先固定的随机流，每个初态选一种实际扰动模式作主统计。下表成本改善相对同初态、同扰动基线，'
      '收益区间单位为成本秒；采用列上行为采用数／512，下行为采用后退化数／采用数。')
    final_rows = []
    for seed in SEEDS:
        record = results['portfolio0']['per_training_seed'][seed]
        m = record['metrics']
        a = m['actual_random_mode']
        lo, hi = m['mean_team_gain_ci95']
        final_rows.append([seed, f"{m['mean_team_gain_percent']:.3f}%", f'[{lo:.5f}, {hi:.5f}]',
                           f"{a['adopted']}/512\n{a['harmful_adopted']}/{a['adopted']}",
                           '全部通过' if record['target_check']['all_conditions_met'] else '风险未认证'])
    table(['模型', '成本改善', '95%收益区间 秒', '采用与退化', '联合目标'], final_rows, [1.4, 2.6, 4.2, 4.6, 3.6])
    p('三份模型均保留至少20%的采用率、至少0.5%的平均收益及正的收益区间。主统计中严格门控均完成512/512次运行，'
      '未采用学习的结果与配对基线一致。四模式组合的经验退化／采用计数依次为10/251、10/237、11/230。'
      '最终经验比例不能改变原校准结论，联合目标仍仅模型1812通过。')

    page('已有准入规则的匹配对照结果')
    p('以下均在同一512个独立初态上比较。先在每个初态内平均三份固定模型，再按初态计算配对自举区间；'
      '三份模型、三辆车和四类扰动不会增加独立样本数。均值区间为近似区间，不能替代条件风险认证。')
    labels = ('原始学习', '名义门控', '严格联合门控', 'Conformal适配', 'LTT适配', 'Mean SPI适配')
    rows = []
    for method, label in zip(('raw', 'nominal', 'portfolio0', 'conformal', 'ltt', 'mean_spi'), labels):
        record = results[method]['fixed_model_average']
        lo, hi = record['mean_team_gain_ci95']
        rows.append([label, f"{record['mean_team_gain_seconds']:.5f}", f'[{lo:.5f}, {hi:.5f}]'])
    table(['方法', '平均团队收益 秒', '95%配对区间 秒'], rows, [5.2, 4.8, 6.4])
    p('名义门控在三个模型上的成本改善为1.468%–1.721%，但采用后的经验退化比例为16.40%–17.78%。'
      '严格门控牺牲部分收益与采用率以降低退化；这种权衡不等于在所有指标上支配其他方法。原始学习模型1802的收益区间为'
      '[−0.06533, 0.10459]秒，跨越零，主统计有3次未完成；模型1812有1次未完成，负结果全部保留。')
    p('LTT在模型1812选择严格门控，在另外两模型回退基线。本文采用LTT进行认证，因此不能将其作为被本文替代的竞争算法。'
      'Mean SPI分别选择名义模型1802、原始模型1812和1822，保留均值收益但未提供采用后的条件风险保证；它是全信息均值检验适配。')
    p('本次保守Conformal适配全部回退基线，产生零采用与零收益。零采用时条件风险未定义，不能写成零条件风险。'
      '这一结果只描述当前共同残差与完整控制器适配，不否定其他共形预测设计；状态相关选择也不能继承某个固定策略的条件风险证书。')
    p('预先冻结的覆盖规则产生156次实际完整控制器部署回放，另执行10次名义影子SUMO。全部被选控制器、逐车指标与采用次数一致，'
      '156条新轨迹均重新检查物理状态。部署选择只读取校准记录和观测初态名义预测，不读取实际未来扰动。回放与影子预测不增加独立样本。')
    p('主研究原冻结严格门控的P95为585.61–640.30 ms，P99为1731.80–1849.07 ms，仍超过500 ms控制周期。'
      '共形适配还需承担整套名义预测准备开销；并行桌面负载下的计时不构成独占部署基准。完整数据和独立报告见公开仓库[12]。')

    assert doc.sections[0]._sectPr.xml == original_section
    assert [tostring(e, method='c14n', exclusive=True) for e in doc.element.body.iter(qn('m:oMath'))] == original_math
    doc.save(out)
    assert digest(reference) == REFERENCE_SHA
    with ZipFile(reference) as old, ZipFile(out) as new:
        assert set(old.namelist()) == set(new.namelist())
        changed = [name for name in old.namelist() if old.read(name) != new.read(name)]
        assert set(changed) <= {'word/document.xml', 'docProps/core.xml'}, changed
    record = dict(document_name=out.name, document_sha256=digest(out), reference_sha256=REFERENCE_SHA,
                  changed_parts=changed, retained_native_equations=len(original_math), source_sha256=digest(__file__),
                  final_report_sha256=digest(root / 'final_report/report.json'), final_comparison_completed=True,
                  real_deployment_replays=156, shadow_sumo_runs=10, all_models_meet_joint_targets=False,
                  passing_strict_models=['1812'], v20_final_risk_certified=False, hard_real_time_certified=False,
                  final_outcomes_used_to_tune_v20=False, expected_pages='approximately 19; render and inspect before delivery')
    evidence_out.parent.mkdir(parents=True, exist_ok=True)
    evidence_out.write_text(json.dumps(record, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(record, ensure_ascii=False))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--reference', type=Path, required=True)
    parser.add_argument('--out', type=Path, required=True)
    parser.add_argument('--evidence-out', type=Path, required=True)
    args = parser.parse_args()
    build(args.reference, args.out, args.evidence_out)
