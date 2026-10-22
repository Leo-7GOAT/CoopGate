"""Confidence-gating manuscript using the retained V11 DOCX design."""
import argparse,json,hashlib,re
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm,Pt
from docx.table import Table
ROOT=Path(__file__).resolve().parents[1]


def build(reference,out):
    digest=hashlib.sha256(reference.read_bytes()).hexdigest()
    assert digest=='0c8662d7c60f2557384163b6a07c21fdf1a85b3eafb4d6a2bbfa689f59275fb9'
    doc=Document(reference);template=deepcopy(doc.tables[0]._tbl)
    s=json.loads((ROOT/'artifacts/v13/ego/summary.json').read_text());base=s['methods']['rollout'];av=s['seed_average']
    perf=s['performance'];gate=s['gate_diagnostics'];ci=s['paired_initial_cluster_intervals']
    title='相对规划基线的学习控制置信门控\n自车双指标校准与可恢复安全验证'
    doc.paragraphs[0].text=title;doc.core_properties.title=title.replace('\n',' ');doc.core_properties.version='12'
    cut=doc.paragraphs[3]._p;remove=False
    for e in list(doc.element.body):
        if e is cut:remove=True
        if remove and e.tag!=qn('w:sectPr'):doc.element.body.remove(e)

    def p(text,style=None):
        a=doc.add_paragraph(style=style);pos=0
        for token in re.finditer(r'([A-Za-zΑ-Ωα-ω]+)_([A-Za-z]+)',text):
            a.add_run(text[pos:token.start()]);math=OxmlElement('m:oMath');sub=OxmlElement('m:sSub')
            for tag,value in [('m:e',token.group(1)),('m:sub',token.group(2))]:
                element=OxmlElement(tag);run=OxmlElement('m:r');literal=OxmlElement('m:t');literal.text=value
                run.append(literal);element.append(run);sub.append(element)
            math.append(sub);a._p.append(math);pos=token.end()
        a.add_run(text[pos:])
        if style is None:a.paragraph_format.first_line_indent=Cm(.74)
        return a
    def h(text):doc.add_heading(text,1)
    def page(text):doc.add_page_break();h(text)
    def eq(text):
        a=doc.add_paragraph();a.alignment=WD_ALIGN_PARAGRAPH.CENTER;m=OxmlElement('m:oMath')
        def run(parent,value):
            z=OxmlElement('m:r');t=OxmlElement('m:t');t.text=value;t.set(qn('xml:space'),'preserve');z.append(t);parent.append(z)
        pos=0
        for token in re.finditer(r'([A-Za-zΑ-Ωα-ωΣΠ≤≥0-9]+)(?:_\{([^}]+)\}(?:\^\{([^}]+)\})?|\^\{([^}]+)\})',text):
            run(m,text[pos:token.start()]);sub=token.group(2);sup=token.group(3) or token.group(4)
            z=OxmlElement('m:sSubSup' if sub and sup else 'm:sSub' if sub else 'm:sSup')
            b=OxmlElement('m:e');run(b,token.group(1));z.append(b)
            if sub:
                t=OxmlElement('m:sub');run(t,sub);z.append(t)
            if sup:
                t=OxmlElement('m:sup');run(t,sup);z.append(t)
            m.append(z);pos=token.end()
        run(m,text[pos:]);a._p.append(m)
    def table(headers,rows,widths):
        element=deepcopy(template);doc.element.body.insert(len(doc.element.body)-1,element);t=Table(element,doc._body)
        while len(t.rows)>1:t._tbl.remove(t.rows[-1]._tr)
        while len(t.columns)<len(headers):t.add_column(Cm(2.5))
        t.autofit=False
        for cell,value in zip(t.rows[0].cells,headers):cell.text=value
        fill=deepcopy(t.rows[0].cells[0]._tc.get_or_add_tcPr().find(qn('w:shd')))
        if fill is not None:
            for cell in t.rows[0].cells:
                props=cell._tc.get_or_add_tcPr();old=props.find(qn('w:shd'))
                if old is not None:props.remove(old)
                props.append(deepcopy(fill))
        for values in rows:
            for cell,value in zip(t.add_row().cells,values):cell.text=str(value)
        for col,width in zip(t.columns,widths):col.width=Cm(width)
        for i,row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell,width in zip(row.cells,widths):
                cell.width=Cm(width)
                for a in cell.paragraphs:
                    a.alignment=WD_ALIGN_PARAGRAPH.CENTER;a.paragraph_format.space_before=a.paragraph_format.space_after=Pt(5)
                    a.paragraph_format.line_spacing=1.05;a.paragraph_format.keep_with_next=i<len(t.rows)-1
                    for run in a.runs:run.font.size=Pt(10);run.bold=i==0
        doc.add_paragraph().paragraph_format.space_after=Pt(2)

    p('针对学习控制器在不同交通交互中收益不均、单一综合指标可能掩盖自车延误的问题，本文研究相对传统规划基线的置信门控。'
      '以自车到达时间及到达前累计指令变化量作为任务目标，分别预测学习方案与规划基线的完整闭环结果。'
      '利用独立校准回合的两维最大残差构造成本和时间的联合上界，仅在成本具有改善余量且到达时间不退化时放行学习方案。'
      '未通过性能门控时采用规划基线，全部执行指令仍由相同的可恢复制动证书检查。')
    pa=perf['augmented'];pg=perf['gate'];ga=gate['all']
    p(f"在三组独立Double DQN训练、1024次独立SUMO校准及7168次冻结后对照中，学习辅助规划器相对强规划基线，"
      f"自车综合代价降低{pa['ego_cost_gain_percent']:.2f}%，到达时间降低{pa['ego_time_gain_percent']:.2f}%，累计指令变化量降低{pa['ego_variation_gain_percent']:.2f}%。"
      f"门控放行率为{ga['acceptance_rate']*100:.2f}%，有害放行占全部策略回合的{ga['joint_harm_rate']*100:.2f}%。"
      '本文将纯网络执行与学习辅助规划分别评价，统计风险结论与物理安全结论分别陈述。')
    p('关键词：自动驾驶；强化学习；控制器切换；共形校准；可恢复安全；SUMO')
    h('1 引言与研究问题')
    p('学习型决策有机会利用复杂交互中的行动规律，但其输出分数不能直接解释为“比传统规划更好”的概率。'
      '对于简单场景，学习方案可能没有可测的额外收益；对于失配场景，即使网络给出确定动作，实际性能仍可能退化。'
      '本文关注的是可解释的采用条件与明确的基线退回机制，而不是强制让学习控制器覆盖所有场景。')
    p('研究包含三个可分开的判断：学习组件是否贡献了实际收益；性能门控是否保留收益并减少退化；执行指令是否满足物理约束。'
      '“复杂场景收益更大”仅作为分组检验假设，不预先当作结论。本文使用三车固定路线合流和一个可学习车辆作为前置验证，尚不声称完成多智能体强化学习。')
    p('已有共形策略学习研究利用分位数切换基础策略，ARMS也探索按环境融合强化学习与解析控制[5,6]。'
      '因此，“RL加传统控制切换”本身不足以作为创新。本文目前的具体工作是明确基线相对的自车双指标退化事件，将其与逐周期安全证书结合，并给出可复核的训练、冻结测试和消融证据。')

    page('2 控制对象与序列强化学习')
    p('环境为两条入口汇入一条出口的固定路线合流，三车路线标识为[0,1,0]。匝道车辆ego1参与学习，其余车辆固定提出最大允许加速度。'
      '公共安全层可以对任一车辆制动，因此这是全受控协同仿真，不是其他驾驶员不可控的开放交通。决策周期0.5秒，物理子步0.1秒，任务时域25秒。')
    eq('J_{e}=T_{e}+λD_{e},   D_{e}=Σ_{k=0}^{Kₑ−1}|u_{e,k}−u_{e,k−1}|')
    p('Kₑ为到达目标的控制周期数，Tₑ=0.5Kₑ秒并截断于任务时域；初始指令为零。Dₑ仅累计自车到达前的指令变化，避免把其他车辆完成前的自车动作混入目标。'
      'λ=0.05s³/m，令综合代价以秒计。未完成、碰撞和违规单独记录，不能用较低的变化量掩盖停车失败。其他两车的平均到达时间单独报告。')
    p('学习器使用dueling Double DQN[1,2]。18维观测包含本车位置、速度、上一执行指令、路线、到达标识和时钟，以及150米内两个邻车的相对位置、相对速度、路线关系、上一指令和可见性。'
      '邻车指令来自本协同仿真的共享状态。此前观测遗漏该记忆，会将安全层转移不同的状态映射成相同输入；本版本显式补入。')
    eq('Q(o,a)=V(o)+A(o,a)−mean_{b}A(o,b)')
    eq('y=r+γ(1−z)Q_{target}(o′,argmax_{b}Q_{online}(o′,b)),   γ=0.995')
    p('z为到达目标或触及25秒边界的终止标志。网络包含两个128维ReLU隐藏层，并分离状态价值与动作优势。动作是−5至2.5m/s²、间隔0.625m/s²的13个加速度提案，先屏蔽违反指令变化率的动作。'
      '回放容量262144，批量256，Adam学习率0.0003，目标网络每250次梯度更新同步。探索率在600000步内由1降至0.03。')
    p('三个新训练种子701、711、721分别进行2097152次模型环境转移，最终检查点全部保留，不挑选中途最优模型，无行为克隆或规划器标签。'
      '每次转移都包含当前提案、安全修正、下一个状态和奖励。Q值对应原始提案，而非把修正后指令错误地当作采样动作。')
    p('奖励使用负的自车阶段时间与指令变化代价，并加入γΦ(x′)−Φ(x)形式的势函数塑形；Φ为负的归一化剩余路线距离，终止势为零。'
      '其折扣和对固定初态仅增加常数，但折扣学习目标与报告的未折扣工程目标仍有差别。训练后端为精确纵向模型，不计作SUMO运行次数。')

    page('3 安全合同与传统规划基线')
    p('纵向子步模型采用速度截断后的前向积分，执行误差w逐子步满足已知界。状态中保留上一执行指令，以反映有限指令变化率。')
    eq('v_{i,t+1}=clip(v_{i,t}+(u_{i,k}+w_{i,t})Δt,0,20)')
    eq('p_{i,t+1}=p_{i,t}+v_{i,t+1}Δt,   |w_{i,t}|≤0.5')
    eq('−5≤u_{i,k}≤2.5,   |u_{i,k}−u_{i,k−1}|≤2.5')
    p('安全合同要求同路车辆保持6.8米间距，不同入口车辆的冲突区[−25,6.8]占用不重叠。检查覆盖物理子步内的连续路线区间，而非只比较决策时刻的位置。'
      '该合同是声明模型下的路线安全约束，不等同于任意道路几何与车辆碰撞包络。')
    p('每辆车提供提案、保持上一指令和备用制动三个候选，形成最多27个联合顶点。按与原提案的平方偏差排序，选择第一个通过鲁棒备用证书的联合指令。'
      '备用制动按记忆递推u←max(u−2.5,−5)，在区间模型中检查当前动作及随后制动序列，直至全部速度上界为零，且指令足以抵消最大正向误差。')
    p('命题1：在精确测量、声明动力学、已知误差界、可执行制动与指令变化率、初态可恢复的条件下，获证书的当前动作保留下一周期的备用可恢复性。'
      '证明要点是将原证书中的剩余备用序列作为后继状态的备用方案，并以保持静止的终端动作续接。实际后继状态包含于预测区间，故合同可递归保持。'
      '学习策略被放行并不解除这项逐周期检查。')
    p('传统基线B每周期枚举相同的自车动作集合，其余两车仍使用固定提案。候选经过同一联合证书后，模拟规则反馈直至目标或时域结束，选择自车综合代价最小的首动作。'
      '下一周期重新测量并计算。这是rollout规划器，不将其称为某篇已发表MPC系统的完整复现。固定规则控制器另作辅助对照。')
    p('基线的安全性来自相同的可恢复证书及其条件，不来自“传统算法一定安全”的假设。超出合同、感知失效、误差超过0.5m/s²等情形不在本证书承诺范围。')

    page('4 学习辅助规划与收益归因')
    p('独立网络策略N直接输出逐周期动作提案。学习辅助规划器L额外模拟N产生的未来反馈序列，与传统规则序列共同评价候选首动作。'
      '两者使用同一套已训练权重，但执行机制不同。因此，本文同时报告纯网络N与学习辅助规划L，不将在线搜索带来的全部收益归因于网络。')
    eq('W(x)=min(V_{0}(x),V_{N}(x))')
    eq('μ(x)=argmin_{a}[ℓ(x,a)+W(f(x,a))]')
    p('V₀和V_N分别是规则尾部与学习尾部的名义剩余代价；状态含指令记忆和时钟。所有候选和尾部指令都经过同一安全修正器。'
      'L保留传统单步rollout得到的当前最优方案，只在学习尾部得到严格更低预测代价时更新；相同代价时保留传统方案。')
    p('命题2：在完全匹配的名义模型、有限时域终端代价为零、基础尾部可行且其首动作包含在候选集内时，J_L(x)≤W(x)。'
      '取实现W(x)的基础策略的首动作，由W在后继状态不大于该基础策略的剩余代价，有minₐ[ℓ+W(f)]≤W(x)。'
      '沿闭环逐步相加并利用终端值为零，即得该结论。这是标准多策略rollout的政策改进性质[3]。')
    p('这不推出J_L≤J_B：传统基线B本身也进行了滚动改进，上述式子没有给出两个改进策略之间的排序。'
      '执行误差下的实际性能也不能由名义不等式替代。因此，L相对B是否更优仍由独立对照及下一节的性能门控检验。')
    table(['执行方式','是否训练','在线决策'],[
        ['规则控制','否','固定提案＋安全层'],['规划基线B','否','规则尾部rollout'],
        ['纯网络N','是','网络提案＋安全层'],['未训练网络辅助规划','否','同构额外尾部'],
        ['学习辅助规划L','是','训练后额外尾部']], [5.4,3.0,8.0])
    p('未训练网络对照与L具有相同的网络结构、候选集合和额外尾部预算。该消融用于辨别训练是否产生额外价值，'
      '并不取代纯网络相对强基线的前置检验，也不构成与已发表完整自动驾驶系统的同条件比较。')

    page('5 双指标置信门控及统计界')
    p('在合流事件开始时，用名义模型分别模拟L、B的完整反馈闭环，预测相对成本差和相对到达时间差。预测不使用将来的误差带。'
      '独立校准数据在相同初态和同一外生执行误差下分别运行两种控制器，得到实际差值。预测向量记为m，实际差值记为d。')
    eq('d=(J_{L}−J_{B}, T_{L}−T_{B}),   R=max(d_{J}−m_{J},d_{T}−m_{T})')
    p('两维均以秒计。初态按冲突区ETA差预先分组：自车与任一主路车的最小ETA差不超过1.25秒为高冲突组，否则为低冲突组。'
      '在每组n个校准分数中取k=ceil((n+1)(1−α))阶统计量q，α=0.1；k>n时取正无穷。')
    eq('U_{J}=m_{J}+q,   U_{T}=m_{T}+q')
    eq('A={U_{J}≤−0.05 s 且 U_{T}≤0 s}')
    p('只有A成立时才采用L，否则采用B。超出已声明的初态校准范围或预测失败同样返回B。'
      '这是性能门控，发生在每次合流开始时；选定控制器随后每0.5秒反馈重算，物理安全证书仍逐周期检查。')
    p('命题3：若训练模型、预测规则和分组固定，校准与新回合在组内可交换，则P{d_J≤U_J且d_T≤U_T}≥1−α。'
      '证明：由共形秩对称性，新分数不超过第k阶校准分数的概率至少为1−α；最大残差不超过q等价于两维同时被覆盖[4]。')
    eq('P(A ∩ ({d_{J}>0} ∪ {d_{T}>0}))≤α')
    p('若A成立而成本或到达时间发生退化，则至少一维差值超过对应上界，故有害放行事件包含在联合不覆盖事件中。'
      '同一论证也覆盖“未达到0.05秒成本改善余量或到达时间退化”的事件。')
    p('该结论是对校准样本与新回合共同随机性的边际风险界，不是给定当前校准集后的条件保证，不是放行子集内错误率≤10%，也不是碰撞概率。'
      '固定测试集的经验覆盖率可能低于90%，应照实报告。任意域外偏移、同回合反复复用阈值或事后挑选模型均不直接满足这一结论。')

    page('6 冻结实验与复核设计')
    p('训练场景在常规和高密度初态间等概率采样，训练误差为每子步独立的±0.5m/s²均匀误差。'
      '正式测试使用无误差、均匀误差、恒定偏置和间歇脉冲四种外生误差带。所有对照共享同一初态和同一误差带，不根据控制器轨迹重新生成环境扰动。')
    table(['阶段','规模','用途'],[
        ['模型训练','3×2097152步','三个固定最终策略'],['独立校准','256初态／1024执行','拟合两维联合余量'],
        ['冻结测试','128初态／7168执行','14方法×4种误差'],['轨迹审计','前2初态全部方法','112条完整轨迹']], [5.2,5.2,6.0])
    p('训练种子、预算、最终检查点使用规则先行登记；校准初态位于31000000至32000000之间，测试初态位于32000000至33000000之间且互不重叠。'
      '源码、模型、校准阈值、测试初态和11项验收在提交3bac214中冻结，随后才执行正式对照。开发集不计入独立测试。')
    p('14种方法包括规则和强规划两个基线，以及三个训练种子分别对应的纯网络、未训练网络辅助规划、训练后网络辅助规划与双指标门控。'
      '未校准门控、仅成本校准门控的结果由已经成对执行的完整轨迹重组，属于门控决策消融，不虚增SUMO执行次数。')
    p('每个回合保留完成、碰撞、合同违规、不可恢复、未认证指令、中止、各车到达时间、各车变化量及计算时间。'
      '正式运行采用一个SUMO工作进程，方法顺序按场景随机打乱；预先暖机后计时。源代码与训练哈希、权重变化、阈值重算、逐子步证书和指令回放分别复核。')
    p('统计以128个独立初态为簇，先平均同初态的四类误差和三个固定模型，再配对重采样10000次，报告95%区间。'
      '这些是近似描述性区间，不把1536个模型—误差组合误称为独立交通场景，也不代表对训练随机种子总体的精确推断。')
    p('辅助对照使用团队平均目标，另有5632次独立SUMO执行。在团队代价改善的同时，纯学习自车到达时间增加8.26%，门控自车增加2.94%。'
      '这说明团队平均指标不能替代自车非退化要求，也为将自车时间单独加入门控提供实验动机。该组数据独立归档，不与本节主测试混合。')

    page('7 单车独立测试结果')
    p('表1列出全部四类误差下的结果。除规划基线外，各学习类别均对三个预先固定的模型平均。每个类别包含1536个执行组合，独立初态数仍为128。')
    table(['方法','自车到达 s','变化量 m/s²','综合代价 s'],[
        ['强规划基线',f"{base['ego_time']:.4f}",f"{base['ego_variation']:.4f}",f"{base['ego_cost']:.4f}"],
        *[[name,f"{av[mode]['ego_time']:.4f}",f"{av[mode]['ego_variation']:.4f}",f"{av[mode]['ego_cost']:.4f}"]
          for mode,name in [('raw','纯Double DQN'),('untrained','未训练网络辅助'),('augmented','训练后网络辅助'),('gate','双指标门控')]]], [5.3,3.7,3.7,3.7])
    p(f"学习辅助规划L相对B的综合代价降低{pa['ego_cost_gain_percent']:.3f}%，到达时间降低{pa['ego_time_gain_percent']:.3f}%，"
      f"变化量降低{pa['ego_variation_gain_percent']:.3f}%。相对同构未训练网络对照，代价降低{s['learning_gain_vs_untrained_percent']:.3f}%，"
      '后者用于隔离训练的额外贡献。')
    p('表2  学习辅助规划相对基线的配对初态簇改善区间')
    table(['指标','平均改善','95%区间'],[[name,f"{ci['augmented'][key]['mean_improvement']:.4f}",
        f"[{ci['augmented'][key]['ci95'][0]:.4f}, {ci['augmented'][key]['ci95'][1]:.4f}]"]
        for key,name in [('ego_time','自车时间 s'),('ego_variation','指令变化 m/s²'),('ego_cost','综合代价 s')]], [6.0,4.0,6.4])
    p(f"其他两车平均到达时间为：基线{base['other_time']:.4f}秒、学习辅助规划{av['augmented']['other_time']:.4f}秒、门控{av['gate']['other_time']:.4f}秒。"
      '该指标用于观察是否把成本转移给其他参与者，不能用团队平均值替代自车主指标。')
    p(f"学习辅助规划和门控分别完成{av['augmented']['completed']}/{av['augmented']['n']}与{av['gate']['completed']}/{av['gate']['n']}回合，"
      f"各记录{av['augmented']['collisions']}与{av['gate']['collisions']}次碰撞、{av['augmented']['violations']}与{av['gate']['violations']}次合同违规。"
      '额外8次检查点缺失或运行中异常的实际SUMO接管均完成，单独验证工程故障路径，不纳入无故障性能校准结论。')
    if s['individual_ego_verified']:
        p('纯网络前置条件通过当前独立测试，但该结果仅限本场景、动作接口与安全层。扩大到多智能体时仍需重新训练、复核和比较。')
    else:
        p('纯网络前置条件未通过。即使学习辅助规划器获得改善，也不能写成纯强化学习策略已经超过强规划基线，'
          '不能以此解除MARL扩展门槛。本文保留两种执行方式的完整数据，继续将其作为未完成的研究条件。')
    ac=s['acceptance'];failed=[k for k,v in ac['checks'].items() if not v]
    p(f"学习辅助规划与门控的冻结验收通过{ac['passed']}/{ac['total']}项。"+
      ('未通过项目：'+', '.join(failed)+'。预先阈值未修改。' if failed else '全部预先声明的此类验收通过；它们不替代纯网络前置条件。'))

    page('8 门控收益保留与消融')
    p('有害放行定义为：已选择学习辅助规划，而实际自车成本或到达时间任一项比基线更差。'
      '统计以同初态、同误差的成对实际执行为依据。门控轨迹另与其所选控制器核对，不仅计算纸面选择。')
    table(['分组','放行率 %','有害占全部 %','保留净收益 %'],[[name,
        f"{gate[group]['acceptance_rate']*100:.2f}",f"{gate[group]['joint_harm_rate']*100:.2f}",
        '无正收益' if gate[group]['retained_gain_percent'] is None else f"{gate[group]['retained_gain_percent']:.2f}"]
        for group,name in [('all','全部'),('low_conflict','低冲突'),('high_conflict','高冲突')]], [5.3,3.7,3.7,3.7])
    p(f"全部策略回合中的有害放行为{ga['harmful_admissions']}/{ga['n']}。始终采用学习辅助规划时，退化比例为{ga['always_augmented_harm_rate']*100:.2f}%。"
      f"门控对强基线的成本改善为{pg['ego_cost_gain_percent']:.3f}%，时间改善为{pg['ego_time_gain_percent']:.3f}%。")
    p('两类冲突组仅表示当前ETA关系，并不代表完整城市道路的难度分级。不同分组的样本量和收益分布不同，不能由放行率较高直接推出学习在所有复杂场景均更优。')
    rr=[r for r in s['counterfactual_gate_ablations'] if r['group']=='all']
    labels={'uncalibrated_cost':'未校准成本门控','calibrated_cost_only':'仅成本校准','joint_cost_time':'成本＋时间联合校准'}
    table(['门控变体','放行率 %','有害占全部 %','保留成本收益 s'],[[labels[r['variant']],
        f"{r['acceptance_rate']*100:.2f}",f"{r['joint_harm_rate']*100:.2f}",f"{r['mean_retained_cost_gain']:.4f}"] for r in rr], [5.3,3.7,3.7,3.7])
    p('两种成本门控消融由已执行的L、B完整轨迹重组，双指标门控额外完成实际执行。'
      '同一初态的多个模型或扰动不作为独立样本；门控减少放行带来的收益损失与风险下降同时报告。')
    cost_only=next(r for r in rr if r['variant']=='calibrated_cost_only');joint=next(r for r in rr if r['variant']=='joint_cost_time')
    p(f"相对仅成本校准，联合时间约束将有害放行从{cost_only['joint_harmful_admissions']}次降至{joint['joint_harmful_admissions']}次，额外改进幅度有限且损失部分收益。"
      '其统计约束的含义更明确，但当前数据不足以证明其显著优于仅成本校准。')
    conditional='无放行' if ga['conditional_harm_rate'] is None else f"{ga['conditional_harm_rate']*100:.2f}%"
    p(f"联合上界的经验覆盖率为{ga['empirical_joint_coverage']*100:.2f}%，低冲突组为{gate['low_conflict']['empirical_joint_coverage']*100:.2f}%，"
      f"高冲突组为{gate['high_conflict']['empirical_joint_coverage']*100:.2f}%。放行子集内的实际退化比例为{conditional}。"
      '这些量的分母和含义不同，不能把α=0.1写成条件退化率或碰撞概率。')

    page('9 工程成本与研究局限')
    p('表3给出单机暖机条件下的计算分位数。设置时间包含加载与门控预测；控制周期时间是算法选择指令的耗时。'
      '该表分别展示门控一次性决策的代价与后续逐周期执行成本，不把两者混为统一的实时性指标。')
    table(['方法','设置时间P95 ms','控制周期P95 ms'],[[name,f"{r['setup_p95_ms']:.2f}",f"{r['controller_p95_ms']:.2f}"]
        for name,r in [('规划基线',base),('纯网络',av['raw']),('学习辅助规划',av['augmented']),('双指标门控',av['gate'])]], [7.2,4.6,4.6])
    p('SUMO在计算期间暂停，普通计算分位数不是最坏执行时间，不能证明真实延迟下的闭环稳定或安全。本文不统计通信延迟，'
      '未完成硬件在环、实车调试或完整车辆横纵向动力学验证。指令变化量降低也不能直接表述为乘坐舒适性已经改善。')
    p('当前证据的边界包括：仅三辆全受控固定路线车辆；状态测量精确；纵向误差界已知；邻车上一指令可用；'
      '统计校准依赖所属初态组内的可交换性；性能门控每次合流选择一次。更大规模交通、非合作驾驶员、未知意图、传感器误差与频繁在线切换均未由本实验覆盖。')
    p('本实验只有一个学习车辆，不能验证多个学习者之间的策略适应。扩展到多智能体时，需要多个车辆逐周期依据各自观测执行策略，'
      '并增加联合训练、集中训练分散执行与独立学习等匹配对照。本阶段不以学习辅助规划器的改善替代独立网络的前置证据。')
    p('本文的研究价值在于将问题变成明确可检验的基线相对决策：放行后是否兑现成本余量、是否拖慢自车、以及物理合同是否成立。'
      '共形最大残差、多策略rollout、Double DQN和备用安全过滤都有先例；本文贡献限定为具体问题的联合门控构造、条件分析与可审计实验，'
      '没有提出新的通用强化学习或共形预测算法。')
    p('与Conformal Policy Learning及ARMS的差别需在相同任务与预算下进一步实证[5,6]。本项目尚未复现并击败这些完整系统，'
      '也未证明所有复杂交通情形中学习都优于规划。后续重点应是更丰富交互下的独立策略收益、动态切换的风险累计处理，以及完整已发表方法对照。')
    p('所有负结果、未通过验收与有限适用条件随数据公开。原始模型、校准回合、测试协议、逐回合指标、轨迹与复核工具可在项目仓库获得[7]。')
    page('参考文献')
    references=[
        '[1] van Hasselt H, Guez A, Silver D. Deep Reinforcement Learning with Double Q-learning. AAAI, 2016. https://arxiv.org/abs/1509.06461',
        '[2] Wang Z, Schaul T, Hessel M, et al. Dueling Network Architectures for Deep Reinforcement Learning. 2016. https://arxiv.org/abs/1511.06581',
        '[3] Bertsekas D. Multiagent Rollout Algorithms and Reinforcement Learning. arXiv:1910.00120. https://arxiv.org/abs/1910.00120',
        '[4] Angelopoulos A N, Bates S. A Gentle Introduction to Conformal Prediction and Distribution-Free Uncertainty Quantification. arXiv:2107.07511. https://arxiv.org/abs/2107.07511',
        '[5] Huang H, Sharma S, Loquercio A, et al. Conformal Policy Learning for Sensorimotor Control Under Distribution Shifts. 2023. https://arxiv.org/abs/2311.01457',
        '[6] Liu N, Shen S, Li Z, et al. Adaptive Reinforcement and Model Predictive Control Switching for Safe Human-Robot Cooperative Navigation. Preprint, 2026. https://arxiv.org/abs/2601.16686',
        '[7] CoopGate. V13 ego-specific frozen protocol and SUMO results. Freeze commit 3bac214. https://github.com/Leo-7GOAT/CoopGate/tree/main/artifacts/v13/ego']
    for text in references:
        a=p(text);a.paragraph_format.first_line_indent=Cm(0);a.paragraph_format.space_after=Pt(4)
        for run in a.runs:run.font.size=Pt(10)
    doc.save(out)
    assert hashlib.sha256(reference.read_bytes()).hexdigest()==digest
    with ZipFile(reference) as old,ZipFile(out) as new:
        assert set(old.namelist())==set(new.namelist())
        changed=[name for name in old.namelist() if old.read(name)!=new.read(name)]
        assert set(changed)<={'word/document.xml','docProps/core.xml'},changed
    print(json.dumps({'output':str(out),'reference_sha256':digest,'changed_parts':changed}))


if __name__=='__main__':
    a=argparse.ArgumentParser();a.add_argument('--reference',type=Path,required=True);a.add_argument('--out',type=Path,required=True)
    v=a.parse_args();build(v.reference,v.out)
