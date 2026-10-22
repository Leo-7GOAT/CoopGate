"""Preserve draft six and append the completed independent V7 evidence."""
import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
from zipfile import ZipFile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table

ROOT = Path(__file__).resolve().parents[1]


def build(reference, out):
    original = hashlib.sha256(reference.read_bytes()).hexdigest()
    doc = Document(reference); template = deepcopy(doc.tables[0]._tbl)
    original_body = [e.xml for e in doc.element.body if e.tag != qn('w:sectPr')]
    main = json.loads((ROOT/'artifacts/v7/summary.json').read_text())
    pair = json.loads((ROOT/'artifacts/v7_pair/summary.json').read_text())
    cal = json.loads((ROOT/'artifacts/v7/calibration.json').read_text())
    doc.core_properties.version = '7'

    def p(text, style=None):
        paragraph = doc.add_paragraph(text, style)
        if style is None:
            paragraph.paragraph_format.first_line_indent = Cm(.74)
        return paragraph

    def eq(text):
        paragraph = doc.add_paragraph(); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        math = OxmlElement('m:oMath')
        def run(parent, text):
            r = OxmlElement('m:r'); value = OxmlElement('m:t'); value.text = text
            value.set(qn('xml:space'), 'preserve'); r.append(value); parent.append(r)
        pos = 0
        for token in re.finditer(r'([A-Za-z]+)(?:_\{([^}]+)\}|\^\{([^}]+)\})', text):
            run(math, text[pos:token.start()]); sub = token.group(2) is not None
            script = OxmlElement('m:sSub' if sub else 'm:sSup'); base = OxmlElement('m:e')
            run(base, token.group(1)); script.append(base)
            index = OxmlElement('m:sub' if sub else 'm:sup'); run(index, token.group(2) or token.group(3))
            script.append(index); math.append(script); pos = token.end()
        run(math, text[pos:]); paragraph._p.append(math)

    def table(headers, rows):
        element = deepcopy(template); doc.element.body.insert(len(doc.element.body)-1, element)
        table = Table(element, doc._body)
        while len(table.rows) > 1:
            table._tbl.remove(table.rows[-1]._tr)
        for cell, value in zip(table.rows[0].cells, headers):
            cell.text = value
        for values in rows:
            for cell, value in zip(table.add_row().cells, values):
                cell.text = str(value)
        for i, row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.space_before = paragraph.paragraph_format.space_after = Pt(5)
                    paragraph.paragraph_format.line_spacing = 1.05
                    paragraph.paragraph_format.keep_with_next = i < len(table.rows)-1
                    for run in paragraph.runs:
                        run.font.size = Pt(10); run.bold = i == 0
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    doc.add_page_break(); doc.add_heading('附录A 接受子集中的性能风险', 1)
    p('本稿保留前述V5与V6证据，新增两项独立验证。首先检查原性能评分能否支持更严格的接受条件风险；'
      '随后针对校准不足，另采新场景检验双方案模型核验。两个试验均在测试前冻结协议，不复用旧测试标签调阈值。')
    p('V6给出的是边际覆盖，而非接受学习方案后的条件可靠度。V7固定原三个模型，新增常规与高密度各512个校准初态。'
      '令a为所提方案，m为预测相对基线代价差，d为真实差；非基线提案在m≤t时被接受。以ε=10⁻⁶ s排除浮点相等误差。')
    eq('A_{t} = {a ≠ 0, m ≤ t},     R_{g}(t) = P_{g}(d > ε | A_{t})')
    p('在13个预先固定阈值−3、−2.75、…、0 s上，分别计算两分布的接受数n与退化数k。'
      '二项Clopper–Pearson单侧上界采用η=0.05/(2×13)，n=0或k=n时上界取1。')
    eq('B_{g}(t) = Beta^{−1}(1 − η; k + 1, n − k)')
    p('仅保留两个分布均满足B≤0.1的阈值，再选校准接受数最多者，同数选更严格者；无合格候选则完全回退。'
      '条件于接受样本数，退化数服从二项分布。对全部26个固定组合取并集上界，以至少95%的校准概率，所有获证候选的真实条件风险同时不超过10%，因此随后选择仍有效。')
    p('该推论假设各分布内独立同分布、模型及评分在校准前固定、校准与部署独立、模型代价与真实执行一致。'
      '若未来分布仅改变两类交通的混合比例，条件风险是各类风险按“混合权重×接受概率”加权的平均，故同样受界约束。'
      '它不涵盖任意新交通分布、未知动力学或接受概率为零时的条件风险。多重检验风险控制遵循已有Learn then Test思想[7]。')
    p('本次没有阈值同时获证。在高密度校准中，t=−1.0 s接受53个样本、退化7个，校正上界约31.5%；'
      't=−1.25 s仅接受13个样本、退化0个，上界仍约38.2%。严格门控因此完全回退。零接受不等于条件风险估计为零，'
      '这说明当前评分与数据不足以支撑所要求的证书，而不是通过降低报告标准获得成功。')

    doc.add_page_break(); doc.add_heading('附录B 新校准门控的独立测试', 1)
    p('主试验另取常规与高密度各80个独立新场景，六种方法共960次真实SUMO执行。对照包括冻结学习器、原V6门控、'
      '使用同一批新数据的两分布最大conformal分位数门控、同时PAC门控及传统基线和完整规划教师。'
      '高密度现已提供校准数据，不能把本轮称作未见分布泛化。')
    names = [('constant', '传统基线'), ('learned', '冻结学习器'), ('v6_conformal', '原V6门控'),
             ('robust_conformal', '两分布最大分位数'), ('robust_pac', '同时PAC门控'), ('teacher', '完整规划教师')]
    p('表A1  新场景中的平均截断到达时间  单位为秒', 'Caption')
    table(['方法', '常规 80场景', '高密度 80场景'],
          [[name, f"{main['summary']['nominal'][m]['mean_arrival_s']:.3f}", f"{main['summary']['dense'][m]['mean_arrival_s']:.3f}"] for m, name in names])
    p(f"两分布单侧残差分位数分别为{cal['conformal_quantiles_s']['nominal']:.3f} s和{cal['conformal_quantiles_s']['dense']:.3f} s，"
      '最大分位数门控使用两者较大值。它约束每个声明分布的边际覆盖，不能替代同时PAC门控的接受条件风险界。')
    for regime, label in [('nominal', '常规'), ('dense', '高密度')]:
        r = main['summary'][regime]['robust_conformal']; old = main['summary'][regime]['v6_conformal']
        interval = r['conditional_wilson95']; ci = '不定义' if interval is None else f'[{interval[0]:.1%}, {interval[1]:.1%}]'
        p(f"{label}分布中，最大分位数门控采用{r['accepted']}/80次，采用后退化{r['accepted_failures']}次，"
          f"条件风险描述性95% Wilson区间为{ci}；平均收益{r['gain_percent']:.2f}%。"
          f"原V6门控采用{old['accepted']}/80次、采用后退化{old['accepted_failures']}次，平均收益{old['gain_percent']:.2f}%。"
          f"最大分位数的提案上界实测覆盖率为{r['proposal_coverage']:.1%}，并未在两类分布都达到名义90%。")
    p('同时PAC门控因无合格阈值，在全部160个场景回退，收益为零。该负结果保留在原始数据中。'
      '测试区间只作描述性诊断，不替代校准阶段的同时置信保证；统计独立单元是场景而不是每个方法回合。'
      '与原V6门控的比较同时改变了校准分布覆盖与样本量，不能把差异全部归因于门控公式。')

    doc.add_page_break(); doc.add_heading('附录C 双方案核验的工程补充', 1)
    p('针对校准阶段发现的不足，补充试验单独冻结协议，采用80个额外独立场景、四种方法，共320次真实SUMO执行。'
      '在线仅展开传统基线与学习提案的完整反馈轨迹，选择到达代价较低者，同分选择传统基线；'
      '随后通过已有计划失效检测与每周期联合制动证书执行。这是在线模型核验，不能称作免费的概率置信度。')
    eq('J_{执行} = min(J_{基线}, J_{提案}) ≤ J_{基线}')
    p('该不退化关系要求模型与执行一致、计划未失效。状态失配后回退只保留适用条件下的低层合同，'
      '不再保证原始完整回合的性能关系。双方案核验也不能找到学习器未提议的更优方案。')
    p('表A2  独立补充试验的平均到达时间  单位为秒', 'Caption')
    table(['方法', '常规 40场景', '高密度 40场景'],
          [[name, f"{pair['summary']['nominal'][m]['mean_arrival_s']:.3f}", f"{pair['summary']['dense'][m]['mean_arrival_s']:.3f}"]
           for m, name in [('constant', '传统基线'), ('learned', '冻结学习器'), ('verified_pair', '双方案模型核验'), ('teacher', '完整规划教师')]])
    for regime, label in [('nominal', '常规'), ('dense', '高密度')]:
        r = pair['summary'][regime]['verified_pair']; ci = r['saved_ci95_s']
        p(f"{label}双方案核验相对基线改善{r['gain_percent']:.2f}%，保留完整教师{r['teacher_gain_retained_percent']:.1f}%的均值收益；"
          f"平均节省时间的描述性95%配对区间为[{ci[0]:.3f}, {ci[1]:.3f}] s，退化回合{r['wins_ties_losses'][2]}/40。")
    t, full = pair['timing']['verified_pair'], pair['timing']['teacher']
    p(f"双方案核验首次响应P95为{t['first_response_p95_ms']:.2f} ms，完整37方案规划为{full['first_response_p95_ms']:.2f} ms。"
      '时间包含在线模型展开，仿真计算期间未推进车辆；这不是包含计算延迟的闭环保证或最坏执行时间证明。')
    complete = sum(r['completed'] for study in (main, pair) for group in study['summary'].values() for r in group.values())
    collision = sum(r['collisions'] for study in (main, pair) for group in study['summary'].values() for r in group.values())
    p(f"两项试验共1280次执行，完成{complete}次，记录碰撞{collision}次。全部方案代价逐回合核对模型标签。"
      '结果仍限于精确模型、三车固定路线和全受控交通；尚无非合作车辆、感知误差或实车验证。')
    p('[7] Angelopoulos A N, Bates S, Candès E J, Jordan M I, Lei L. Learn then Test: Calibrating Predictive Algorithms to Achieve Risk Control. arXiv:2110.01052. https://arxiv.org/abs/2110.01052')
    assert [e.xml for e in list(doc.element.body)[:len(original_body)]] == original_body
    doc.save(out)
    assert hashlib.sha256(reference.read_bytes()).hexdigest() == original
    with ZipFile(reference) as before, ZipFile(out) as after:
        changed = [n for n in before.namelist() if before.read(n) != after.read(n)]
        assert set(changed) <= {'word/document.xml', 'docProps/core.xml'}, changed
    print(json.dumps({'output': str(out), 'changed_parts': changed}))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--reference', type=Path, required=True); parser.add_argument('--out', type=Path, required=True)
    args = parser.parse_args(); build(args.reference, args.out)
