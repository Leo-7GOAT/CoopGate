"""Audited joint-gating manuscript, preserving the retained V12 DOCX design."""
import argparse
import hashlib
import json
import re
from copy import deepcopy
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table

ROOT = Path(__file__).resolve().parents[1]


def build(reference, out):
    digest = hashlib.sha256(reference.read_bytes()).hexdigest()
    assert digest == 'e37213cca4aa5b563b6b7fe61ec83b37ad866f6b455f56916b9dd27ab234c704'
    study = ROOT / 'artifacts/v18/independent_study'
    report = json.loads((ROOT / 'artifacts/v18/independent_analysis/analysis.json').read_text())
    audit = json.loads((study / 'audit.json').read_text())
    assert report['source_sha256']['episodes.json'] == audit['episodes_sha256']
    assert hashlib.sha256((study / 'episodes.json').read_bytes()).hexdigest() == audit['episodes_sha256']
    result = report['results']; raw = result['raw_average']['groups']; gate = result['gate_average']['groups']
    single = json.loads((ROOT / 'artifacts/v14/single/summary.json').read_text())
    doc = Document(reference); template = deepcopy(doc.tables[0]._tbl)
    title = '多车协同学习决策的基线相对联合门控\n名义性能保证与扰动下闭环验证'
    doc.paragraphs[0].text = title
    doc.core_properties.title = title.replace('\n', ' '); doc.core_properties.version = '13'
    cut = doc.paragraphs[3]._p; remove = False
    for element in list(doc.element.body):
        if element is cut: remove = True
        if remove and element.tag != qn('w:sectPr'): doc.element.body.remove(element)

    def p(text, style=None):
        a = doc.add_paragraph(text, style=style)
        if style is None: a.paragraph_format.first_line_indent = Cm(.74)
        return a

    def h(text): doc.add_heading(text, 1)
    def page(text): doc.add_page_break(); h(text)

    def eq(text):
        a = doc.add_paragraph(); a.alignment = WD_ALIGN_PARAGRAPH.CENTER
        math = OxmlElement('m:oMath')
        def run(parent, value):
            r = OxmlElement('m:r'); t = OxmlElement('m:t'); t.text = value
            t.set(qn('xml:space'), 'preserve'); r.append(t); parent.append(r)
        pos = 0
        pattern = r'([A-Za-zΑ-Ωα-ωΣΠ≤≥0-9]+)(?:_\{([^}]+)\}(?:\^\{([^}]+)\})?|\^\{([^}]+)\})'
        for match in re.finditer(pattern, text):
            run(math, text[pos:match.start()]); sub = match.group(2); sup = match.group(3) or match.group(4)
            node = OxmlElement('m:sSubSup' if sub and sup else 'm:sSub' if sub else 'm:sSup')
            base = OxmlElement('m:e'); run(base, match.group(1)); node.append(base)
            if sub:
                element = OxmlElement('m:sub'); run(element, sub); node.append(element)
            if sup:
                element = OxmlElement('m:sup'); run(element, sup); node.append(element)
            math.append(node); pos = match.end()
        run(math, text[pos:]); a._p.append(math)

    def table(headers, rows, widths):
        element = deepcopy(template); doc.element.body.insert(len(doc.element.body) - 1, element)
        t = Table(element, doc._body)
        while len(t.rows) > 1: t._tbl.remove(t.rows[-1]._tr)
        while len(t.columns) < len(headers): t.add_column(Cm(2.5))
        # Prototype has three columns; every table in this draft has at least three.
        assert len(t.columns) == len(headers)
        t.autofit = False
        for cell, value in zip(t.rows[0].cells, headers): cell.text = value
        fill = deepcopy(t.rows[0].cells[0]._tc.get_or_add_tcPr().find(qn('w:shd')))
        if fill is not None:
            for cell in t.rows[0].cells:
                props = cell._tc.get_or_add_tcPr(); old = props.find(qn('w:shd'))
                if old is not None: props.remove(old)
                props.append(deepcopy(fill))
        for values in rows:
            for cell, value in zip(t.add_row().cells, values): cell.text = str(value)
        for col, width in zip(t.columns, widths): col.width = Cm(width)
        for i, row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
                for para in cell.paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    para.paragraph_format.space_before = para.paragraph_format.space_after = Pt(5)
                    para.paragraph_format.line_spacing = 1.05
                    para.paragraph_format.keep_with_next = i < len(t.rows) - 1
                    for run in para.runs: run.font.size = Pt(10); run.bold = i == 0
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def ci(values): return f'[{values[0]:.4f}, {values[1]:.4f}]'

    p('本文研究协同车辆在学习型决策与传统规划之间的逐周期选择。多车各自判断自己的学习动作有利，'
      '并不意味着这些动作同时执行后仍有利；团队平均收益也可能掩盖某辆车的延误。'
      '为此，本文对完整联合学习动作计算相对同一完整基线后续策略的逐车成本与到达时间差，'
      '仅在所有分量不退化且团队成本严格改善时采用学习动作，否则执行基线当前动作。'
      '门控以基线剩余价值为势函数，在精确名义模型下得到逐车性能不退化关系；物理安全则由独立的有界扰动备用制动证书检查。')
    p(f"在单车前置验证通过后，本文训练三组具有三个独立actor的模型式协同策略，并对256个新初态、四类扰动完成7,168次冻结SUMO对照。"
      f"常开学习策略平均团队成本相对强规划基线改善{raw['all']['team_gain_percent']:.2f}%，联合门控改善{gate['all']['team_gain_percent']:.2f}%，"
      f"后者配对初态簇95%区间为{ci(gate['all']['paired_team_gain_ci95'])}秒。"
      '同时公开逐车指标、每个训练种子、退化事件和计算代价。结果只支持所声明仿真条件下的结论，扰动下的条件性能风险尚未取得统计证书。')
    p('关键词：协同驾驶；多智能体强化学习；控制器切换；基线相对性能；可恢复安全；SUMO')
    h('1 引言')
    p('研究重点是判断学习策略何时值得采用。学习模块首先应有独立的性能贡献；性能门控再决定是否保留这种贡献。'
      '对简单场景，明确退回基线比强制使用学习控制更合理；对交互强的场景，学习潜力需要在冻结对照中验证。'
      '网络动作概率或价值差不能直接解释成相对传统规划的可靠性。')
    p('本文将可检验贡献限定为：逐车独立采用所造成的闭环组合问题；对完整联合动作实施双指标基线相对检查；'
      '以及单车前置验证、真实三actor训练和实际仿真执行组成的复核链。'
      '名义策略改进恒等式属于已有rollout思想，本文不将其声称为新的通用强化学习定理。')

    page('2 相关工作与研究定位')
    p('多智能体rollout通过分解联合动作或依次更新个体策略降低规划复杂度[1,2]。'
      '基线遗憾最小化也已在多智能体规划中研究相对既定策略的非退化结果[3]。'
      '因此，“多车学习优于基线”及“用基线价值检验动作”均有先例。本文的实际问题是：'
      '在安全过滤和车辆反应相互耦合时，采用判据是否对应真正执行的完整联合动作。')
    p('MAPPO使用集中价值估计与分散actor，是协作强化学习的重要参照[4]。本文实际实现过MAPPO和IPPO，'
      '但先前开发数据中二者均未击败较强的团队规划基线。最终训练采用模型式近似策略迭代：'
      '针对一个actor枚举当前动作，其他actor在未来继续闭环反应，以完整团队回报构造更新目标。'
      '随后采用数据聚合减少离线状态分布与实际actor访问分布的偏差[5]。该方法不是MAPPO，也没有纳什均衡保证。')
    p('共形预测与Learn Then Test为统计不确定性和风险检验提供了工具[6,7]。'
      'MA-COPP进一步处理多智能体策略改变导致的联合轨迹分布变化[8]。'
      '阈值型安全策略改进、多个候选的校准检验及深度模型式安全策略改进也已有研究[9,10]。'
      '本文对每个冻结策略直接采集新配对仿真，不实现离策略密度比迁移；当前最终门控使用确定性名义价值差，不能借用其他论文的统计保证。')
    table(['已有研究方向', '本文采用的要素', '尚需比较的范围'], [
        ['多智能体rollout', '完整后续回报与顺序更新', '已发表方法的同预算复现'],
        ['基线相对策略改进', '逐车成本及时间判据', '非名义模型下的改进条件'],
        ['共形与风险检验', '明确事件与分母', '有效采用率下的条件风险证书'],
        ['运行时安全监督', '逐步物理证书与故障接管', '感知误差和硬件实时验证']
    ], [4.4, 5.1, 6.9])
    p('本文的创新性应由组合机制的实际影响、有效采用收益和复核证据评价，不能仅凭“强化学习加门控”成立。'
      '目前还不足以声称全面优于上述完整系统，投稿前需要补充匹配预算的公开方法对照与更丰富场景。')

    page('3 控制对象与性能指标')
    p('实验使用SUMO中的三车固定路线合流：两辆主路车和一辆匝道车，路线标识为[0,1,0]。'
      '所有车辆都受共同物理过滤器控制，邻车状态和上一执行指令可用。'
      '这是一项协同纵向控制实验，尚不包含摄像头感知、横向轨迹跟踪或不可控人类驾驶员。'
      '决策周期0.5秒，物理子步0.1秒，时域25秒，目标路线位置60米。')
    eq('v_{i,t+1}=clip(v_{i,t}+(u_{i,k}+w_{i,t})Δt,0,20)')
    eq('p_{i,t+1}=p_{i,t}+v_{i,t+1}Δt,   |w_{i,t}|≤0.5')
    eq('−5≤u_{i,k}≤2.5,   |u_{i,k}−u_{i,k−1}|≤2.5')
    p('位置以米计，速度以米每秒计，加速度和误差以米每二次方秒计。指令变化约束对应5米每三次方秒的有限变化率。'
      '状态包含所有车辆位置、速度、上一执行指令与时钟。有限时域时钟和指令记忆必须保留，否则相同位置速度可能具有不同的后续代价和可行指令。')
    eq('J_{i}=T_{i}+λD_{i},   D_{i}=Σ_{k=0}^{Kᵢ−1}|u_{i,k}−u_{i,k−1}|,   λ=0.05')
    eq('J_{team}=(J_{0}+J_{1}+J_{2})/3,   Y_{i}=(J_{i},T_{i})')
    p('各车在到达前累计自身指令变化；到达之后停止自己的成本和时间累计。λ的单位为s³/m，因此成本以秒计。'
      '未完成时到达时间截断于25秒，并单独记录未完成、碰撞、约束违规和行政中止，不能用低变化量掩盖停车失败。'
      '学习训练优化团队平均成本，性能门控另外要求每车成本和到达时间均不退化。')
    p('“高冲突”依据初始状态预先定义：匝道车与任一主路车进入冲突区的预计到达时间差不超过1.25秒。'
      '其余为低冲突。该分组是本任务的交互强度代理，不能等同于所有真实道路的复杂程度。')

    page('4 真实多智能体策略训练')
    p('三个车辆分别持有独立actor参数，均根据18维本地组织的观测逐周期输出加速度提案。'
      '观测包含本车状态、上一指令、路线和时钟，以及150米范围内邻车的相对状态、路线关系、上一指令和可见性。'
      '共享状态是本协同仿真的前提；安全过滤器仍是集中模块，因此整套执行系统并非完全无通信、无协调的分散系统。')
    p('单车阶段先以模型式近似策略迭代学习匝道车策略，并对三个训练种子完成独立冻结验证。'
      '多车阶段使用对应已通过验证的单车模型初始化匝道actor，另两车新建actor。'
      '三组复现使用不同随机数、访问数据和主路参数初始化，但其单车模型共享早先声明的DQN预训练祖先，故不能声称从零开始的全流程独立训练。')
    eq('Q^{π}(z,a_{i},π_{−i}(z))=E[Σ_{k=t}^{H−1}c_{team}(z_{k},u_{k})]')
    p('每次固定完整参考联合策略，仅更新一个actor。针对该actor的全部可行当前动作，在精确名义模型中展开到目标或时域终点；'
      '后续三个actor均按参考策略闭环反应，共同安全过滤器也参与每一步。由完整团队回报得到改善目标，再用神经网络拟合。'
      '这包含多次环境转移、序列回报和相互响应，是模型式协作强化学习；不是把三个独立单车网络并排执行。')
    eq('J(μ)−J(π)=E_{μ}[Σ_{t}(Q^{π}(z_{t},μ(z_{t}))−V^{π}(z_{t}))]')
    p('上述性能差分关系说明，精确的逐状态改善可带来策略改善。但有限采样、局部观测、神经逼近和分布变化会破坏逐状态条件。'
      '因此本文不从训练损失下降直接推出闭环性能提升，而由独立SUMO结果检验。')
    p('第一阶段依次按[0,2,1]更新actor，共四个循环，每次8,192个状态、120个训练epoch。'
      '第二阶段共两个循环，每次actor更新包含三次数据聚合，参考改进动作混合率依次为1、0.5、0；每次聚合4,096个状态、80个epoch。'
      '固定使用最终检查点，不依据测试结果选择中间模型。最终原始网络策略在线只做前向推理和物理过滤，性能推演属于外部门控。')
    p('actor使用两个128维ReLU隐藏层，输出13个加速度等级，范围为−5至2.5m/s²、间隔0.625m/s²。'
      '输入在18维观测基础上加入八项确定性交通特征。训练使用温度0.04的回报遗憾软目标、交叉熵与期望遗憾损失，'
      'Adam学习率0.0005、批量上限256、梯度范数裁剪5；不可行动作预先屏蔽。保存的critic字段在此算法中冻结且不使用。')

    page('5 基线相对联合动作门控')
    p('强团队基线对当前联合动作进行两轮确定性坐标搜索，按[匝道、主路0、主路2]顺序更新候选，'
      '以共同物理过滤后的完整规则后续成本排序，下一周期重新计算。它是近似离散rollout规划器，不是全局最优求解器，也不冒称完整MPC复现。'
      '门控评价后续时，必须重复执行这一完整反馈基线，不能用更弱的规则序列代替。')
    eq('V_{t,i}^{b}(z)=c_{i}(z,b_{t}(z))+V_{t+1,i}^{b}(f(z,b_{t}(z)))')
    eq('δ_{t,i}(z,a)=c_{i}(z,a)+V_{t+1,i}^{b}(f(z,a))−V_{t,i}^{b}(z)')
    p('这里每个成本分量同时包含综合代价和到达时间；当前学习候选a是三个actor提案经同一物理过滤器后的联合指令。'
      '门控要求所有车辆、两个指标的δ均不大于零，并要求团队综合代价的平均δ严格小于零。'
      '条件不成立或候选与基线相同，则选择基线当前指令。数值实现使用1e−10容差和1e−9严格改善阈值。')
    eq('a_{t}=a_{t}^{learn}  if  δ_{t,i}≤0 ∀i and mean_{i}δ_{t,i,J}<0;  else  b_{t}(z_{t})')
    p('命题1：若名义动力学和基线后续价值计算精确，状态含必要指令记忆与时钟，终端价值为零，'
      '并逐周期使用上述门控，则同一初态下各车完整名义成本与到达时间均不高于基线。')
    eq('Y_{i}^{gate}(z_{0})−Y_{i}^{b}(z_{0})=Σ_{t=0}^{H−1}δ_{t,i}(z_{t},a_{t})≤0')
    p('证明：沿实际名义状态转移累加δ，所有中间基线剩余价值相消，得到上式。学习动作获准时逐分量非正，'
      '基线自身动作由Bellman恒等关系使δ等于零。各车到达后令对应阶段成本为零，故同一求和覆盖各自不同的完成时刻。'
      '这是已有策略改进思想在当前双指标联合门控中的应用。')
    p('该判据给出可解释的名义性能裕量，没有输出已校准的成功概率。'
      '它允许行程中多次选择学习或基线，保证的是所声明名义模型下的全程相对性能，不能直接推及有扰动的实际行程。')

    page('6 物理证书与安全接管')
    p('性能门控回答“采用学习是否有益”，物理过滤器回答“当前动作是否保留安全备用方案”。'
      '两者约束不同。所有对照共用相同加速度、指令变化率与路线安全合同，学习获准也不能绕过物理检查。')
    p('同路车合同要求至少6.8米间距，不同入口车在冲突区[−25,6.8]的占用不能重叠。'
      '检查覆盖每个物理子步内的路线运动区间。该合同只对应声明的纵向路线模型，不代表任意道路几何、车辆碰撞包络或感知误差模型。')
    eq('u_{i,k+1}^{backup}=max(u_{i,k}^{backup}−2.5,−5)')
    p('每车提供提案、保持上一指令与备用制动三个候选，最多形成27个联合顶点。'
      '按偏离原提案的平方距离排序，选择通过鲁棒备用证书的联合指令。'
      '证书以区间动力学涵盖每个子步±0.5m/s²误差，检查当前动作及后续制动，直到速度上界全部归零并能抵消最大正向误差。')
    p('命题2：在初态可恢复、精确测量、已知动力学与误差界、制动和指令变化率可实现等条件下，'
      '获证书的当前动作保留后继状态的安全备用序列。证明采用剩余备用序列移位，并以保持静止的终端动作续接。'
      '实际后继状态包含于预测区间，所以安全合同能够递归保持。超出这些前提时不作绝对物理安全承诺。')
    p('工程接管模块校验策略输出的形状、有限数值、幅值、变化率及物理证书。'
      '发生异常或不合法输出时，基线从当前实测状态、真实上一指令和同步时钟接管，不重置环境。'
      '故障后保持基线到本回合结束；初始采用标签保留，防止事后改变风险统计分母。')
    p('已有独立故障注入验证覆盖8个初态、四类扰动、两类控制器和六种分支，共384次实际SUMO执行；'
      '192次异常、非有限输出或非法指令注入均完成接管。4,898条接管指令重建通过，64条拒绝轨迹与基线一致。'
      '该实验使用先前冻结控制器并强制采用，证明接管模块可执行，不等于最终三组门控模型获得统计风险证书。')
    p('正常性能拒绝与运行故障不同。前者只在当前周期选择基线，下周期可重新判断；后者采用粘性接管。'
      '两类事件分别记录，不把物理修正次数、性能拒绝次数和故障接管次数混为同一干预指标。')

    page('7 冻结实验设计与复核')
    p('全部训练先完成，再冻结模型哈希、门控条件、测试初态范围、扰动和验收标准。'
      '最终使用256个新初态，前128个为名义交通分布，后128个为密集交通分布。'
      '每个初态以相同扰动序列比较强基线、三组常开学习和对应三组门控，共7种方法、7,168次实际SUMO运行。'
      '各初态内部随机化方法执行顺序；四个计算分片只划分初态，不拆散方法或扰动配对。')
    table(['阶段', '每组训练或评价规模', '证据用途'], [
        ['单车冻结验证', '512初态 × 4模式 × 7方法', '14,336次实际SUMO'],
        ['多车训练A', '12次actor更新 × 8,192状态', '98,304访问状态'],
        ['多车训练B', '6次更新 × 3聚合 × 4,096状态', '73,728访问状态'],
        ['最终联合对照', '256初态 × 4模式 × 7方法', '7,168次实际SUMO'],
        ['最终独立重放', '2初态 × 4模式 × 7方法', f"{audit['physical_audit']['additional_sumo_replays']}次额外SUMO"],
        ['名义恒等式复核', '256初态 × 3个门控模型', f"{audit['nominal_reconstruction']['identities']}条模型闭环重建"]
    ], [4.0, 7.0, 5.4])
    p('四类误差为无误差、逐子步均匀误差、固定偏置与脉冲误差，幅值均在声明界内。'
      '模型训练和门控预测不读取实际未来扰动序列。训练中的反事实分支转移与名义闭环重建不计入实际SUMO运行次数。')
    p('区间估计的采样单位是初态。同一个初态的四种扰动和三个固定模型先求平均，再按名义和密集分层固定数量进行5,000次配对bootstrap。'
      '报告95%百分位区间，条件是三个已训练模型；它不估计任意训练随机种子的总体不确定性。'
      '同时列出每个模型结果，避免用平均值掩盖失败训练种子。')
    p('严格退化事件定义为同一初态的任一种扰动下，任一车辆的成本或到达时间高于配对基线。'
      '该事件会捕获小幅延误，不表示碰撞。获准定义为该初态至少一种模式下执行过学习动作。'
      '分别报告全部策略初态对与获准策略初态对中的经验比例；多个模型共享初态，这些计数不能当作独立二项试验。')

    page('8 单车前置验证与原始多车收益')
    p('单车阶段只学习匝道actor，其他车辆提出固定动作并由共同物理层过滤。三个固定训练种子均改善自车成本和时间，'
      '冻结的九项前置标准全部通过，之后才进入多车训练。')
    table(['单车方法', '到达时间 s', '变化量 m/s²', '综合成本 s'], [
        ['强自车规划基线', f"{single['methods']['rollout']['ego_time']:.4f}", f"{single['methods']['rollout']['ego_variation']:.4f}", f"{single['methods']['rollout']['ego_cost']:.4f}"],
        ['三模型学习平均', '9.4364', '5.5315', '9.7129'],
        ['相对改善', '1.93%', '23.91%', '2.73%']
    ], [5.3, 3.7, 3.7, 3.7])
    p('单车结果为多车初始化提供有效学习前提，但不说明其与任意学习邻车组合后仍有优势。'
      '多车比较因此使用更强的团队规划基线，以三车平均成本为主要目标，同时保留逐车指标。')
    table(['最终模型', '常开成本 s', '改善比例', '配对95%区间 s'], [
        [s, f"{result['raw_'+s]['groups']['all']['method_cost']:.4f}",
         f"{result['raw_'+s]['groups']['all']['team_gain_percent']:.2f}%", ci(result['raw_'+s]['groups']['all']['paired_team_gain_ci95'])]
        for s in ('1802', '1812', '1822')
    ] + [['三模型平均', f"{raw['all']['method_cost']:.4f}", f"{raw['all']['team_gain_percent']:.2f}%", ci(raw['all']['paired_team_gain_ci95'])]], [3.2, 3.3, 3.1, 6.8])
    p(f"强团队基线成本为{raw['all']['baseline_cost']:.4f}秒。常开学习在高冲突组改善{raw['high_conflict']['team_gain_percent']:.2f}%，"
      f"低冲突组改善{raw['low_conflict']['team_gain_percent']:.2f}%。这些结果检验了本任务中的交互收益，不能推广为所有复杂场景下学习必胜。")
    p('团队均值不能说明利益分配。常开学习使主路2的平均成本增加0.1713秒，成本改善区间为[−0.2811, −0.0665]秒。'
      '这一逐车退化为联合双指标门控提供了明确的检验对象；模型1812的门控团队均值也略差于其常开执行，不能声称每个模型都被门控进一步加速。')
    p('真实多智能体训练的证据还包括：三个actor独立参数均更新；每个角色的训练目标来自其他车辆闭环响应后的共同回报；'
      '最终检查点与所有中间权重、访问状态和反事实目标均公开。复核逐次确认未更新角色保持冻结，并重新计算采样状态的分支回报。')

    page('9 联合门控结果')
    table(['最终模型', '门控成本 s', '改善比例', '配对95%区间 s'], [
        [s, f"{result['gate_'+s]['groups']['all']['method_cost']:.4f}",
         f"{result['gate_'+s]['groups']['all']['team_gain_percent']:.2f}%", ci(result['gate_'+s]['groups']['all']['paired_team_gain_ci95'])]
        for s in ('1802', '1812', '1822')
    ] + [['三模型平均', f"{gate['all']['method_cost']:.4f}", f"{gate['all']['team_gain_percent']:.2f}%", ci(gate['all']['paired_team_gain_ci95'])]], [3.2, 3.3, 3.1, 6.8])
    gains = gate['all']['mean_individual_gain']
    table(['车辆', '平均成本改善 s', '平均到达时间改善 s'], [
        [label, f'{gains[i][0]:.4f}', f'{gains[i][1]:.4f}'] for i, label in enumerate(('主路0', '匝道1', '主路2'))
    ], [4.4, 6.0, 6.0])
    p(f"门控在高冲突组改善{gate['high_conflict']['team_gain_percent']:.2f}%，在低冲突组改善{gate['low_conflict']['team_gain_percent']:.2f}%。"
      f"共{gate['all']['admitted_policy_initial_pairs']}/{gate['all']['policy_initial_pairs']}个策略初态对曾采用学习动作，"
      f"采用步数为{result['gate_average']['learned_action_steps']}/{result['gate_average']['control_steps']}。"
      '这说明门控在多数控制步使用基线，选择少量有名义改善依据的联合动作，而非试图让学习持续接管。')
    p(f"严格退化事件在常开学习中为{raw['all']['harmful_policy_initial_pairs']}/{raw['all']['policy_initial_pairs']}，"
      f"在门控中为{gate['all']['harmful_policy_initial_pairs']}/{gate['all']['policy_initial_pairs']}；"
      f"门控获准后的经验退化比例为{100*gate['all']['empirical_harm_among_admitted']:.2f}%。"
      '这些事件按各模型分别判定后计数，没有先平均不同模型来抵消个体损失。它们仍是描述性结果，不构成选择后的风险上界。')
    p(f"全部{audit['nominal_reconstruction']['identities']}条名义门控闭环均重建并与实际无扰动SUMO结果匹配，"
      '逐车实际成本差与获准步骤的名义优势之和一致。扰动模式中的性能退化不违背命题1，因为其状态转移前提不同。')

    page('10 组合效应与扰动失配')
    p('为检验独立采用是否可组合，另固定开发模型1701，在64个新初态、四种误差下比较强基线、完整学习、'
      '三个单独学习组件、逐车独立门控与完整联合门控，共1,792次实际SUMO执行。'
      '逐车门控先预测“只有我学习、其他车使用基线”的完整结果，再实际执行同时选出的组合。')
    table(['开发门控', '获准初态', '严格退化初态', '平均成本 s'], [
        ['逐车独立', '42/64', '36/64', '9.8572'],
        ['完整联合', '18/64', '13/64', '9.8968'],
        ['始终强基线', '0/64', '0/64', '9.9812']
    ], [5.2, 3.2, 4.0, 4.0])
    p('320条名义预测均与对应无扰动SUMO结果一致。但同时选中至少两车学习的9个初态中，'
      '3个出现已选车辆的无扰动退化。问题不是预测器算错了原问题，而是同时改变其他车辆策略后，'
      '原来“其他车保持基线”的反事实条件失效。完整联合判断的无扰动退化为零，支持对实际联合动作作整体判断。')
    p('逐步联合门控的扩大开发对照使用另128个新初态，完成1,536次SUMO执行，团队平均改善1.88%，'
      '三车平均成本均改善；但44/128个初态仍有扰动下的严格退化。开发证据没有隐藏，最终冻结测试也没有据此更换评判分母。')
    eq('ξ_{t,i}=V_{t+1,i}^{b}(z_{t+1}^{actual})−V_{t+1,i}^{b}(f(z_{t},a_{t}))')
    p('在有扰动轨迹上，名义δ之外还出现基线价值的转移残差ξ。若再与基线的有扰动运行比较，还要扣除基线自身的对应残差累计。'
      '有限时域到达时间、动作离散搜索和安全投影可能产生不光滑变化，因此不能未经证明假设一个小的全局Lipschitz常数来消去误差。')
    p('此前对完整策略使用最大残差共形校准，以及17个预先固定的扰动预测情形。'
      '在一个64初态独立开发测试中，校准后的两个候选均未放行。该负结果说明严密风险条件可能消耗全部学习收益，'
      '而不是支持降低阈值或把名义裕量重新命名为置信概率。')

    page('11 工程代价与适用边界')
    table(['控制方法', '实测计算p95 ms', '在线性能推演'], [
        ['强团队基线', f"{result['baseline']['p95_computation_ms']:.2f}", '候选动作及规则后续'],
        ['常开学习平均', f"{result['raw_average']['p95_computation_ms']:.2f}", '无'],
        *[[f'联合门控 {s}', f"{result['gate_'+s]['p95_computation_ms']:.2f}", '完整反馈基线后续'] for s in ('1802', '1812', '1822')]
    ], [5.1, 5.0, 6.3])
    p('计算时间来自本机CPU执行，统计含实际门控计算，不包含通信延迟。'
      '完整基线后续推演增加了开销，p95低于500毫秒决策周期也不等于最坏情况实时保证。'
      '工程部署还需要冷启动、最坏状态、处理器负载和超时接管的独立测试。')
    checks = report['checks']; labels = {
        'every_raw_seed_positive': '每个原始模型平均改善',
        'raw_average_positive_ci': '原始平均改善区间为正',
        'raw_high_conflict_positive': '高冲突原始平均改善',
        'gate_one_percent_positive_ci': '门控改善至少百分之一且区间为正',
        'gate_all_individual_means_nonnegative': '门控逐车双指标均值不退化',
        'nominal_identities': '名义恒等式',
        'gate_regressions_reduced_both_denominators': '两个分母下退化频率均降低',
        'all_actual_runs_complete_and_no_recorded_events': '所有方法全部完成且无记录事件',
        'each_gate_p95_below_decision_period': '门控计算p95小于决策周期'}
    failed = [labels[k] for k, value in checks.items() if not value]
    p(f"预先冻结的九项平均性能与执行检查通过{sum(checks.values())}/9项。"
      + ('未通过项为：' + '、'.join(failed) + '。' if failed else '所有平均性能检查通过。')
      + '检查涵盖每个原始训练种子的平均收益、配对区间、门控净收益、逐车均值、名义恒等式、退化频率、实际完成及计算p95；'
      '通过这些检查不等于取得扰动下条件风险证书。')
    physical = result['baseline']['runs'] + sum(result[k]['runs'] for k in ('raw_1802', 'raw_1812', 'raw_1822', 'gate_1802', 'gate_1812', 'gate_1822'))
    events = sum(result[k]['collisions'] + result[k]['contract_violations'] + result[k]['aborted'] for k in ('baseline', 'raw_1802', 'raw_1812', 'raw_1822', 'gate_1802', 'gate_1812', 'gate_1822'))
    completed = sum(result[k]['completed'] for k in ('baseline', 'raw_1802', 'raw_1812', 'raw_1822', 'gate_1802', 'gate_1812', 'gate_1822'))
    p(f'最终实际运行完成数为{completed}/{physical}，记录碰撞、合同违规和行政中止总数为{events}。'
      '零记录事件描述样本，不证明开放道路中零事故概率。当前模型覆盖固定三车合流，无法据此声称可扩展到任意车数或复杂路口。')
    p('未完成的三次均来自常开模型1812的偏置扰动；三个门控模型3,072次执行全部完成。'
      '对全部三个失败初态追加基线、常开及门控的九次解释性重放，三次超时重现，基线和门控均完成。'
      '门控前两例各采用一次学习动作，第三例全程退回基线；这些事后解释案例不构成新的确认性效果样本。')
    p('学习与门控使用已知纵向模型，SUMO提供实际仿真闭环执行和可复核状态。'
      '这不是跨动力学仿真器的迁移验证；要增强工程可信度，应增加参数失配、感知误差、更多车辆、横向跟踪及独立动力学后端。'
      '由于三个策略共享早期单车祖先，训练随机性结论也应保持当前条件范围。')

    page('12 讨论与结论')
    p('本研究将论文重点落在采用决策上：学习模块先提供可测的序列决策收益，门控再以完整联合动作和相同基线后续策略检验收益。'
      '相比逐车单独判断，联合判断处理了其他车辆策略同时变化造成的反事实条件失效；相比只优化团队均值，逐车双指标检查限制了名义场景中的利益转移。')
    p('实验已包含单车通过后的真实多actor训练、冻结的新初态对照、逐步名义恒等式重建和额外SUMO重放。'
      '这些证据支持本任务中学习与门控的具体表现。论文没有把简单场景无收益、个别车辆退化或校准完全拒绝等结果删去，'
      '也不把物理可恢复性与相对基线的性能保证混为一个概率。')
    p('当前主要理论结论是条件明确的名义非退化关系和备用可恢复性；扰动下只给出经验比较。'
      '若要保留“置信门控”作为核心术语，需要在独立校准与独立测试之间严格分离选择，'
      '并证明足够采用率下的性能退化风险。对所有初态的边际有害采用界，不能直接解释成获准之后的条件风险。')
    p('进一步工作应优先解决有效采用与风险控制之间的矛盾，而不是增加没有独立收益的模型规模。'
      '可检验方向包括状态相关的转移残差界、面向逐车损失的保守价值近似，以及把门控计算预算显式纳入执行决策。'
      '任何新预测器或阈值都需要新的冻结测试集，不能继续用本轮最终数据调参后维持同一确认性结论。')
    p('投稿价值仍取决于与已发表安全策略改进、联合预测和运行时监督方法的匹配比较，以及更丰富场景的外部验证。'
      '当前成果适合作为完整的可复现研究初稿和工程作品集；“方法组合本身足以达到顶刊或顶会创新标准”尚无依据。'
      '全部代码、模型、冻结协议、负结果、逐回合记录和复核命令公开于CoopGate仓库[11]。')

    page('参考文献')
    refs = [
        '[1] Bertsekas D. Multiagent Rollout Algorithms and Reinforcement Learning. 2019. https://arxiv.org/abs/1910.00120',
        '[2] Multiagent Reinforcement Learning Rollout and Policy Iteration. IEEE/CAA Journal of Automatica Sinica, 2021. https://www.ieee-jas.com/en/article/doi/10.1109/JAS.2021.1003814',
        '[3] Multi-Agent Planning with Baseline Regret Minimization. IJCAI, 2017. https://www.ijcai.org/proceedings/2017/63',
        '[4] Yu C, et al. The Surprising Effectiveness of PPO in Cooperative Multi-Agent Games. NeurIPS, 2022. https://arxiv.org/abs/2103.01955',
        '[5] Ross S, Gordon G, Bagnell D. A Reduction of Imitation Learning and Structured Prediction to No-Regret Online Learning. AISTATS, 2011. https://proceedings.mlr.press/v15/ross11a.html',
        '[6] Angelopoulos A N, Bates S. A Gentle Introduction to Conformal Prediction and Distribution-Free Uncertainty Quantification. 2021. https://arxiv.org/abs/2107.07511',
        '[7] Learn then Test: Calibrating Predictive Algorithms to Achieve Risk Control. https://arxiv.org/abs/2110.01052',
        '[8] Kuipers T, et al. Conformal Off-Policy Prediction for Multi-Agent Systems. CDC, 2024. https://arxiv.org/abs/2403.16871',
        '[9] Cho B M, et al. CSPI-MT: Calibrated Safe Policy Improvement with Multiple Testing for Threshold Policies. 2024. https://arxiv.org/abs/2408.12004',
        '[10] Deep SPI: Safe Policy Improvement via World Models. ICLR, 2026. https://proceedings.iclr.cc/paper_files/paper/2026/hash/0212b0dff7b4aa3bd8b8cc4ce24fd384-Abstract-Conference.html',
        '[11] CoopGate. Frozen three-model SUMO protocol and raw records. Protocol freeze 79beb44. https://github.com/Leo-7GOAT/CoopGate/tree/main/artifacts/v18/independent_study'
    ]
    for text in refs:
        a = p(text); a.paragraph_format.first_line_indent = Cm(0); a.paragraph_format.space_after = Pt(6)
        for run in a.runs: run.font.size = Pt(10)
    doc.save(out)
    assert hashlib.sha256(reference.read_bytes()).hexdigest() == digest
    with ZipFile(reference) as old, ZipFile(out) as new:
        assert set(old.namelist()) == set(new.namelist())
        changed = [name for name in old.namelist() if old.read(name) != new.read(name)]
        assert set(changed) <= {'word/document.xml', 'docProps/core.xml'}, changed
    print(json.dumps(dict(output=str(out), reference_sha256=digest, changed_parts=changed), ensure_ascii=False))


if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--reference', type=Path, required=True)
    parser.add_argument('--out', type=Path, required=True)
    a = parser.parse_args(); build(a.reference, a.out)
