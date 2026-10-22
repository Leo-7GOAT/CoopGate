"""Update the retained manuscript with completed V19 evidence, not live outcomes."""
import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
from zipfile import ZipFile
from lxml.etree import tostring

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table

ROOT = Path(__file__).resolve().parents[1]
REFERENCE_SHA = '295adef770b35213760f37ff880da6faa4f1897e85b30b1bf72abeb8c4a23cf1'


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def build(reference, out, evidence_out):
    assert reference.resolve() != out.resolve() and digest(reference) == REFERENCE_SHA
    evidence = {}
    for name in ('portfolio_probe_analysis', 'composition_ablation', 'shifted_study', 'shifted_analysis',
                 'shifted_replay', 'parallel_gate_audit', 'parallel_analysis', 'certification_power',
                 'actuator_bounds', 'deployment_runtime_audit'):
        directory = ROOT / 'artifacts/v19' / name
        for relative, expected in json.loads((directory / 'sha256.json').read_text()).items():
            assert digest(directory / relative) == expected, (name, relative)
        evidence[f'artifacts/v19/{name}/sha256.json'] = digest(directory / 'sha256.json')
    def load(name, file='analysis.json'):
        return json.loads((ROOT / 'artifacts/v19' / name / file).read_text())
    composition = load('composition_ablation')
    transfer = load('shifted_analysis')
    replay = load('shifted_replay', 'audit.json')
    runtime = load('parallel_analysis')
    assert transfer['actual_sumo_runs'] == 6400 and replay['physical_failures_reproduced'] == 68
    assert runtime['actual_sumo_replays'] == 58 and runtime['paired_decisions'] == 667
    assert load('certification_power')['minimum_zero_harm_adoptions'] == 115
    assert composition['adoption_transitions']['portfolio_only'] == 7

    doc = Document(reference)
    prototype = deepcopy(doc.tables[0]._tbl)
    original_section = doc.sections[0]._sectPr.xml
    original_math = [deepcopy(e) for e in doc.element.body.iter(qn('m:oMath'))]
    title = '多车协同学习决策的基线相对联合门控\n有限扰动保证与动力学迁移验证'
    doc.paragraphs[0].text = title
    doc.core_properties.title = title.replace('\n', ' ')
    doc.core_properties.version = '14'

    def replace(prefix, content):
        matches = [p for p in doc.paragraphs if p.text.startswith(prefix)]
        assert len(matches) == 1, prefix
        matches[0].text = content
        return matches[0]

    replace('本文研究协同车辆在学习型决策与传统规划之间的逐周期选择。',
        '本文研究协同车辆在学习决策与传统规划之间的逐周期选择。多车分别认可的学习动作同时执行后可能相互干扰，团队均值也可能掩盖个别车辆延误。'
        '本文对学习与基线动作的完整联合组合检查逐车成本和到达时间，并将检查扩展至固定扰动库。'
        '通过基线剩余价值的望远镜求和，给出真实扰动属于该库且模型精确时的性能不退化关系；物理可恢复性由独立安全过滤器约束。')
    replace('在单车前置验证通过后，本文训练三组具有三个独立actor的模型式协同策略，',
        '在256个新初态、7,168次SUMO运行的冻结对照中，原始三车学习与名义门控的平均团队成本分别改善1.63%和2.01%。'
        '严格扰动组合门控在开发集保留0.82%收益，并在6,400次直接迁移中显示密集交通和紧密跟车收益，同时保留低速让行退化和强惯性失败。'
        '完整回放中的等价并行计算降低15.22%耗时，但未满足最坏情况实时要求。采用后的条件风险校准及独立最终比较尚未完成。')
    replace('本文将可检验贡献限定为：',
        '本文考察完整联合组合对实际采用、逐车性能和计算代价的影响，并给出有限扰动条件下的保证与动力学迁移边界。'
        '风险校准采用已有的Learn then Test工具，策略改进恒等式也属于已有rollout思想。'
        '独立校准完成后，仍须同时检验条件风险、学习采用率和实际净收益，才能评价门控的部署价值。')
    replace('共形预测与Learn Then Test为统计不确定性和风险检验提供了工具',
        '共形预测与Learn then Test为不确定性和风险检验提供了工具[6,7]；MA-COPP处理多智能体策略改变后的联合轨迹分布[8]。'
        '阈值型安全策略改进与深度模型式改进已有相关研究[9,10]。本文使用新采集的配对仿真，能够观察候选与基线的完整结果。'
        '第12节给出同候选、同观测、同样本预算的共形、条件风险和均值改进适配；这些适配的保证对象不同，最终对照结果尚未产生。')
    replace('7 冻结实验设计与复核', '7 名义门控的独立实验设计')
    replace('全部训练先完成，再冻结模型哈希、门控条件、测试初态范围、扰动和验收标准。',
        '针对第5节名义门控，训练完成后冻结模型哈希、门控条件、初态范围、扰动和验收标准。'
        '独立比较使用256个新初态，名义与密集交通各128个。每个初态配对四种扰动，比较强基线、三个原始学习模型和对应三个门控，共7,168次SUMO执行。'
        '方法顺序在初态内随机化；计算分片只划分初态，保留全部方法和扰动配对。')
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text in ('最终联合对照', '最终独立重放'):
                    content = {'最终联合对照': '名义联合对照', '最终独立重放': '独立物理重放'}[cell.text]
                    assert len(cell.paragraphs) == 1
                    cell.paragraphs[0].text = content
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs: run.font.size = Pt(10)
    replace('9 联合门控结果', '9 名义联合门控结果')
    end_nominal = next(p for p in doc.paragraphs if p.text.startswith('全部768条名义门控闭环'))
    extra = doc.add_paragraph('该独立研究完成7,165/7,168次运行，预注册检查通过8/9项。三个未完成案例均来自原始模型1812的偏置扰动，门控完成3,072/3,072次。'
        '三个失败在追加回放中重现，相应基线和门控均完成。零记录碰撞描述有限样本，不构成开放道路的零事故保证。')
    extra.paragraph_format.first_line_indent = Cm(.74)
    end_nominal._p.addnext(extra._p)

    # Retain chapters 1–10 and the original reference elements; update the
    # superseded engineering/discussion block without rebuilding the package.
    reference_heading = next(p for p in doc.paragraphs if p.text == '参考文献')
    children = list(doc.element.body)
    ref_index = children.index(reference_heading._p)
    retained_references = [deepcopy(e) for e in children[ref_index:] if e.tag != qn('w:sectPr')]
    chapter11 = next(p for p in doc.paragraphs if p.text == '11 工程代价与适用边界')
    cut = chapter11._p.getprevious()  # existing page break
    removing = False
    for element in list(doc.element.body):
        if element is cut: removing = True
        if removing and element.tag != qn('w:sectPr'): doc.element.body.remove(element)

    def p(text):
        paragraph = doc.add_paragraph(text)
        paragraph.paragraph_format.first_line_indent = Cm(.74)
        return paragraph
    def page(text):
        doc.add_page_break()
        doc.add_heading(text, 1)
    def eq(text):
        paragraph = doc.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        math = OxmlElement('m:oMath')
        def run(parent, value):
            r = OxmlElement('m:r'); t = OxmlElement('m:t'); t.text = value
            t.set(qn('xml:space'), 'preserve'); r.append(t); parent.append(r)
        position = 0
        pattern = r'([A-Za-zΑ-Ωα-ωΣΠ≤≥0-9]+)(?:_\{([^}]+)\}(?:\^\{([^}]+)\})?|\^\{([^}]+)\})'
        for match in re.finditer(pattern, text):
            run(math, text[position:match.start()]); sub = match.group(2); sup = match.group(3) or match.group(4)
            node = OxmlElement('m:sSubSup' if sub and sup else 'm:sSub' if sub else 'm:sSup')
            base = OxmlElement('m:e'); run(base, match.group(1)); node.append(base)
            if sub:
                e = OxmlElement('m:sub'); run(e, sub); node.append(e)
            if sup:
                e = OxmlElement('m:sup'); run(e, sup); node.append(e)
            math.append(node); position = match.end()
        run(math, text[position:]); paragraph._p.append(math)
    def table(headers, rows, widths):
        element = deepcopy(prototype); doc.element.body.insert(len(doc.element.body)-1, element)
        t = Table(element, doc._body)
        while len(t.rows) > 1: t._tbl.remove(t.rows[-1]._tr)
        while len(t.columns) < len(headers): t.add_column(Cm(2.5))
        assert len(t.columns) == len(headers) and abs(sum(widths)-16.4) < 1e-9
        t.autofit = False
        for cell, value in zip(t.rows[0].cells, headers): cell.text = value
        fill = deepcopy(t.rows[0].cells[0]._tc.get_or_add_tcPr().find(qn('w:shd')))
        for cell in t.rows[0].cells:
            props = cell._tc.get_or_add_tcPr()
            old = props.find(qn('w:shd'))
            if old is not None: props.remove(old)
            props.append(deepcopy(fill))
        for values in rows:
            for cell, value in zip(t.add_row().cells, values): cell.text = str(value)
        for col, width in zip(t.columns, widths): col.width = Cm(width)
        for index, row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for para in cell.paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    para.paragraph_format.space_before = para.paragraph_format.space_after = Pt(5)
                    para.paragraph_format.line_spacing = 1.05
                    para.paragraph_format.keep_with_next = index < len(t.rows)-1
                    for run in para.runs: run.font.size = Pt(10); run.bold = index == 0
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    page('11 有限扰动下的联合组合门控')
    p('对三个actor的学习提案与同一反馈基线动作作二元组合，至多产生七个非基线候选。'
      '每个完整组合先经过共同物理过滤器，再检查全部车辆的成本与时间。扰动库包含一条无扰动、八条恒定异号偏置和八条固定符号脉冲。'
      '候选先按名义团队收益排序，只有全部17条预测都通过的候选才能采用。')
    eq('δ_{t,ω,i,m}(z,a)=c_{i,m}(z,a)+V_{t+1,ω,i,m}^{b}(f_{t,ω}(z,a))−V_{t,ω,i,m}^{b}(z)')
    p('这里m表示综合成本或到达时间，V表示从当前时刻到固定终点持续执行同一反馈基线的剩余指标。'
      '严格准入要求所有扰动、车辆和指标的δ非正，且名义团队成本严格下降。未获准时执行基线，其对应差值为零。')
    eq('J_{ω,i,m}^{g}−J_{ω,i,m}^{b}=Σ_{t}δ_{t,ω,i,m}(z_{t},a_{t})≤0')
    p('命题3：若真实外生扰动恰为库中某条序列，模型与基线续跑精确，状态包含时钟与上一命令，则任意逐步切换的已通过组合都满足上述逐车不退化关系。'
      '证明沿真实轨迹逐步展开同一时变基线价值，中间项抵消；终端价值为零。候选提前拒绝与排序不影响该条件。'
      '均匀噪声和每次脉冲重新抽符号的实际过程不必属于此有限库，因而不能继承该路径保证。'
      '上述命题采用精确的非正差值；若精确评估下每步允许ε正容差，累计上界应为Hε，浮点评估误差还需单独计入。')
    table(['开发比较', '团队成本改善', '采用后退化记录'], [
        ['名义联合门控', '1.964%', '39 / 85'],
        ['严格整动作扰动库', '0.545%', '1 / 39'],
        ['严格联合组合扰动库', '0.822%', '1 / 46']], [6.4, 4.2, 5.8])
    p('三个固定模型在32个相同初态上配对比较。组合版相对整动作版多保留七个采用记录，配对团队收益增加0.0279秒，初态自举95%区间为[0.0057, 0.0562]秒。'
      '同一个有害采用案例仍存在，新增组合没有消除全部扰动退化。上述模型—初态记录有关联，1/46不构成条件风险证书；这些数据只用于开发与机制消融。')

    page('12 条件风险校准与匹配方法对照')
    p('定义A为闭环确实采用过学习动作，H为任一车辆成本或时间高于配对基线，或发生未完成、碰撞、中止、约束违反。'
      '对每个冻结完整控制器，分别考察事先随机选定扰动模式的一次运行，以及四种模式组成的压力组。压力组以任一模式采用为分母，以任一模式退化或失败为事件。')
    eq('r=P(H=1∣A=1)')
    p('两个分母不能互换，压力组风险也不能自动解释成每种模式的采用后风险。统计协议固定九个候选、两个目标，风险上限为5%，总错误认证预算为5%。'
      '每个候选与目标分配1/360预算，通过精确二项检验构造Clopper–Pearson单侧上界；这是Learn then Test的应用[7]。')
    eq('P(Binomial(N,0.05)≤K)≤1/360')
    p('N为采用过学习的独立初态数，K为其中退化或失败的数量。零退化时也至少需要115个采用初态；零采用不能认证。'
      '未通过检验可能来自证据不足，不能直接判为真实风险超过5%。不能将同一初态的车辆、模型、扰动或重复回放当作更多独立样本。')
    table(['部署选择器', '共同输入与选择', '保证对象'], [
        ['Split conformal', '全候选最大残差校准\n观测初态的完整名义预测', '新初态压力组的边际覆盖'],
        ['Learn then Test', '同训练模型内选择\n两类条件风险都通过的候选', '采用后的条件风险'],
        ['Mean SPI适配', '同配对样本的\n联合均值下界', '总体均值改善的渐近检验']], [4.1, 6.3, 6.0])
    p('三个选择器共用九个完整控制器、观测、基线、物理层及数据预算，不在训练种子之间挑优。候选包含原始学习、名义门控和严格组合门控。'
      '在线门控逐步选择动作组合，部署选择器决定运行哪一份完整控制器；完整轨迹不能按最终结果拼接。Mean SPI采用全信息配对仿真，未复现CSPI-MT的观察数据双重稳健估计[9]。')
    p('校准从名义与密集分布的等权混合中抽取1,024个IID初态，固定40,960次SUMO；最终确认使用另512个IID初态和20,480次SUMO。当前校准尚未完成，最终对照尚未执行。'
      '最终需同时报告两个风险目标、至少20%的实际采用目标、至少0.5%的团队收益目标及正的配对收益区间。'
      '56次人工校准接线审计仅验证部署选择与直接执行的一致性，不替代真实校准。')

    page('13 跨场景与动力学直接迁移')
    p('冻结权重与门控参数后，在名义、密集、低速让行、速度不对称、紧密跟车五类初态中，各取32个新初态，配对四类动力学与十个控制器，完成6,400次SUMO。'
      '四类动力学为名义模型、增益与二次阻力、轻度惯性、强惯性。各控制器使用相同隐藏参数及外生创新；实际加速度误差可以随命令变化。')
    p('名义动力学的外生创新范围为±0.5 m/s²，其余三类为±0.1 m/s²；跨动力学差异因此包含创新强度变化。'
      '轻度惯性τ为0.03–0.05秒，强惯性为0.3–0.5秒。解析误差界给出的统一充分条件为τ不大于约0.05456秒；它约束物理误差，不建立迁移后的性能风险证书。')
    table(['原误差范围内的控制器', '完成数', '未完成数'], [
        ['强反馈基线', '480 / 480', '0'], ['原始学习', '1429 / 1440', '11'],
        ['名义门控', '1440 / 1440', '0'], ['严格组合门控', '1440 / 1440', '0']], [7.0, 4.8, 4.6])
    p('前三类动力学共4,800次运行，实际误差均处于原±0.5 m/s²假设内。原始学习在低速让行中使平均成本增加8.14%–9.05%，因此独立名义研究中的收益不能推广到所有新场景。')
    table(['严格组合的成本改善', '名义模型', '增益阻力', '轻度惯性'], [
        ['名义初态', '−0.019%', '+0.111%', '+0.312%'], ['密集交通', '+0.898%', '+1.105%', '+0.770%'],
        ['低速让行', '+0.020%', '+0.212%', '+0.578%'], ['速度不对称', '+0.004%', '+0.015%', '+0.015%'],
        ['紧密跟车', '+1.385%', '+1.040%', '+1.735%']], [6.2, 3.4, 3.4, 3.4])
    p('密集交通与紧密跟车六个格子的逐格95%收益区间下界为正；区间只按每格32个独立初态自举，未作跨格同时校正。'
      '名义初态与名义动力学格存在微小负均值，不能声称处处不退化。严格组合在前三类动力学的214次采用记录中有12次退化，这些相关计数不证明5%的风险限额。')

    page('14 强惯性失败与工程适用范围')
    table(['强惯性控制器', '完成数', '中止数', '其余未完成'], [
        ['强反馈基线', '152 / 160', '8', '0'], ['原始学习', '456 / 480', '20', '4'],
        ['名义门控', '468 / 480', '12', '0'], ['严格组合门控', '467 / 480', '13', '0']], [6.1, 4.1, 3.1, 3.1])
    p('强惯性下1,600次运行均至少一次超出原物理误差范围，最大实际残差为3.10757 m/s²。共有57次未完成，其中53次在离开可恢复集合或找不到认证命令后中止。'
      '没有记录碰撞或路线约束违反，但中止后车辆的持续运动没有模拟，不能把终止仿真解释为已安全停稳。')
    p('全部运行保留原25秒时域截断成本。部分强惯性格的相对收益包含失败基线的截断成本，不能单凭较大的改善比例宣称抗失配有效。'
      '继续扩展这一范围需要显式建模执行器状态、重新构建可恢复集合，并验证接管后的物理过程。')
    p('按预先冻结的覆盖规则追加336次实际SUMO回放，全部68次未完成运行原样重现。'
      '291个保存有原始轨迹的案例逐步比较命令、状态、扰动、外生创新与执行器状态；另外45个案例核对完整指标并重建物理轨迹。'
      '回放验证执行一致性，不增加独立有效样本。')
    p('当前任务仍是三车固定路线纵向协同合流。物理过滤与名义预测假设状态可观测，系统没有覆盖视觉感知、自由换道、大车队、人类驾驶员或实车。'
      '共享状态与集中物理过滤器也是部署前提，不能把三个actor的独立参数解释成完全无需协调的系统。')

    page('15 完整门控调用的计算代价')
    p('严格扰动门控需要多次完整反馈续跑。精确命令缓存避免重复规划；单调运动模型给出的乐观到达时间下界可以提前拒绝必然失败的候选。'
      '该下界忽略车辆交互，只用于拒绝；采用动作仍须完成全部17场景检查。双线程版本将独立扰动续跑并行计算，保留候选和逻辑检查顺序。')
    measured = runtime['summary']
    serial, parallel = measured['bounded'], measured['parallel']
    table(['58次配对回放的测量', '串行下界版', '双线程版'], [
        ['总墙钟时间 ms', f"{serial['total_wall_ms']:.2f}", f"{parallel['total_wall_ms']:.2f}"],
        ['控制器进程CPU时间 ms', f"{serial['total_cpu_ms']:.3f}", f"{parallel['total_cpu_ms']:.3f}"],
        ['P95 ms', f"{serial['p95_ms']:.2f}", f"{parallel['p95_ms']:.2f}"],
        ['P99 ms', f"{serial['p99_ms']:.2f}", f"{parallel['p99_ms']:.2f}"],
        ['最大值 ms', f"{serial['max_ms']:.2f}", f"{parallel['max_ms']:.2f}"],
        ['超过500 ms的调用', '50 / 667', '34 / 667']], [7.2, 4.6, 4.6])
    p('冻结29个案例后，两种实现各执行29次SUMO，667个配对决策的动作、轨迹及准入证据完全一致。'
      '12个学习采用步都保留完整检查。总墙钟时间降低15.22%，进程CPU时间增加10.03%，额外计算28个未参与最终决定的扰动预测。'
      '单步最大缓存条目数由790增至2,262；条目计数不等同于实测内存字节数。')
    p('计时覆盖线程池建立、缓存复制、任务提交、等待、退出和返回信息构造；不计通信延迟，CPU计时不包括独立SUMO进程。'
      '函数预热、执行顺序交替，机器同时运行其他冻结研究。汇总P95虽略低于500 ms，仍有34次超时，第二个模型P95为548.19 ms。'
      '因此没有最坏情况实时保证，也不能把不同实验的缓存、下界和并行加速比相乘。')
    p('共形部署适配还需要从观测初态预测全部十个完整控制器。独立接线审计中的两次准备分别耗时28.75秒和17.73秒；仅报告所选策略的执行P95会遗漏这部分开销。'
      '主风险校准与最终测试仍使用原冻结控制器，等价工程回放不替代独立风险数据。')

    page('16 讨论与结论')
    p('已有结果支持联合检查的必要性及其在本任务中的实际收益。名义联合门控提高三个固定模型的平均团队收益，并使逐车均值改善；'
      '严格扰动组合进一步减少开发集中的有害采用，同时付出收益和计算代价。迁移实验显示收益集中于密集交通与紧密跟车，低速让行和强惯性揭示了适用边界。')
    p('理论保证分为三个层次：可恢复物理安全依赖误差与初始集合假设；逐车性能不退化依赖精确模型和有限扰动序列；'
      '实际采用后的条件风险需要独立统计校准。三者不能互相替代。当前完成前两类条件推导与闭环实验，第三类认证和最终匹配比较仍待完整数据。')
    p('后续结论应以冻结校准和最终确认结果为准。若所有候选都被拒绝，应报告零采用和零收益；若未通过风险检验，应结合有效采用数及上界说明证据是否不足。'
      '不能为了保留学习收益而修改当前阈值，也不能用最终测试数据继续调参后仍称其为独立确认。')
    p('本文将统计校准、策略改进恒等式和模型式多智能体训练作为已有工具。可检验的研究贡献在于完整联合组合的采用机制、逐车相对性能约束、'
      '有限扰动条件与实际风险的衔接，以及包含负结果的迁移和执行证据。这些机制的独立有效性仍需由完成的匹配对照和有效采用收益支撑。')
    p('代码、冻结模型、实验协议、逐回合记录和复核入口见公开仓库[11,12]。本研究尚未形成可用于开放道路的安全承诺。')

    doc.add_page_break()
    for element in retained_references:
        doc.element.body.insert(len(doc.element.body)-1, element)
    ref = p('[12] CoopGate. Disturbance composition, direct transfer, complete-call runtime and frozen risk protocol. '
            'https://github.com/Leo-7GOAT/CoopGate/tree/main/artifacts/v19')
    ref.paragraph_format.first_line_indent = Cm(0)
    for run in ref.runs: run.font.size = Pt(10)
    assert doc.sections[0]._sectPr.xml == original_section
    # All original equations occur before the replaced block and remain intact.
    new_math = list(doc.element.body.iter(qn('m:oMath')))
    assert [tostring(e, method='c14n', exclusive=True) for e in new_math[:len(original_math)]] == [
        tostring(e, method='c14n', exclusive=True) for e in original_math]
    doc.save(out)
    assert digest(reference) == REFERENCE_SHA
    with ZipFile(reference) as old, ZipFile(out) as new:
        assert set(old.namelist()) == set(new.namelist())
        changed = [name for name in old.namelist() if old.read(name) != new.read(name)]
        assert set(changed) <= {'word/document.xml', 'docProps/core.xml'}, changed
    record = dict(document_name=out.name, document_sha256=digest(out), reference_sha256=REFERENCE_SHA,
                  changed_parts=changed, completed_evidence=evidence, source_sha256=digest(__file__),
                  calibration_outcomes_read=False, conditional_risk_certified=False, final_comparison_completed=False,
                  hard_real_time_certified=False, expected_pages='approximately 17; requires rendered page review')
    evidence_out.parent.mkdir(parents=True, exist_ok=True)
    evidence_out.write_text(json.dumps(record, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    print(json.dumps(record, ensure_ascii=False))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--reference', type=Path, required=True)
    parser.add_argument('--out', type=Path, required=True)
    parser.add_argument('--evidence-out', type=Path, required=True)
    args = parser.parse_args()
    build(args.reference, args.out, args.evidence_out)
