"""Substantive V11 manuscript using the retained document's native design."""
import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
from zipfile import ZipFile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table

ROOT=Path(__file__).resolve().parents[1]


def build(reference,out):
    before_hash=hashlib.sha256(reference.read_bytes()).hexdigest();doc=Document(reference)
    template=deepcopy(doc.tables[0]._tbl)
    r=json.loads((ROOT/'artifacts/v11/summary.json').read_text());main=r['overall']['smooth_learned'];gain=r['contrasts_perturbed']['all']['v10_rollout']
    title='有界执行误差下的合流控制\n联合安全证书与指令平顺优化'
    doc.paragraphs[0].text=title;doc.core_properties.title=title.replace('\n',' ');doc.core_properties.version='11'
    cut=doc.paragraphs[3]._p;remove=False
    for e in list(doc.element.body):
        if e is cut:remove=True
        if remove and e.tag!=qn('w:sectPr'):doc.element.body.remove(e)

    def p(text,style=None):
        a=doc.add_paragraph(text,style)
        if style is None:a.paragraph_format.first_line_indent=Cm(.74)
        return a
    def h(text):doc.add_heading(text,1)
    def page(text):doc.add_page_break();h(text)
    def eq(text):
        a=doc.add_paragraph();a.alignment=WD_ALIGN_PARAGRAPH.CENTER;m=OxmlElement('m:oMath')
        def run(parent,value):
            z=OxmlElement('m:r');t=OxmlElement('m:t');t.text=value;t.set(qn('xml:space'),'preserve');z.append(t);parent.append(z)
        pos=0
        for token in re.finditer(r'([A-Za-zΣΠ≤≥0-9]+)(?:_\{([^}]+)\}(?:\^\{([^}]+)\})?|\^\{([^}]+)\})',text):
            run(m,text[pos:token.start()]);sub=token.group(2);sup=token.group(3) or token.group(4)
            z=OxmlElement('m:sSubSup' if sub and sup else 'm:sSub' if sub else 'm:sSup')
            base=OxmlElement('m:e');run(base,token.group(1));z.append(base)
            if sub:
                index=OxmlElement('m:sub');run(index,sub);z.append(index)
            if sup:
                index=OxmlElement('m:sup');run(index,sup);z.append(index)
            m.append(z);pos=token.end()
        run(m,text[pos:]);a._p.append(m)
    def table(headers,rows,widths=None):
        element=deepcopy(template);doc.element.body.insert(len(doc.element.body)-1,element);t=Table(element,doc._body)
        while len(t.rows)>1:t._tbl.remove(t.rows[-1]._tr)
        while len(t.columns)<len(headers):t.add_column(Cm(2.5))
        t.autofit=False;widths=widths or [8.2,4.1,4.1]
        for cell,value in zip(t.rows[0].cells,headers):cell.text=value
        header_fill=deepcopy(t.rows[0].cells[0]._tc.get_or_add_tcPr().find(qn('w:shd')))
        if header_fill is not None:
            for cell in t.rows[0].cells:
                props=cell._tc.get_or_add_tcPr();old_fill=props.find(qn('w:shd'))
                if old_fill is not None:props.remove(old_fill)
                props.append(deepcopy(header_fill))
        for values in rows:
            for cell,value in zip(t.add_row().cells,values):cell.text=str(value)
        for col,width in zip(t.columns,widths):col.width=Cm(width)
        for i,row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell,width in zip(row.cells,widths):
                cell.width=Cm(width)
                for a in cell.paragraphs:
                    a.alignment=WD_ALIGN_PARAGRAPH.CENTER;a.paragraph_format.space_before=a.paragraph_format.space_after=Pt(5)
                    a.paragraph_format.line_spacing=1.05;a.paragraph_format.keep_with_next=i<len(t.rows)-1
                    for run in a.runs:run.font.size=Pt(10);run.bold=i==0
        doc.add_paragraph().paragraph_format.space_after=Pt(2)

    p('针对安全约束下合流控制中加减速反复的问题，本文在有界执行误差的联合备用证书内，研究到达时间与累计指令变化量的分层优化。'
      '以固定指令格点将两个代价写为整数，优先缩短到达时间，在到达时间相同时减少指令变化。'
      '为原有加速或制动基控制器增加保持指令的类型，两类分别进行坐标前瞻，再选择较优的获证书方案。')
    p(f"在200个全新初态、四类执行误差、九种方法共7200次SUMO执行中，完整方法完成{main['completed']}/800回合，"
      f"碰撞{main['collision_episodes']}回合、路线合同违规{main['contract_violation_episodes']}回合、控制器中止{main['aborted']}回合。"
      f"非零误差下相对上一版到达时间优化器，到达代价减少{gain['time']['reduction_percent']:.2f}%，累计指令变化量减少{gain['variation']['reduction_percent']:.2f}%。"
      f"暖机控制计算周期P95为{main['warm_controller_cycle_p95_ms']:.2f} ms。本文分别证明条件鲁棒安全、名义字典序性能界与剪枝等价性；实际扰动下的双指标改善由配对实验检验。")
    p('关键词  合流控制  安全过滤  字典序优化  指令变化量  多策略前瞻')
    h('1 引言与研究定位')
    p('学习策略与传统控制之间的切换可以借助物理证书约束风险，但安全过滤和到达时间优化并不自动给出平顺动作。'
      '此前控制器只在到达周期数严格减少时更新方案，到达时间相同的温和动作不能取代原指令。'
      '遇到合流冲突时，加速与备用制动之间的粗粒度选择又可能引发不必要的减速和再加速。')
    p('预测安全过滤已有备用安全控制思想[1]，多智能体坐标rollout已具有基策略改进性质[2,3]，带容差的字典序方法也已用于车辆多目标规划[5]。'
      '本文沿用这些基础，在带指令记忆、联合冲突合同和有界执行误差的任务中，实现可核验的分层前瞻及两类基控制器改进。'
      '字典序优化和rollout不是本文首次提出，学习模块也仍是此前冻结的集中式上下文bandit。')

    page('2 动力学与两个任务指标')
    p('三车路线标签为(0,1,0)，同路线车辆0位于车辆2前方。状态包括位置s、速度v、上一指令及全局决策周期k。'
      '控制周期Δ=0.5 s，子步h=0.1 s，周期内保持指令u，执行误差w在每个子步可变化。')
    eq('v_{i,t+1} = clip(v_{i,t} + h(u_{i,k} + w_{i,t}), 0, v_{max})')
    eq('s_{i,t+1} = s_{i,t} + h v_{i,t+1},     |w_{i,t}| ≤ ε')
    table(['参数','含义','取值'],[
        ['速度上限','v上限','20 m/s'],['指令范围','制动与加速','[−5, 2.5] m/s²'],
        ['指令变化率上限','每周期限幅','5 m/s³'],['已知执行误差界','ε','0.5 m/s²'],
        ['路线间距与冲突区','合同参数','6.8 m / [−25, 6.8] m'],
        ['任务终点与截止时间','路线坐标及时间','60 m / 25 s'],['固定指令量子','q','0.625 m/s²']])
    p('任务终点为s≥60。令每周期未到达车辆数为ℓ，令指令变化的格点单位数为d；累计目标向量为：')
    eq('ℓ(x) = Σ_{i} 1[s_{i} < 60],     d(u,u^{-}) = Σ_{i} |u_{i} − u_{i}^{-}| / q')
    eq('J = (T,D) = Σ_{k=0}^{N−1} (ℓ(x_{k}), d(u_{k},u_{k}^{-}))')
    p('三车全部到达后剩余代价置零。T乘Δ/3为截断平均到达时间，D乘q为三车累计指令变化量，单位m/s²。'
      '初始指令为零；已到达车辆继续受控，变化量计到最后一车到达或截止。该指标不除以回合长度，避免用更慢行驶稀释变化量。'
      '另外报告平均指令变化率，单位m/s³；二者都不等于含执行误差的实际加速度变化或实测乘坐舒适度。')

    page('3 联合安全证书与保持指令')
    p('同路线车辆保持至少6.8 m间距；异路线车辆不能同时占据冲突区，双方进入受保护进路后还需保持路线间距。'
      '每个子步按位置端点间的仿射插值检查整段合同。SUMO碰撞事件单独记录，合同保证限定于这个模型。')
    p('沿用冻结V9区间证书。以u−ε和u+ε传播速度与位置上下界，对任意界内执行误差包络真实状态。候选指令执行一个周期后，追加联合制动备用律：')
    eq('B_{i}(u^{-}) = max(−b, u_{i}^{-} − j_{max} Δ)')
    p('候选及最多24个备用周期都需通过区间合同检查，终端要求所有速度上界为零且最后指令不大于−ε。'
      '此时任何允许误差都不能使车辆离开静止状态，得到停止不变集。区间管为保守外包，不是精确可达集。')
    p('命题1  条件鲁棒安全。若初态可恢复、测量准确且执行误差满足声明界，获证书动作保持路线合同，下一实测状态仍有可行制动后缀。')
    p('证明：基动作及每次坐标替换均核验完整三车指令。真实后继包含于获证书动作的一周期区间内；由该实测点重建的备用后缀被旧证书后缀包络包含，'
      '合同与停止终端条件继续成立。逐周期归纳即可。证明依赖声明动力学及数值实现条件，浮点误差没有经过形式化验证。')
    p('原二值基控制器每车只能从提案动作a与备用制动B中选择。新增保持型控制器还允许保持上一指令u⁻：')
    eq('A_{binary} = Π_{i} {a_{i},B_{i}},     A_{hold} = Π_{i} {a_{i},u_{i}^{-},B_{i}}')
    p('两类分别至多8和27个联合动作，均按优先级加权提案偏离选择，并逐个接受同一联合证书。'
      '保持动作提供较小的变化选择，但只有真实通过联合检查时才能使用。所有安全参数沿用旧版，不以放松间距或误差界换取收益。')

    page('4 两类基策略的分层动作改进')
    p('冻结网络根据初态提出策略索引pθ，与传统索引0组成集合P；索引指定车间优先级和初期滑行时长。'
      '同一P用于二值与保持型两类基控制器。网络不读取未来误差带，也不在测试阶段更新。所有预测使用全局k，不重启滑行时钟。')
    p('采用字典序比较两分量代价：先最小化到达周期数T，只有T相等才最小化指令变化单位数D。'
      '加速、制动和最大指令增量均为q的整数倍，控制记忆与候选也在同一格点上，两个目标均为整数。')
    eq('(T_{1},D_{1}) ≺ (T_{2},D_{2})  ⇔  T_{1}<T_{2}  or  (T_{1}=T_{2} and D_{1}<D_{2})')
    p('在每种控制器类型f内部，评价传统和学习提案的完整名义余程，选字典序较优的p*f。'
      '两类各自以该基策略首动作作为incumbent，按优先级权重降序逐车扫描；同权重按车号排列。')
    eq('L_{i} = max(−b,u_{i}^{-}−j_{max}Δ),     U_{i} = min(a_{max},u_{i}^{-}+j_{max}Δ)')
    p('枚举允许区间内的固定格点qn，从大到小试探；每次仅改变一个分量，重新检查整个联合指令。'
      '通过安全证书后执行名义一步，沿当前类型的p*f评价完整尾部：')
    eq('Q_{k}^{p,f}(z,u) = g(z,u) + V_{k+1}^{p,f}(F_{0}(z,u))')
    p('只接收严格字典序改进，完全平局保留旧动作。最后比较原始最小基策略值和两类扫描结果，执行最优获证书动作。'
      '每车每类最多9个格点，两类单轮扫描最多54个试探；已到达车与重复指令跳过。没有联合网格全局最优保证，尾部评价和证书成本另计。')
    p('先比较未经改进的尾部、只优化胜出的类型，可能错过另一类型经动作改进后的更优方案。'
      '因此完整方法保留两类各自的扫描，实验以只扫描原始较优类型作为消融，衡量额外计算是否值得。')

    page('5 名义分层性能界')
    p('命题2  名义字典序不退化。令W为同一时刻所有基策略及类型的最小完整尾部代价，ρ为上述闭环控制器。'
      '在名义零误差、相同截止时间、精确评价及全局时钟下，ρ的两分量闭环代价不高于W的字典序。')
    eq('W_{k}(z) = min_{p∈P,f}^{lex} V_{k}^{p,f}(z)')
    eq('V_{k}^{ρ}(z) ≤_{lex} W_{k}(z)')
    p('证明：每类扫描的初始incumbent等于该类较优基策略首动作，其完整Q等于该基策略值；所有替换只接受严格字典序改进。'
      '最终还保留全局原始incumbent，故存在一个获选尾部策略(p*,f*)满足：')
    eq('g(z,ρ(z)) + V_{k+1}^{p*,f*}(F_{0}(z,ρ(z))) ≤_{lex} W_{k}(z)')
    p('截止时刻闭环与所有基策略剩余代价均为零。假设下一时刻闭环代价不高于W，而W不高于获选基策略的尾部，'
      '字典序在向量加法下保持次序，代入上式即得当前不等式。反向归纳完成。命题1保证证书内的后续基策略可执行。')
    p('原V9二值基策略仍在集合内，所以名义到达代价不高于V9学习双方案。若闭环到达代价恰等于最优基策略的到达代价，'
      '总变化量也不大于对应最小值。保留全局时钟与完整截止时域，是这个递推成立的条件。')
    p('该命题有三个直接边界。第一，字典序改善不等于两个分量分别下降：时间更短时，变化量仍可能变大。'
      '第二，比较对象为基策略集合，不包括上一版V10的完整在线优化闭环，因此不能由命题推出对V10的支配。'
      '第三，真实含扰动后继与名义后继不同，扰动下实际性能也不受该界保证。双指标改善及对V10的优势必须由独立实验检验。')
    p('安全保证与性能界对应不同命题：前者允许声明界内逐子步扰动，后者仅限名义模型。'
      '学习器即使提出较差策略，也不能跳过物理证书；网络分数没有被解释为安全概率。')

    page('6 保持选择的计算缩减')
    p('命题3  字典序下界剪枝。尾部已累计代价为(T,D)，令L为忽略车间冲突时、受相同变化率约束的独立最早到达周期总数。'
      '则(T+L,D)为完整尾部的字典序下界。当它不优于当前incumbent时，停止评价不改变最终选择。')
    eq('(T+L,D) ≥_{lex} Q_{best}  ⇒  reject')
    p('证明：每车采用最大允许加速包络，速度饱和及位置更新单调，因此独立车辆的到达时间不晚于受车间约束的车辆。'
      '累计时间加L给出第一分量下界，未来变化量非负给出第二分量下界。若下界已不能严格改善，完整尾部也不能严格改善。')
    p('与只优化时间不同，到达下界相等时不能立即剪掉候选，还要比较变化量下界。'
      '否则会错误删除时间相同但动作更平顺的方案。实现以两个整数比较，关闭剪枝的同一算法作为差分参考。')
    p('命题4  稳定排序后的联合核验。有限基动作集合先按性能分数稳定排序，再返回第一个通过证书的动作，'
      '与完整核验所有动作后取最小分数完全等价。')
    p('证明：排在返回动作前的候选全部不满足证书，之后的候选分数不更小；稳定排序保持原枚举平局顺序。'
      '这项缩减只改变核验顺序，不假设安全集合关于加速度单调，也没有用低概率评分免除核验。')
    p('二值动作分数保持与冻结V9相同的浮点运算次序，并直接对照原实现；保持型以独立Python枚举全部27动作作为参考。'
      '开发阶段对48个旧初态、四类误差共192回合复跑，计算缩减前后任务代价、变化量、完成情况和最大指令变化率完全相同。')
    audit=json.loads((ROOT/'artifacts/v11/algorithm_audit.json').read_text())
    p(f"正式实验后，按固定抽取规则审计首初态的八条轨迹，共{audit['diagnostic_states']}个相关状态。"
      f"剪枝开关的指令不一致为{audit['command_mismatches']}，值不一致为{audit['value_mismatches']}；有限基动作核验不一致为{audit['projection_mismatches']}。"
      f"剪枝减少{audit['tail_step_reduction_percent']:.2f}%的尾部反馈步数。这些状态来自已有轨迹，不计作新的独立场景。")

    page('7 冻结实验与统计方法')
    p('开发使用此前公开的V10初态：先16个初态320回合，再48个初态1152回合；随后对相同48个初态分别进行两类扫描和等价计算缩减验证，各192回合。'
      '代码、源码哈希、验收标准及200个新测试种子在正式实验前冻结，提交为7fc085e[4]。这些旧场景不进入主要结果。')
    p('正式实验为常规与高密度各100个全新初态。生成器初速度分别为8至14和12至17 m/s，前两车插入位置分别位于距汇入口70至110和70至85 m范围，'
      '同路线后车位于前车之后30至45 m。SUMO插入步后的实际初态另存。每个初态测试四类执行误差、九种方法，共7200次实际执行。')
    table(['执行误差','子步规则','每方法回合数'],[
        ['零误差','始终为0',200],['随机误差','独立均匀于[−0.5,0.5]',200],
        ['持续偏差','每车固定正或负极值',200],['间歇脉冲','1秒极值与1秒空隙',200]])
    p('V9、V10、完整V11、无学习、固定提案、随机提案，以及移除保持动作、移除变化量目标和只优化较优基策略族的消融均使用同一安全证书。'
      '各方法共享初态及外生误差带，执行顺序随机化。程序单工作进程运行，预热后测时；不读取未来误差，不删除失败回合。')
    p('两个共同主要终点预先指定为非零误差下V11对V10的到达时间与总变化量。每个初态先对三种非零误差取配对差均值，'
      '再对200个初态整簇bootstrap一万次。两个终点各报97.5%百分位区间，以Bonferroni式分配置信水平；这是近似推断，并非分布无关的有限样本保证。'
      '其余对照及95%区间仅作描述性分析。')
    p('十项验收标准为：零碰撞、合同违规及中止；完成率100%；扰动时间改善至少1%；总变化量减少至少25%；'
      '两个主要终点97.5%区间下界各大于0；平均指令变化率减少至少25%；零误差下相对V9更差初态为0；'
      '暖机控制周期P95不超过20 ms；首响应P95不超过35 ms。所有标准在新实验前固定。')

    page('8 独立测试结果与不利场景')
    names=[('v9_learned_pair','V9学习双方案'),('v10_rollout','V10时间优化'),('smooth_constant','无学习'),
        ('smooth_fixed_pair','固定提案'),('smooth_random_pair','随机提案'),('smooth_binary','移除保持动作'),
        ('smooth_time_only','移除变化量目标'),('smooth_single_family','只优化较优类型'),('smooth_learned','V11完整方法')]
    p('表1  三类非零执行误差下的平均任务指标  每种方法600回合','Caption')
    table(['方法','到达时间 s','总变化量 m/s²','变化率 m/s³'],[
        [name,f"{r['perturbed'][m]['mean_cost_s']:.3f}",f"{r['perturbed'][m]['mean_total_variation']:.3f}",
            f"{r['perturbed'][m]['mean_command_jerk']:.3f}"] for m,name in names],[6.2,3.4,3.4,3.4])
    p(f"完整方法完成{main['completed']}/800回合，碰撞{main['collision_episodes']}回合、合同违规{main['contract_violation_episodes']}回合、"
      f"中止{main['aborted']}回合、超时{main['timeouts']}回合。故障事件取并集，避免重复计数。十项预设标准全部通过：{'是' if r['all_criteria_passed'] else '否'}。")
    p(f"预设标准实际通过{sum(r['criteria'].values())}/10项。到达时间减少{gain['time']['reduction_percent']:.4f}%，"
      '低于预先规定的至少1%，这一项保留为未通过；不以四舍五入或修改阈值将其改判。')
    for endpoint,label,unit in [('time','到达时间','s'),('variation','累计指令变化量','m/s²')]:
        e=gain[endpoint];ci=e['paired_cluster_ci97_5'];counts=e['better_tied_worse_clusters']
        p(f"相对V10，{label}平均减少{e['mean_reduction']:.3f} {unit}，相对减少{e['reduction_percent']:.2f}%；"
          f"97.5%配对整簇区间为[{ci[0]:.3f}, {ci[1]:.3f}] {unit}。按各初态三类扰动平均差，更好、持平、更差分别为{counts[0]}、{counts[1]}、{counts[2]}个。")
    p(f"相对V10，常规与高密度场景的到达代价分别减少{r['contrasts_perturbed']['nominal']['v10_rollout']['time']['reduction_percent']:.2f}%和"
      f"{r['contrasts_perturbed']['dense']['v10_rollout']['time']['reduction_percent']:.2f}%。零误差下相对V9更差初态为{r['clean_worse_than_v9']}/200。"
      '所有扰动退化初态保留在原始数据中，均值改进不能解释为逐场景支配。')

    page('9 消融结果与工程计算成本')
    nm=dict(names);ablations=['smooth_binary','smooth_time_only','smooth_single_family','smooth_constant','smooth_fixed_pair','smooth_random_pair']
    p('表2  完整方法相对各消融的平均减少百分比  正数表示完整方法更优','Caption')
    table(['对照方法','到达时间 %','总变化量 %'],[[nm[m],f"{r['contrasts_perturbed']['all'][m]['time']['reduction_percent']:.2f}",
        f"{r['contrasts_perturbed']['all'][m]['variation']['reduction_percent']:.2f}"] for m in ablations])
    learning=r['contrasts_perturbed']['all']['smooth_constant']['time'];ci=learning['paired_cluster_ci95']
    p(f"相对无学习版本，学习提案的额外时间收益为{learning['mean_reduction']:.3f} s，描述性95%区间为[{ci[0]:.3f}, {ci[1]:.3f}] s。"+
      ('区间跨过零，当前证据未可靠建立学习模块的独立通行收益。' if ci[0]<=0<=ci[1] else '该对照单独衡量学习选择，不能代表完整优化器的全部收益。')+
      '固定与随机提案使用相同的两策略预算；学习网络没有重新训练。')
    p('表3  单机暖机计算分位数  单位为毫秒','Caption')
    table(['方法','首响应P95','控制周期P95'],[[name,f"{r['overall'][m]['warm_first_response_p95_ms']:.2f}",
        f"{r['overall'][m]['warm_controller_cycle_p95_ms']:.2f}"] for m,name in names])
    p('首响应包含初始网络提案、V9参考方案评价和第一次TraCI执行；控制周期只统计算法。'
      '全部TraCI周期另存。SUMO在计算期间暂停，分位数不是最坏执行时间，也没有加入通信延迟统计。')

    page('10 讨论与研究边界')
    p(f"当前改进针对此前更快但指令反复更明显的问题。在同一批新场景中，到达代价与总变化量分别减少"
      f"{gain['time']['reduction_percent']:.2f}%和{gain['variation']['reduction_percent']:.2f}%，平均指令变化率减少{gain['jerk']['reduction_percent']:.2f}%。"
      '两级目标、保持指令和逐类扫描的作用由消融拆分；计算缩减另以不改变选择的差分审计验证。')
    p('工程实现提供冻结策略、同场景误差带、完整回合记录、原生SUMO演示及指令复跑。'
      '源码与实验协议先于新实验提交，旧版本及其不利结果保留。原始数据同时记录中止、超时、碰撞、合同违规、指令变化和计算时间，'
      '而不是仅保留成功展示轨迹。')
    p('安全仍依赖已知执行误差界、精确状态测量和三车合作固定路线。没有验证更多参与者、非合作车辆、感知误差、未知误差界或实车。'
      '使用TraCI速度接口的子步模型也没有涵盖完整车辆横纵向动力学，指令变化下降不能直接宣称实车舒适度提升。')
    p('名义字典序性能界与误差下的双指标实测分别成立在各自条件下。更快不必然更平顺，总体平均改善也不意味着每个初态更好。'
      '置信区间基于当前场景生成分布及有限独立初态，不能推出开放交通环境中的同等收益。')
    p('Hu等已以带容差的字典序方法处理车辆效率、能耗与平顺性，当前实验使用记录轨迹开环回放[5]。'
      '本文关注三车合作反馈、执行误差及备用证书递推，没有覆盖其完整重型车能耗与横纵向任务，也没有同预算复现，不能直接比较数值或声称领先。'
      '投稿前仍需补充已发表方法对照和更大规模交通场景。')
    h('参考文献')
    references=[
        '[1] Wabersich K P, Zeilinger M N. A predictive safety filter for learning-based control of constrained nonlinear dynamical systems. Automatica, 2021, 129:109597. https://arxiv.org/abs/1812.05506',
        '[2] Bertsekas D. Multiagent Rollout Algorithms and Reinforcement Learning. arXiv:1910.00120, revised 2020. https://arxiv.org/abs/1910.00120',
        '[3] Bhattacharya S, Kailas S, Badyal S, Gil S, Bertsekas D. Multiagent Rollout and Policy Iteration for POMDP with Application to Multi-Robot Repair Problems. CoRL 2020, PMLR 155:1814–1828, 2021. https://proceedings.mlr.press/v155/bhattacharya21a.html',
        '[4] CoopGate. Frozen V11 protocol, paired SUMO study and raw records. Source freeze 7fc085e. https://github.com/Leo-7GOAT/CoopGate/tree/main/artifacts/v11',
        '[5] Hu X, Tong J, E W, Liu M, Hua M, Wang X. Rolling Multi-Objective Trajectory Planning for Connected and Automated Heavy Vehicles Based on the Tolerance Lexicographic Method. IET Intelligent Transport Systems, 2026, 20(1):e70259. https://doi.org/10.1049/itr2.70259']
    for text in references:
        a=p(text);a.paragraph_format.first_line_indent=Cm(0);a.paragraph_format.space_after=Pt(4)
        for run in a.runs:run.font.size=Pt(9)
    doc.save(out)
    assert hashlib.sha256(reference.read_bytes()).hexdigest()==before_hash
    with ZipFile(reference) as old,ZipFile(out) as new:
        changed=[name for name in old.namelist() if old.read(name)!=new.read(name)]
        assert set(changed)<={'word/document.xml','docProps/core.xml','word/_rels/document.xml.rels','[Content_Types].xml'},changed
    print(json.dumps({'output':str(out),'reference_sha256':before_hash,'changed_parts':changed}))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--reference',type=Path,required=True);parser.add_argument('--out',type=Path,required=True)
    args=parser.parse_args();build(args.reference,args.out)
