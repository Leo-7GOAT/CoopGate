"""Substantive V9 manuscript using the retained document's native design."""
import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import re
from zipfile import ZipFile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table

ROOT=Path(__file__).resolve().parents[1]


def build(reference,out):
    before_hash=hashlib.sha256(reference.read_bytes()).hexdigest();doc=Document(reference)
    template=deepcopy(doc.tables[0]._tbl)
    r=json.loads((ROOT/'artifacts/v9/summary.json').read_text());main=r['overall']['robust_learned_pair'];gain=r['contrasts_perturbed']['all']['robust_constant']
    title='有界执行误差下的合流规划\n学习提案与鲁棒反馈切换'
    doc.paragraphs[0].text=title;doc.core_properties.title=title.replace('\n',' ');doc.core_properties.version='9'
    cut=doc.paragraphs[3]._p;remove=False
    for e in list(doc.element.body):
        if e is cut:remove=True
        if remove and e.tag!=qn('w:sectPr'):doc.element.body.remove(e)

    def p(text,style=None):
        a=doc.add_paragraph(text,style)
        if style is None:a.paragraph_format.first_line_indent=Cm(.74)
        return a
    def h(text):doc.add_heading(text,1)
    def page(text):doc.add_page_break();h(text)
    def eq(text):
        a=doc.add_paragraph();a.alignment=WD_ALIGN_PARAGRAPH.CENTER;m=OxmlElement('m:oMath')
        def run(parent,value):
            z=OxmlElement('m:r');t=OxmlElement('m:t');t.text=value;t.set(qn('xml:space'),'preserve');z.append(t);parent.append(z)
        pos=0
        for token in re.finditer(r'([A-Za-zΣ0-9]+)(?:_\{([^}]+)\}(?:\^\{([^}]+)\})?|\^\{([^}]+)\})',text):
            run(m,text[pos:token.start()]);sub=token.group(2);sup=token.group(3) or token.group(4)
            z=OxmlElement('m:sSubSup' if sub and sup else 'm:sSub' if sub else 'm:sSup')
            base=OxmlElement('m:e');run(base,token.group(1));z.append(base)
            if sub:
                index=OxmlElement('m:sub');run(index,sub);z.append(index)
            if sup:
                index=OxmlElement('m:sup');run(index,sup);z.append(index)
            m.append(z);pos=token.end()
        run(m,text[pos:]);a._p.append(m)
    def table(headers,rows,widths=None):
        element=deepcopy(template);doc.element.body.insert(len(doc.element.body)-1,element);t=Table(element,doc._body)
        while len(t.rows)>1:t._tbl.remove(t.rows[-1]._tr)
        while len(t.columns)<len(headers):t.add_column(Cm(2.5))
        t.autofit=False;widths=widths or [8.2,4.1,4.1]
        for cell,value in zip(t.rows[0].cells,headers):cell.text=value
        header_fill=deepcopy(t.rows[0].cells[0]._tc.get_or_add_tcPr().find(qn('w:shd')))
        if header_fill is not None:
            for cell in t.rows[0].cells:
                props=cell._tc.get_or_add_tcPr();old_fill=props.find(qn('w:shd'))
                if old_fill is not None:props.remove(old_fill)
                props.append(deepcopy(header_fill))
        for values in rows:
            for cell,value in zip(t.add_row().cells,values):cell.text=str(value)
        for col,width in zip(t.columns,widths):col.width=Cm(width)
        for i,row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell,width in zip(row.cells,widths):
                cell.width=Cm(width)
                for a in cell.paragraphs:
                    a.alignment=WD_ALIGN_PARAGRAPH.CENTER;a.paragraph_format.space_before=a.paragraph_format.space_after=Pt(5)
                    a.paragraph_format.line_spacing=1.05;a.paragraph_format.keep_with_next=i<len(t.rows)-1
                    for run in a.runs:run.font.size=Pt(10);run.bold=i==0
        doc.add_paragraph().paragraph_format.space_after=Pt(2)

    p('针对模型偏差导致学习规划缓存失效、传统回退失去适用保证的问题，本文研究三车固定路线合流中的有界加速度执行误差。'
      '将误差包络传播与带指令记忆的联合制动后缀结合，为路线间隔和冲突区占用构造鲁棒可恢复性证书。'
      '学习器提出反馈方案，与传统基线进行名义模型比较；实际执行按实测状态重新核验，不因轨迹偏差永久丢弃反馈策略。')
    p(f"在200个全新初态、四类误差及七种方法共5600次SUMO执行中，学习提案双方案完成{main['completed']}/800回合，"
      f"记录中止{main['aborted']}次、合同违规{main['contract_violation_episodes']}回合、碰撞{main['collision_episodes']}回合。"
      f"三类非零扰动下，相对同证书传统基线的平均截断到达代价改善{gain['gain_percent']:.2f}%，"
      f"暖机首次响应P95为{main['warm_first_response_p95_ms']:.2f} ms。证明限于已知误差界、精确测量和声明的路线合同；"
      '扰动下的实际通行性能仍由配对实验评估。')
    p('关键词  合流规划  有界执行误差  区间可达集  联合安全证书  学习提案')
    h('1 引言与相关研究')
    p('学习控制与传统规划的切换至少包含两个问题：提案是否值得执行，以及执行时是否仍满足安全条件。'
      '已有原型采用完整模型核验和缓存执行，在精确动力学中可避免候选库内的性能退化；但状态与预测不一致后永久回退，'
      '其备用保证也依赖精确模型。仅降低计算时间不能弥补这一模型假设缺口。')
    p('预测安全过滤器已利用备用策略处理学习控制的约束与模型不确定性[1]，自适应模型预测安全认证已研究有界扰动与鲁棒管[2]。'
      '动态模型预测屏障还关注安全恢复期间的任务进展[3]。本文不把鲁棒过滤概念作为首次创新，'
      '而给出适配联合合流合同的可检查区间证书，并通过相同过滤器下的固定、随机和学习提案消融，分离学习与安全机制的贡献。')

    page('2 问题定义与安全合同')
    p('车辆路线标签为(0,1,0)，同路线车辆0始终在车辆2之前；位置s为距汇入口的路线坐标。'
      '控制周期Δ=0.5 s，子步h=0.1 s，速度上限20 m/s，指令范围[−5,2.5] m/s²，指令变化率不超过5 m/s³。'
      '每个周期保持指令u不变，每个子步可有独立执行误差w。')
    eq('v_{i,t+1} = clip(v_{i,t} + h(u_{i,k} + w_{i,t}), 0, v_{max})')
    eq('s_{i,t+1} = s_{i,t} + h v_{i,t+1},     |w_{i,t}| ≤ ε')
    p('取ε=0.5 m/s²，属于事先指定的模型误差界。它不是由测试样本估出的置信区间。'
      '指令限幅和变化率限制施加于u；实际加速度u+w及其变化率不享有相同的界。速度饱和是模型的一部分。')
    p('同路线车辆保持至少6.8 m间距。不同路线车辆不能同时占据冲突区Z=[−25,6.8]；'
      '双方进入s≥−25的受保护进路后，还需保持至少6.8 m路线坐标间距。子步内位置按端点进行仿射插值，'
      '安全检查覆盖该插值整段，不仅检查子步采样点。')
    p('仿真使用Eclipse SUMO和TraCI逐子步速度接口，记录车辆累计路程和速度。三车均由本系统管理，路线固定且关闭内部路口连接。'
      '证明对象是上述路线合同；SUMO碰撞事件另行记录。连续空间车身几何、测量误差和未知车辆行为尚未纳入证明。')
    p('任务代价为三车首次在决策时刻到达s≥60的平均时间；25 s未到达者记为25 s。'
      '已到达车辆继续巡航，避免因停车占据出口而获得不合理优势。控制器中止按25 s计入整回合任务代价，并单独报告。')

    page('3 区间可达包络与占用时间检查')
    p('从当前精确实测状态初始化退化区间。以下使用下标L、U分别表示下界和上界。'
      '速度饱和与位置递推均单调，可用u−ε和u+ε逐步传播区间端点。')
    eq('v_{L}^{+} = clip(v_{L} + h(u − ε), 0, v_{max})')
    eq('v_{U}^{+} = clip(v_{U} + h(u + ε), 0, v_{max})')
    eq('s_{L}^{+} = s_{L} + h v_{L}^{+},     s_{U}^{+} = s_{U} + h v_{U}^{+}')
    p('性质1：任意满足逐子步误差界的实际状态均包含在该区间中。证明：起点包含关系成立；'
      '每一步的真实输入位于[u−ε,u+ε]，单调的饱和速度与位置更新保持包含关系，逐步归纳得到结论。'
      '不要求扰动独立、同分布或均值为零。')
    p('同路线安全检查比较前车下界与后车上界，在子步两端均要求至少6.8 m间距。'
      '任意实际相对位置被这两个端点构成的仿射界约束，故覆盖整个子步。')
    p('不同路线的冲突占用检查使用上界轨迹求最早可能进入时间，使用下界轨迹求最晚可能离开时间。'
      '两个外包占用区间不相交，足以排除实际轨迹同时占据冲突区。'
      '从双方均可能进入受保护进路的最早时刻开始，还要求相对位置区间在时段两端保持同一顺序和足够间距。')
    p('这些条件是充分条件。区间舍弃车辆状态之间和位置速度之间的相关性，可能拒绝实际可行的方案；'
      '本文未把区间证书当作精确可达集。包络与物理约束的实现分别接受随机包含测试、极值扰动测试和零半径差分测试。')

    page('4 联合备用与递归可恢复性')
    p('上一控制指令作为扩展状态的一部分保留。备用律逐周期降低指令，直至三车停止：')
    eq('B_{i}(u^{-}) = max(−b, u_{i}^{-} − j_{max} Δ)')
    p('一个候选动作获得许可，需要先通过自身一个周期的区间合同检查，再通过最多24个联合制动备用周期的区间检查。'
      '只有所有速度上界精确为0、最终指令均不大于−ε时才接受终端集合。此时任意允许误差下速度仍被饱和在0，'
      '继续备用不会改变位置，因而终端集合保持不变。')
    p('性质2：若初始扩展状态存在通过证书的联合制动后缀，每次执行的动作均通过相同证书，且误差界、测量和路线假设成立，'
      '则每个执行子步保持路线合同，下一决策时刻仍存在可认证备用后缀。')
    p('证明：由性质1，实际第一周期轨迹包含于候选动作的区间包络，区间合同检查保证该周期安全。'
      '下一实测状态位于原预测包络内；从这一点沿同一备用后缀重新传播，所得包络包含于原后缀包络。'
      '缩小区间不会破坏已成立的占用不相交和间距条件。指令记忆已更新为实际执行指令，'
      '下一备用动作因此恰好接续原证书。到达终端集合后可无限延长备用，由归纳得到结论。')
    p('该性质允许逐周期重新选择任何获得许可的联合模式，不要求继续执行原始预测轨迹。'
      '状态与计划不一致时，控制器重新核验反馈提案；若实测状态在鲁棒备用集合之外，立即中止并报告，'
      '不将集合外动作伪装为安全回退。')
    p('实现采用float64，合同比较保留明确数值容差。数学论证针对声明的模型；数值实现通过差分和仿真审计，'
      '尚不是机器验证的浮点形式化证明。停止可恢复性也不推出有限时间任务完成，因此实验另外统计完成率与超时。')

    page('5 学习提案与反馈方案切换')
    p('候选库保持37项：传统基线一项，其余36项由三种先行偏好、三个权重比和四个滑行前缀组合。'
      '每项方案都是反馈控制器，在当前状态构造提案，并在提案或备用构成的八个联合模式中选择通过证书且偏离提案最小者。'
      '控制权重和前缀保持固定，状态和安全核验逐周期更新。')
    p('学习器仍为三个两层128单元网络组成的集成，输入三车位置、速度及交互特征，输出候选方案。'
      '检查点来自此前768个训练初态和128个验证初态，本轮没有重训。训练属于完整信息离线上下文bandit，'
      '不能将其称为新的分布式多智能体强化学习。')
    p('得到学习提案后，在w=0的名义模型中，用新的区间证书过滤反馈动作，分别展开传统基线与学习方案至完成或截止。'
      '名义代价严格更小时选择学习方案，否则选择传统基线；随后实际执行的每一步仍独立通过区间证书。')
    eq('j^{*} = arg min J_{nom}(j),     j ∈ {0, j_{learned}}')
    p('性质3：当实际误差为零、初态和反馈执行逻辑与名义评价一致时，双方案选择的任务代价不劣于保留的基线。'
      '证明：基线属于候选集，取最小值不会更差；一致执行复现被选方案的到达时间。'
      '存在扰动时，安全性质仍由性质2提供，但该性能结论不再成立，实际到达时间只能通过实验或另行证明评价。')
    p('完整37方案搜索用于检验候选覆盖能力，沿用最早独立到达时间下界进行精确剪枝[4]。'
      '剪枝只针对名义代价，不把名义最优当成扰动下的最优。学习提案、固定提案和随机提案均只比较最多两个候选，'
      '以便比较学习是否提供有价值的方案选择。')

    page('6 独立实验与评价协议')
    p('实验使用AMD Ryzen 5 9600X和64 GB内存的Windows工作站，Python 3.11.9、NumPy 2.4.6、Numba 0.65.1。'
      '控制和仿真采用单工作进程，编译及模型先预热。首次响应包含初始方案计算与第一个SUMO执行周期；不含编译冷启动。')
    p('开发阶段在24个已公开初态上运行480个回合，并检查此前200个初态的鲁棒可恢复性。随后冻结算法文件、模型哈希和六项验收标准，'
      '另抽常规、高密度各100个全新初态。每个初态运行四类误差和七种控制器，共5600个实际SUMO回合，公开记录见[5]。')
    table(['误差模式','子步误差规则','每种方法回合数'],[
        ['零误差','w=0',200],['随机误差','独立均匀于[−0.5,0.5]',200],['持续偏差','每车固定取±0.5',200],['间歇脉冲','1秒极值脉冲与1秒空隙',200]])
    p('同一初态及误差模式使用相同的外生误差带，方法次序按冻结随机种子打乱。误差带不传给控制器。'
      '对照为V8缓存计划、名义反馈、鲁棒传统基线、固定提案双方案、随机提案双方案、学习提案双方案和鲁棒完整搜索。'
      '固定提案索引1来自此前训练集，随机提案从36个非基线方案均匀抽取。')
    p('预先要求学习提案版本：无碰撞、合同违规或中止；完成率至少98%；三类非零扰动合并后相对鲁棒传统基线改善至少3%，'
      '且95%配对区间下界大于0；相对V8的故障回合减少至少80%；暖机首次响应P95不超过20 ms。')
    p('多个误差模式共享初态，不视为独立样本。合并评价时先对每个初态的三个非零模式配对差取平均，再按200个初态整簇bootstrap。'
      '中止按25 s计入代价，另报中止、超时、合同违规及碰撞；没有删除不利回合。区间为描述性区间，未作多比较显著性承诺。')

    page('7 完成情况与到达性能')
    names=[('v8_cached','V8缓存计划'),('nominal_feedback','名义反馈'),('robust_constant','鲁棒传统基线'),
           ('robust_fixed_pair','固定提案双方案'),('robust_random_pair','随机提案双方案'),('robust_learned_pair','学习提案双方案'),('robust_full','鲁棒完整搜索')]
    p('表1  全部800个条件中的完成情况与合同违规', 'Caption')
    table(['方法','完成回合','中止 / 合同违规回合'],[[name,r['overall'][m]['completed'],f"{r['overall'][m]['aborted']} / {r['overall'][m]['contract_violation_episodes']}"] for m,name in names])
    p('中止、合同违规和碰撞可能发生在同一回合，不应相加作为独立故障数。旧版在有界执行误差下已超出其精确模型保证条件；'
      '上述对照用于暴露适用范围差异，而非否定其原有条件性结论。')
    p('表2  平均截断到达代价  单位为秒', 'Caption')
    table(['方法','零误差','随机','持续','脉冲'],[[name,*[f"{r['summary_by_error'][mode][m]['mean_cost_s']:.3f}" for mode in ('clean','uniform','bias','pulse')]] for m,name in names],[5.2,2.8,2.8,2.8,2.8])
    p(f"学习提案版本记录碰撞{main['collision_episodes']}回合、超时{main['timeouts']}回合；"
      f"三类非零扰动下相对同证书传统基线平均节省{gain['mean_saved_s']:.3f} s，改善{gain['gain_percent']:.2f}%。"
      '这项比较不依赖将旧方法中止记为25秒所产生的代价差。')

    page('8 学习贡献与计算成本')
    p('表3  三类非零扰动中学习提案相对各对照的节省', 'Caption')
    comparisons=['robust_constant','robust_fixed_pair','robust_random_pair','robust_full']
    name_map=dict(names)
    table(['对照方法','平均节省 秒','95%整簇配对区间 秒'],[[name_map[m],f"{r['contrasts_perturbed']['all'][m]['mean_saved_s']:.3f}",
        f"[{r['contrasts_perturbed']['all'][m]['paired_cluster_ci95_s'][0]:.3f}, {r['contrasts_perturbed']['all'][m]['paired_cluster_ci95_s'][1]:.3f}]"] for m in comparisons])
    p('正数表示学习提案更快。固定、随机与学习提案使用同一证书和最多两个候选，差异因此更接近提案质量本身。'
      '完整搜索使用37个候选，在零误差模型中的最优性并不意味着扰动中必然更好；实际不利结果同样保留。')
    counts=gain['better_tied_worse_clusters']
    p(f"相对鲁棒传统基线，按每个初态的三类扰动平均差计，更好/持平/更差分别为{counts[0]}/{counts[1]}/{counts[2]}。"
      f"常规与高密度的平均改善分别为{r['contrasts_perturbed']['nominal']['robust_constant']['gain_percent']:.2f}%和"
      f"{r['contrasts_perturbed']['dense']['robust_constant']['gain_percent']:.2f}%。完整搜索平均更快，学习双方案的优势主要在计算开销。")
    p('表4  暖机首次响应P95  单位为毫秒', 'Caption')
    table(['方法','首次响应P95','候选数量上限'],[[name,f"{r['overall'][m]['warm_first_response_p95_ms']:.2f}",1 if m=='robust_constant' else 2 if m.endswith('_pair') else 37] for m,name in names])
    p(f"预先冻结的六项标准全部通过：{'是' if r['all_criteria_passed'] else '否'}。"
      '运行时间为单机实测分位数，不是最坏执行时间。仿真在控制计算期间暂停，因此不能据此宣称含计算延迟的闭环稳定性或实车实时性。')

    page('9 讨论与研究边界')
    p('本轮扩展了原型的模型适用范围：从精确动力学下的缓存执行，推进到已知加速度误差界下的反馈方案切换。'
      '区间证书的价值在于覆盖声明界内任意误差序列，而非仅记住测试过的扰动；通行收益则需由对照实验单独支持。')
    p('以下问题仍未解决：误差边界未知或被违反、感知误差、非合作车辆、自由换道、更多交通参与者、延迟以及实车验证。'
      '当前安全集可能因区间过度保守而缩小。名义双方案核验不提供扰动下逐场景性能不退化保证；'
      '此前统计置信门控未获得合格风险证书的结果也没有被本轮物理证书修复。')
    p('理论依据属于鲁棒可达分析与备用安全过滤的具体实现，不能把这些已有思想重新包装为首次提出。'
      '当前增量是带指令记忆的联合路线证书、反馈执行路径和同证书提案消融形成的可复现证据。'
      '距离投稿仍需与已发表同类方法在统一观测、动作和计算预算下对照，并扩大交互任务和物理误差来源。')
    h('参考文献')
    refs=[
        '[1] Wabersich K P, Zeilinger M N. A predictive safety filter for learning-based control of constrained nonlinear dynamical systems. Automatica, 2021, 129:109597. https://arxiv.org/abs/1812.05506',
        '[2] Didier A, Wabersich K P, Zeilinger M N. Adaptive Model Predictive Safety Certification for Learning-based Control. CDC, 2021. https://arxiv.org/abs/2109.13033',
        '[3] Banerjee A, Rahmani K, Biswas J, Dillig I. Dynamic Model Predictive Shielding for Provably Safe Reinforcement Learning. NeurIPS, 2024. https://arxiv.org/abs/2405.13863',
        '[4] CoopGate. V8 optimistic arrival lower bound and exact finite-portfolio pruning. Public implementation and derivation. https://github.com/Leo-7GOAT/CoopGate/blob/main/docs/theory_v8.md',
        '[5] CoopGate. Frozen V9 protocol, source hashes and raw SUMO episodes. https://github.com/Leo-7GOAT/CoopGate/tree/main/artifacts/v9']
    for text in refs:
        a=p(text);a.paragraph_format.first_line_indent=Cm(0);a.paragraph_format.space_after=Pt(4)
        for run in a.runs:run.font.size=Pt(9)
    doc.save(out)
    assert hashlib.sha256(reference.read_bytes()).hexdigest()==before_hash
    with ZipFile(reference) as old,ZipFile(out) as new:
        changed=[name for name in old.namelist() if old.read(name)!=new.read(name)]
        assert set(changed)<={'word/document.xml','docProps/core.xml','word/_rels/document.xml.rels','[Content_Types].xml'},changed
    print(json.dumps({'output':str(out),'reference_sha256':before_hash,'changed_parts':changed}))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--reference',type=Path,required=True);parser.add_argument('--out',type=Path,required=True)
    args=parser.parse_args();build(args.reference,args.out)
