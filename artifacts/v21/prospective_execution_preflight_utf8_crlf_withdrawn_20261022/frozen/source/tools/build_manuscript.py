"""Build the Chinese manuscript from measured pilot artifacts.

Run with a Python environment providing python-docx. No fabricated results or
hand-entered performance improvements are inserted into the manuscript.
"""
from pathlib import Path
import argparse
import json
from statistics import mean

from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
METHODS=['traditional','mappo','independent','mamps_style','joint_nominal','joint_robust']
NAMES={'traditional':'传统规控','mappo':'MAPPO','independent':'独立门控',
       'mamps_style':'单向迭代','joint_nominal':'联合名义','joint_robust':'联合多假设'}


def create(output):
    summaries={s:json.loads((ROOT/'artifacts'/f'seed{s}'/'summary.json').read_text()) for s in (7,19)}
    manifests={s:json.loads((ROOT/'artifacts'/f'seed{s}'/'manifest.json').read_text()) for s in (7,19)}
    doc=Document()
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=Cm(2.25)
    sec.left_margin=sec.right_margin=Cm(2.35)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Caption'):
        st=doc.styles[name]
        st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
        st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    doc.styles['Normal'].font.size=Pt(11)
    doc.styles['Normal'].paragraph_format.line_spacing=1.25
    doc.styles['Normal'].paragraph_format.space_after=Pt(6)
    doc.styles['Title'].font.size=Pt(19)
    doc.styles['Title'].paragraph_format.space_after=Pt(12)
    doc.styles['Subtitle'].font.italic=False
    doc.styles['Heading 1'].font.size=Pt(14)
    doc.styles['Heading 2'].font.size=Pt(12)
    doc.styles['Caption'].font.size=Pt(9)
    for name in ('Heading 1','Heading 2'):
        doc.styles[name].paragraph_format.keep_with_next=True
    doc.core_properties.title='基于SUMO的多智能体强化学习与传统规控切换研究'
    doc.core_properties.author='李东旭'

    def p(text,style=None):
        para=doc.add_paragraph(text,style)
        if style is None:
            para.paragraph_format.first_line_indent=Cm(.74)
        if style == 'Caption' and text.startswith('表'):
            para.paragraph_format.keep_with_next=True
        return para

    def h(text):
        doc.add_heading(text,1)

    def equation(text):
        para=doc.add_paragraph()
        para.alignment=WD_ALIGN_PARAGRAPH.CENTER
        math=OxmlElement('m:oMath')
        run=OxmlElement('m:r');node=OxmlElement('m:t');node.text=text
        run.append(node);math.append(run);para._p.append(math)

    def table(headers,rows,widths):
        t=doc.add_table(rows=1, cols=len(headers))
        t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
        for cell,w,label in zip(t.rows[0].cells,widths,headers):
            cell.width=Cm(w);cell.text=label
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            cells=t.add_row().cells
            for c,w,text in zip(cells,widths,row):
                c.width=Cm(w);c.text=str(text)
        for rid,row in enumerate(t.rows):
            no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
            for cell in row.cells:
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcPr=cell._tc.get_or_add_tcPr()
                borders=OxmlElement('w:tcBorders')
                for side in ('top','left','bottom','right'):
                    b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
                tcPr.append(borders)
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E2E8F0' if rid==0 else ('FFFFFF' if rid%2 else 'F8FAFC'));tcPr.append(shade)
                margins=OxmlElement('w:tcMar')
                for side in ('top','left','bottom','right'):
                    m=OxmlElement('w:'+side);m.set(qn('w:w'),'80');m.set(qn('w:type'),'dxa');margins.append(m)
                tcPr.append(margins)
                for para in cell.paragraphs:
                    para.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    para.paragraph_format.space_after=Pt(2)
                    para.paragraph_format.line_spacing=1.1
                    for run in para.runs:
                        run.font.size=Pt(9);run.bold=(rid==0)
        p('')
        return t

    p('基于SUMO的多智能体强化学习与传统规控切换研究','Title')
    p('李东旭','Subtitle')
    p('摘要','Heading 1')
    joint=[summaries[s]['joint_robust'] for s in (7,19)]
    base=[summaries[s]['traditional'] for s in (7,19)]
    p('针对多车协同决策中学习策略与传统规控的切换问题，本文建立基于Eclipse SUMO的可复现实验平台。'
      '平台采用集中训练与分散策略执行的MAPPO，配合规则回退规划和纵向模型预测控制器，比较独立门控、'
      '单向迭代门控、联合模式枚举及有限背景行为假设下的联合门控。'
      '实验包含三辆受控车、两辆背景车及正常交通、较密交通和前车制动三类条件。'
      '两次独立训练各执行50000个决策步，冻结策略后使用新场景种子进行六种方法的配对评估。'
      f'联合多假设门控在两个策略种子下的任务成功率分别为{joint[0]["success_rate"]:.1%}和{joint[1]["success_rate"]:.1%}，'
      f'经验碰撞率分别为{joint[0]["collision_rate"]:.1%}和{joint[1]["collision_rate"]:.1%}。'
      '结果用于分析风险、通行效率和计算成本的权衡，不构成普适安全保证或优于已有方法的充分证据。'
      '本文明确区分预测裕度分数与校准安全概率，并讨论终端可恢复性、模型误差和有限实验规模带来的限制。')
    p('关键词  多智能体强化学习  模型预测控制  策略切换  安全门控  交通仿真')

    h('1 引言')
    p('强化学习能够通过交互学习多车行为，但部署阶段的安全性受到训练覆盖、预测误差和控制约束影响。'
      '传统规划控制可以提供结构明确的候选动作，却不能从任意状态恢复安全。将两类方法组合时，关键在于判断'
      '何时允许学习策略执行，以及切换之后其他车辆依赖的协作是否仍然成立。')
    p('例如，汇入车辆根据主路车辆让行的预测开始并线。如果主路车辆随后改用另一种规划方式，原来的让行预测可能失效。'
      '只对每辆车的原始学习动作独立评估，再分别替换部分动作，无法自动确保替换后的联合动作仍满足原来的预测约束。')
    p('这一问题已有明确先例。MAMPS已研究多智能体回退策略的连锁影响，不能把联合切换本身作为本文首次提出的贡献[1]。'
      '本文的工作定位为可复现的交通场景比较研究：在同一冻结策略、相同车辆约束和相同测试种子下，分析预测假设和模式选择方式'
      '对系统表现的影响，并记录未获得优势的情形。')

    h('2 相关研究与本文定位')
    p('MAMPS在确定性动力学及可恢复性假设下，迭代选择需要回退的智能体，并给出相应安全论证[1]。'
      '集中式和因子化屏障也被用于多智能体安全学习[2]。因此，本项目采用的联合状态检查、回退控制和模式选择具有已有理论基础。')
    p('MAPPO为合作任务提供了较强的策略梯度基线，本文使用共享Actor与集中Critic实现一个紧凑版本[3]。'
      '面向网联自动驾驶的研究已将多智能体学习、共享信息和控制障碍函数安全检查结合[4]。'
      '基于共形方法的策略学习则研究了利用不确定性信息进行策略切换[5]。这些结果说明，仅将MAPPO、MPC和置信度阈值组合并不足以确认新颖性。')
    p('相对上述基础，本文当前实施了两项工程扩展。第一，在不可直接控制的背景车存在时，对有限的加速度假设分别检查候选动作。'
      '第二，在三辆受控车形成的八种模式组合中进行显式搜索，并与单向迭代方式比较。前者不等于完整随机安全控制，后者也不具备一般规模下的高效性。'
      '两项扩展是否形成可投稿的算法贡献，需要更充分的查新、强基线与等预算验证。')

    h('3 问题描述')
    p('考虑N等于3辆受控车辆及两辆背景车辆。受控车状态包含纵向位置、横向位置、速度和加速度。'
      '仿真器负责交通演化和碰撞检测；受控车接收纵向加速度与目标车道指令。背景车辆保留SUMO行为模型，门控器不读取其未来真实动作。')
    equation('x(t+1) = F(x(t), u(t), w(t))')
    p('其中，w表示背景行为和预测模型未描述的影响。Actor输入为自车状态及邻车相对状态；联合监督器使用共享组状态。'
      '训练时的集中Critic不参与部署，但联合监督器依赖集中信息，因此整体系统不能描述为完全分布式。')
    equation('b ∈ {0,1}³,    uᵢ(b) = bᵢuᵢᴿᴸ + (1−bᵢ)uᵢᴮ')
    p('b表示各车使用学习候选或传统候选的模式。这里的二元表达表示候选选择，不是将两个连续控制量按任意比例混合。'
      '目标是在约束检查通过的条件下保留有效的学习决策，同时减少切换开销。')

    h('4 方法')
    doc.add_heading('4.1 多智能体决策与纵向控制',2)
    p('Actor和Critic均采用两层64维隐藏层及双曲正切激活。Actor参数在三辆车之间共享，Critic使用联合观测计算团队价值。'
      '每辆车有九个离散动作，由三档目标速度和保持、向左、向右三种车道意图组合。越界车道意图裁剪到可用车道。'
      '使用团队回报构造广义优势估计，并采用PPO裁剪目标更新策略；时间上限截断与真正终止采用不同的自举处理。')
    equation('J = Σₖ [(vₖ−vref)² + 0.1aₖ² + 0.3(aₖ−aₖ₋₁)²]')
    p('纵向控制器用OSQP求解有限时域二次规划，约束速度、加速度和加速度变化率。预测步长0.5秒，预测步数12。'
      '正常模式最大加速度2.5米每二次方秒、最大减速度5米每二次方秒、最大加加速度5米每三次方秒。'
      '紧急制动允许突破舒适性变化率约束，并单独计数。横向运动由SUMO的连续换道模型执行；本文未实现完整横纵向MPC。')
    doc.add_heading('4.2 传统候选与预测检查',2)
    p('传统规划根据前车间距调节目标速度，汇入车辆仅在主路间隙满足规则条件时请求换道，否则减速。'
      '该策略是具有明确结构的回退候选，不被假设为全状态空间上的安全控制器。')
    p('门控器预测一次候选决策之后转入规则回退的轨迹。检查内容包括车辆间几何间距和合流区域约束，'
      '并加入随预测时间增长的位置裕度。当前预测器近似纵向闭环响应和换道持续过程，没有与SUMO动力学严格等价。')
    equation('M(b) = min over ω,k,i  clearanceᵢ(x̂ₖ(b,ω))')
    equation('C(b) = 1 / [1 + exp(−M(b)/2)]')
    p('C是裕度的单调变换，仅用于阈值门控。它没有经过数据校准，不能解释为动作安全概率。'
      '当前检查步数为10，对应5秒。联合多假设版本分别使用背景车减速、匀速及加速假设，并取最不利裕度。'
      '有限假设集无法覆盖所有驾驶行为，且当前没有定义和证明不变终端安全集合。')
    doc.add_heading('4.3 联合模式选择',2)
    equation('maximize over b   Σᵢbᵢ − λ Σᵢ|bᵢ−bᵢprev| + ε Σᵢvᵢref')
    p('对三辆车枚举八种模式组合，在通过检查的候选中最大化上述目标。λ为0.08，ε为0.001。'
      '保持学习模式的分数阈值为0.8，重新进入学习模式的阈值为0.9。全回退候选仍须通过非负几何裕度检查。'
      '如果没有候选通过，则执行并记录紧急制动；这只是一种最低风险尝试，不记作获得了安全证书。')
    p('单向迭代基线沿用MAMPS的模式收缩思想，但使用本文相同的近似有限时域检查器和阈值。'
      '由于缺少原方法要求的终端条件，该基线称为MAMPS风格适配，不继承原论文的安全定理。'
      '联合枚举的检查次数更大，因此性能变化不能直接解释为相同计算预算下的优势。')

    h('5 实验设置')
    p('实验在真实Eclipse SUMO中运行，使用TraCI进行闭环控制。道路为双车道收敛到单车道的短合流场景，'
      '不是完整复杂匝道地图。仿真步长0.1秒，决策周期0.5秒，每回合最长30秒，三辆车均到达目标位置才记为任务成功。'
      '任一受控车碰撞则记回合碰撞。达到时间上限而未完成任务不会记作成功。')
    p('为了避免把仿真器隐式避撞效果归功于门控，受控车设置speedMode等于96、laneChangeMode等于0。'
      '背景车辆仍使用SUMO的跟驰与换道模型。所有方法共享相同输入信息、控制约束和MPC。'
      '实验没有引入通信延迟指标，也没有声称存在视觉感知或真实车辆验证。')
    p('开发阶段首先进行12000步短训练，观察到明显的策略种子差异。随后固定训练预算为50000步，'
      '以种子7和19分别从头训练，使用训练结束时的模型，未根据最终测试结果挑选检查点。'
      '最终测试种子从20000开始，三类场景使用不同的种子段，每类5回合；各方法使用相同场景种子。'
      '训练中出现过同类场景，所以这里验证新种子下的表现，不属于严格的分布外测试。')
    p('软件版本为SUMO '+manifests[7]['sumo']+'及PyTorch '+manifests[7]['torch']+'。'
      '背景车加速度预测假设为−2.5、0和1.5 m/s²。每种方法在每个策略种子下测试15回合，'
      '总计180次执行，共用15个固定场景配置。训练和测试采用CPU执行。')

    h('6 初步实验结果')
    p('表1逐个报告训练种子下的结果。相同策略的各方法共享测试场景，两个策略种子也复用相同场景集合，'
      '因此不能把全部车辆、时间帧或重复背景场景当作独立样本。仓库同时保留每回合原始CSV和分组结果。')
    p('表1 六种方法在两个策略种子下的实测结果','Caption')
    rows=[]
    for method in METHODS:
        for seed in (7,19):
            r=summaries[seed][method];n=r['episodes']
            rows.append([NAMES[method],seed,f'{round(r["collision_rate"]*n)}/{n}',f'{round(r["success_rate"]*n)}/{n}',
                         f'{r["rl_fraction"]*100:.1f}',f'{r["gate_mean_ms"]:.2f}'])
    table(['方法','策略种子','碰撞','成功','RL比例 %','门控均值 ms'],rows,[3.3,1.7,2.0,2.0,3.2,3.8])
    p(f'传统规控在种子7和19的成功率分别为{base[0]["success_rate"]:.1%}与{base[1]["success_rate"]:.1%}，'
      f'成功回合平均时间均为{base[0]["successful_time_s"]:.2f}秒。传统方法不依赖策略种子，两行是相同场景的重复运行，不能当作30个独立样本。'
      '这是必须保留的强基线；如果混合方法没有同时改善安全和效率，就不能声称系统整体优于传统规控。')
    for seed in (7,19):
        r=summaries[seed]
        p(f'策略种子{seed}下，MAPPO、独立门控、单向迭代、联合名义、联合多假设的成功率依次为'
          + '、'.join(f'{r[m]["success_rate"]:.1%}' for m in METHODS[1:])+'。'
          +f'其中联合多假设的平均RL使用比例为{r["joint_robust"]["rl_fraction"]:.1%}，'
          +f'紧急制动步数均值为{r["joint_robust"]["emergency_steps"]:.2f}。这些指标需要联合解读，'
          +'不能仅根据较低碰撞率推断方法更优。')
    image=ROOT/'artifacts'/'seed19'/'comparison.png'
    if image.exists():
        doc.add_picture(str(image),width=Cm(16))
        doc.paragraphs[-1].paragraph_format.keep_with_next=True
        p('图1 策略种子19下的碰撞率 任务成功率与RL使用比例','Caption')
    p('图中区间为回合级Wilson描述性区间。每组只有15回合，零碰撞仍对应非零风险上界；'
      '不同场景的分层混合和仅有两个训练种子进一步限制了统计解释。当前结果不足以建立显著优于前人的结论。'
      '门控耗时也受到本机同时运行实验的负载影响，不能作为部署最坏执行时间保证。')
    p('在策略种子19下，独立门控取得0次碰撞和10次成功，联合多假设门控则出现2次碰撞和6次成功。'
      '这直接否定了“增加预测假设和联合搜索就必然更安全”的简单推断。当前检查器与执行器存在模型差异，'
      '且紧急回退不具备已证明的可恢复性，因此较保守的预测检查也可能改变交互过程并进入新的失败状态。'
      '这些因素是需要进一步排查的机制，不能仅由汇总结果断定具体碰撞原因。')
    p('策略种子7下，纯MAPPO与多种门控均没有碰撞，但全部未在30秒内完成三车通行。'
      '这说明安全裕度本身不足以度量学习决策是否值得执行：速度过低或长期等待也可能取得较高安全分数。'
      '因此，后续门控需同时检验可恢复性与任务进展，而不能仅提高RL使用比例。')

    h('7 讨论与局限')
    p('第一，传统回退的质量决定门控能够恢复到何种状态。过度保守的检查可能牺牲完成率，而过于简单的预测器可能错误放行。'
      '当前安全分数由裕度直接变换，尚未验证条件概率校准，也未对自适应模式选择后的整回合风险给出界。')
    p('第二，本文采用有限时域检查，并未证明终端集合的不变性，不能从有限预测轨迹无碰撞推出无限时域安全。'
      '背景行为假设只覆盖三个加速度值，并非对任意人类驾驶行为的鲁棒包络。')
    p('第三，当前MPC仅控制纵向，车辆横向由SUMO执行；没有完成CARLA闭环动力学验证、实车实验、'
      '观测噪声消融、完整驾驶风格迁移实验或等计算预算比较。这些内容不能写成已经完成的结果。')
    p('后续应先复现具有明确终端条件的强安全过滤基线，选择风险校准或固定预算搜索中的一个问题深入。'
      '如果更充分实验仍显示传统规控更优，应保留这一负面结果并调整研究问题，而不是通过弱化基线或选择有利种子形成优势。')

    h('8 结论')
    p('本文实现了基于SUMO的多车学习策略与传统规控切换平台，提供可复现的训练、六类方法对照、原始数据与轨迹回放。'
      '相对已有工作的当前差异在于混合交通下有限背景行为假设的测试，以及同一检查器下不同模式选择策略的比较。'
      '这些差异目前属于工程扩展和待验证研究假设。要形成有竞争力的投稿，仍需建立更明确的方法贡献、充分的对照证据及对应的理论或统计依据。')

    h('参考文献')
    refs=[
      '[1] Zhang W, Bastani O, Kumar V. MAMPS Safe Multi-Agent Reinforcement Learning via Model Predictive Shielding. arXiv:1910.12639, 2019. https://arxiv.org/abs/1910.12639',
      '[2] Elsayed-Aly I, Bharadwaj S, Amato C, et al. Safe Multi-Agent Reinforcement Learning via Shielding. AAMAS, 2021. https://arxiv.org/abs/2101.11196',
      '[3] Yu C et al. The Surprising Effectiveness of PPO in Cooperative Multi-Agent Games. NeurIPS, 2022. https://arxiv.org/abs/2103.01955',
      '[4] Zhang Z, Han S, Wang J, Miao F. Spatial-Temporal-Aware Safe Multi-Agent Reinforcement Learning of Connected Autonomous Vehicles in Challenging Scenarios. ICRA, 2023. https://arxiv.org/abs/2210.02300',
      '[5] Huang H et al. Conformal Policy Learning for Sensorimotor Control Under Distribution Shifts. arXiv:2311.01457, 2023. https://arxiv.org/abs/2311.01457',
      '[6] Eclipse SUMO. Simulation of Urban MObility Documentation. https://eclipse.dev/sumo/docs/'
    ]
    for ref in refs:
        para=p(ref)
        para.paragraph_format.first_line_indent=Cm(0)
        for run in para.runs:run.font.size=Pt(9)
    footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    output=Path(output);output.parent.mkdir(parents=True,exist_ok=True)
    doc.save(output)
    print(output)


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--out',required=True)
    args=parser.parse_args()
    create(args.out)
