"""Build the paper's completed OOD appendix using the retained manuscript style."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
import shutil

from docx import Document
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from summarize_completed_ood_20261027 import read_complete, ROOT, SOURCE, digest

OUT=ROOT/'runs/ood_paper_appendix_20261027/revision1'
NAME='CoopGate_扩展场景完整结果_20261027.docx'
REFERENCE=ROOT/'runs/manuscript_v17/candidate_1.docx'


def build():
    if OUT.exists():raise FileExistsError('Preserve the previous document revision')
    report=read_complete()
    assert digest(REFERENCE)=='1f8e040594d25ce5b61fd2e2a41794ef191b6d7952698114117b0a3e23f3f795'
    receipt=json.loads((ROOT/'artifacts/recovery_20261027/ood_finalization_receipt.json').read_text(encoding='utf-8'))
    figure=ROOT/'artifacts/recovery_20261027/ood_scope_summary/failure_scope.png'
    assert digest(figure)==receipt['figure_sha256']
    failures=json.loads((ROOT/'artifacts/recovery_20261027/ood_failure_types/summary.json').read_text(encoding='utf-8'))
    assert digest(SOURCE/'analysis/report.json')==receipt['report_sha256']
    OUT.mkdir(parents=True)
    shutil.copy2(REFERENCE,OUT/'reference.docx')
    (OUT/'artifact.md').write_text('''# Completed OOD paper appendix

Retained reference.docx is the style authority, unchanged from the manuscript V17 snapshot.
Preserve its paragraph styles, fonts and portrait page geometry for the text page.
Use one landscape section for the full nine-factor failure figure so labels remain legible.
This landscape figure page is an intentional deviation from the reference geometry.
Lead with the failed all-model decision; retain every failure and the distinction between composite failures and collisions.
Expected output: one DOCX, two pages; render and inspect every page before delivery.
''',encoding='utf-8')
    doc=Document(REFERENCE)
    for child in list(doc.element.body):
        if child.tag!=qn('w:sectPr'):doc.element.body.remove(child)
    for name in ('Title','Heading 1','Heading 2'):
        doc.styles[name].font.color.rgb=RGBColor(0,0,0)
    normal=doc.styles['Normal'];normal.font.size=Pt(10.5)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'),'宋体')
    normal.paragraph_format.space_after=Pt(5)
    normal.paragraph_format.line_spacing=1.12
    doc.core_properties.title='CoopGate 复合失配实验完整结果'
    doc.core_properties.author='李东旭'
    doc.core_properties.subject='论文实验结果附录'
    doc.core_properties.version='20261027'

    def paragraph(text,style=None):
        p=doc.add_paragraph(text,style)
        p.paragraph_format.first_line_indent=Inches(0)
        return p

    def table(headers,rows):
        t=doc.add_table(rows=1,cols=len(headers));t.style='Table Grid'
        for cell,value in zip(t.rows[0].cells,headers):cell.text=value
        tr=t.rows[0]._tr.get_or_add_trPr();repeat=OxmlElement('w:tblHeader');tr.append(repeat)
        for row in rows:
            for cell,value in zip(t.add_row().cells,row):cell.text=str(value)
        for ri,row in enumerate(t.rows):
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(2)
                    p.paragraph_format.line_spacing=1
                    for run in p.runs:run.font.size=Pt(9);run.bold=(ri==0)
        return t

    paragraph('CoopGate 复合失配实验完整结果','Title')
    paragraph('论文实验附录  2026 年 10 月 27 日')
    paragraph('完整复核未满足三个固定模型全部通过的预设判据。模型 1812 和 1822 通过相对比较，1802 的失败惩罚代价收益未达显著性要求。基线和门控在这组超出假设范围的压力测试中均有大量失败，当前结果不能支持可安全部署的结论。')
    paragraph('样本与原定判定','Heading 1')
    paragraph('原定 22,464 次 SUMO 运行及完整物理与门控回放全部结束。独立单位为 192 个初始条件，每个初始条件重复九个动力学与感知组合、十三种方法；运行总数不能作为独立样本量。')
    rows=[]
    for model in ('1802','1812','1822'):
        m=report['primary']['models'][model]
        rows.append([model,f"{100*m['failure_risk_difference']:.4f}",
            f"{m['failure_holm_adjusted_p']:.6f}",f"{m['failure_penalized_cost_gain_seconds']:.5f}",
            f"{m['loss_holm_adjusted_p']:.6f}",'通过' if m['frozen_thresholds_pass'] else '未通过'])
    table(['模型','失败率差\n百分点','失败率\nHolm p','惩罚代价\n增益 秒','增益\nHolm p','判定'],rows)
    paragraph('失败率差为门控减基线。原非劣性界限为 1 个百分点，惩罚代价增益需严格为正并通过检验。表中沿用冻结的分层配对初始条件 bootstrap 和 Holm 判定，没有新增检验；这些比较不构成有限样本总体安全证书。')
    paragraph('失败类型与工程边界','Heading 1')
    selected=[next(x for x in failures['methods'] if (x['rule'],x['model'])==k)
              for k in [('baseline','none'),*[('bank_performance',m) for m in ('1802','1812','1822')]]]
    table(['方法','复合失败\n共 1,728 次','软件中止','记录碰撞','有几何违规\n的运行'],
          [[('基线' if x['model']=='none' else 'bank '+x['model']),x['counts']['failure'],
            x['counts']['aborted'],x['counts']['collision'],x['counts']['episodes_with_route_violation']]
           for x in selected])
    paragraph('复合失败率约为 55%，不是碰撞率；以上标志可以重叠，不能相加。中止主要对应区间或执行器域检查失败，以及当前状态缺少执行器感知停车证书。软件中止不等于已安全停车，缺少当前制动策略族的证书也不证明碰撞不可避免。')
    paragraph('所有复合动力学实现均超出原物理证书的声明类别，带观测误差的条件还违反精确状态假设。混合总体的相对判据不能推广成每个因素格均更安全；失败惩罚代价改善也不能直接写成无失败任务的通行效率改善。')
    p=paragraph('复现资料  docs/ood_complete_results_20261027.md 及其完整报告和原始轨迹清单。联合风险条件证书与公平对照仍待完成。')
    for run in p.runs:run.font.size=Pt(9)

    section=doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.orientation=WD_ORIENT.LANDSCAPE
    section.page_width,section.page_height=Inches(11.69),Inches(8.27)
    section.left_margin=section.right_margin=Inches(.6)
    section.top_margin=section.bottom_margin=Inches(.5)
    p=paragraph('复合失配条件下的经验失败率','Heading 1')
    p.paragraph_format.space_after=Pt(3)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.add_run().add_picture(str(figure),width=Inches(8.8))
    p=paragraph('图 1  各格均保留失败，使用同样的 192 个初始条件。逐格数据为描述性结果，没有另外检验显著性。mps_shared 与 mamps_shared 为本项目接口适配版本，不能据此声称超越原作者完整方法。')
    for run in p.runs:run.font.size=Pt(9)
    path=OUT/NAME;doc.save(path)
    meta=dict(path=str(path),sha256=digest(path),source_report_sha256=receipt['report_sha256'],
              figure_sha256=digest(figure),original_reference_sha256=digest(REFERENCE),
              actual_sumo_runs=22464,independent_initials=192,
              all_three_fixed_models_pass=False,new_safety_certificate=False)
    (OUT/'build.json').write_text(json.dumps(meta,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(meta,ensure_ascii=False),flush=True)


if __name__=='__main__':build()
