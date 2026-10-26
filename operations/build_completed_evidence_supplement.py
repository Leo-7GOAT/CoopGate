"""Build a paper supplement from completed evidence; never read live outcomes."""
from copy import deepcopy
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT/'runs/manuscript_v17/candidate_1.docx'
REFERENCE_SHA = '1f8e040594d25ce5b61fd2e2a41794ef191b6d7952698114117b0a3e23f3f795'
OUT = ROOT/'runs/completed_evidence_supplement_20261026/revision2'
NAME = 'CoopGate_论文实验与理论补充_20261026.docx'
REPORT = 'artifacts/v38/interrupted_block_grid_recovery/v36/analysis/report.json'
REPORT_SHA = '1e13918534046bc71b3a39a0ee340c5852357778930dc5cbc8b96cf4eb175162'


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def build():
    if OUT.exists():
        raise FileExistsError('Retain every manuscript iteration')
    assert sha(REFERENCE) == REFERENCE_SHA
    assert sha(ROOT/REPORT) == REPORT_SHA
    interpretation = ROOT/'artifacts/v38/interrupted_block_grid_recovery/v36/analysis/interpretation.json'
    result = json.loads(interpretation.read_text(encoding='utf-8'))
    assert result['source_report_sha256'] == REPORT_SHA
    assert result['actual_sumo_runs'] == 8320 and result['reused_independent_initials'] == 160
    assert result['new_independent_initials'] == 0 and not result['new_risk_certificate']
    aggregate = result['aggregate']
    assert aggregate['bank_performance']['harmful_and_adopted'] == 5
    assert aggregate['nominal_performance']['harmful_and_adopted'] == 145
    assert aggregate['bank_performance']['adopted'] == 1150
    assert aggregate['nominal_performance']['adopted'] == 1728
    doc = Document(REFERENCE)
    section_xml = doc.sections[0]._sectPr.xml
    for child in list(doc.element.body):
        if child.tag != qn('w:sectPr'):
            doc.element.body.remove(child)
    for name in ('Title', 'Heading 1', 'Heading 2'):
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    doc.core_properties.title = 'CoopGate 已完成实验与条件性能保证'
    doc.core_properties.subject = '论文实验及理论补充'
    doc.core_properties.version = '20261026'
    doc.core_properties.author = '李东旭'

    def p(text, style=None):
        paragraph = doc.add_paragraph(text, style)
        paragraph.paragraph_format.first_line_indent = Cm(0 if style else .74)
        if not style:
            paragraph.paragraph_format.space_after = Pt(5)
        return paragraph

    def h(text):
        return p(text, 'Heading 1')

    def eq(text):
        paragraph = p('')
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_after = Pt(7)
        math = OxmlElement('m:oMath')
        run = OxmlElement('m:r')
        value = OxmlElement('m:t')
        value.text = text
        run.append(value)
        math.append(run)
        paragraph._p.append(math)

    def table(headers, rows, widths):
        t = doc.add_table(rows=1, cols=len(headers))
        t.autofit = False
        for column, width in zip(t.columns, widths):
            column.width = Cm(width)
        for cell, text in zip(t.rows[0].cells, headers):
            cell.text = text
        for row in rows:
            for cell, text in zip(t.add_row().cells, row):
                cell.text = str(text)
        borders = OxmlElement('w:tblBorders')
        for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            node = OxmlElement('w:'+edge)
            for key, value in {'val':'single', 'sz':'4', 'color':'D9D9D9'}.items():
                node.set(qn('w:'+key), value)
            borders.append(node)
        t._tbl.tblPr.append(borders)
        for index, row in enumerate(t.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            if index == 0:
                row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            for column, (cell, width) in enumerate(zip(row.cells, widths)):
                cell.width = Cm(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                props = cell._tc.get_or_add_tcPr()
                margins = OxmlElement('w:tcMar')
                for edge in ('top', 'left', 'bottom', 'right'):
                    n = OxmlElement('w:'+edge)
                    n.set(qn('w:w'), '80')
                    n.set(qn('w:type'), 'dxa')
                    margins.append(n)
                props.append(margins)
                if index == 0:
                    fill = OxmlElement('w:shd'); fill.set(qn('w:fill'), 'DCE6F1'); props.append(fill)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if column == 0 else WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.first_line_indent = Cm(0)
                    paragraph.paragraph_format.space_after = Pt(3)
                    paragraph.paragraph_format.space_before = Pt(3)
                    paragraph.paragraph_format.line_spacing = 1.05
                    for run in paragraph.runs:
                        run.font.size = Pt(10)
                        run.bold = index == 0
        p('')

    p('CoopGate 已完成实验与条件性能保证', 'Title')
    p('论文实验及理论补充  2026 年 10 月 26 日', 'Subtitle')
    h('1 完整闭环结果及研究结论')
    p('学习决策相对传统控制器的实用价值，需要同时考察收益、实际采用和采用后的逐车损害。我们在固定反馈恢复基线上，对完整联合动作进行物理过滤和有限扰动性能评估。已完成的全网格实验表明，扰动集合门控保留正团队收益并减少退化，但有限路径评估仍会误放有害动作，不能直接等同于真实扰动下的性能保证。')
    p('全部 8,320 次真实 SUMO 执行与独立完整重放覆盖五类场景各 32 个旧初态、四类动力学及三个固定模型，共 640 个封存块。独立初态为 160 个；重复动力学、模型和恢复执行不增加独立样本量。下表是完整开发范围的描述性结果，不作新的显著性或总体风险认证。')
    rows = []
    for label, key in [('共同反馈基线','baseline'),('标称性能门控','nominal_performance'),('扰动集合性能门控','bank_performance')]:
        item = aggregate[key]
        rows.append([label, f"{item['completed']}/{item['runs']}", str(item['adopted']), str(item['harmful_and_adopted']), f"{item['all_scope_cost_gain_percent']:.4f}%"])
    table(['控制器','完成／运行','采用回合','采用且退化','成本改善'], rows, [4.5,3.1,2.7,2.9,3.2])
    p('集合门控的回合采用率为 59.8958%，采用内经验损害比例为 5/1,150，即 0.4348%；标称门控为 145/1,728，即 8.3912%。两者采用集合不同，比例差异不能解释为同一人群上的随机处理效果。相同完整分母下的联合事件频率分别为 5/1,920 与 145/1,920。')
    p('成本改善按配对范围的基线成本总和计算，不是逐回合百分比的平均值。集合门控在三个固定模型上的改善分别为 1.7671%、1.6852% 和 1.7182%；所有模型均保留，未选择最佳模型报告。该结果没有证明原始学习策略普遍优于基线，也不替代真实 MARL 训练及协作机制的专项证据。')

    doc.add_page_break()
    h('2 性能保证的数学条件与保留反例')
    p('令 z 包含使闭环具有 Markov 性的全部状态：车辆状态、执行器记忆与命令队列、上一命令、到达标记、时钟及控制器所需历史。每车成本 J 为登记完成时间 T 加 0.05 倍累计命令变化量 D。D 是控制命令变化量，不能当作实车 jerk。固定实际外生序列 ω，并以同一反馈基线 b 定义有限时域剩余值函数。')
    eq('Vᵇₜ,ω(z) = c(z,bₜ(z)) + Vᵇₜ₊₁,ω(fₜ,ω(z,bₜ(z)))')
    eq('δₜ,ω(z,a) = c(z,a) + Vᵇₜ₊₁,ω(fₜ,ω(z,a)) − Vᵇₜ,ω(z)')
    p('命题 1。若实际转移与评估模型完全一致，实际 ω 属于验证集合，剩余值函数与成本计算精确，且每个执行的完整联合动作均对全部车辆和指标满足 δ≤0，那么门控闭环相对同初态、同 ω 下的独立反馈基线不退化。物理约束仍需另由其有效假设下的过滤器保证。')
    eq('Jᵍω,i − Jᵇω,i = Σₜ₌₀ᴴ⁻¹ δₜ,ω,i(zₜ,aₜ) ≤ 0')
    p('证明。将 δ 的定义沿实际门控轨迹求和，连续两步的基线值函数相消；终端值为零，初始值等于对应反馈基线的总成本。回退动作使用同一 b，因此其 δ 为零。证明必须评价完整联合动作；逐车单独替换可行不能推出组合后仍可行。这是标准策略改进恒等式的条件应用，不作为新定理。若每步只保证真实 δ≤ε，则只能得到 Hε；数值评估误差还须另计。')
    p('已完成的 16 次原生分支重放覆盖五个集合门控退化案例的八个采用点，并准确重构原回合差值。这些案例来自四个旧初态，属于事后机制解释，不产生新的独立样本或发生率估计。真实参数及记忆只用于离线 oracle，不进入部署控制器。')
    p('即使名义模型精确，有限路径仍有反例。低速让行初态 18、模型 1822 中，诊断插值越线时间由 9.984358 s 变为 10.007796 s，延迟仅 0.0234375 s，却使原 0.5 s 网格完成时间由 10.0 s 变为 10.5 s。命令变化量减少 2.5，因此原成本增加 0.5−0.05×2.5=0.375。相位误差小于 0.05 s 不能推出原成本不退化；不能改变主指标或放宽容差删除该事件。')

    doc.add_page_break()
    h('3 风险与收益联合判定的统计逻辑')
    p('在线通过有限扰动 bank 表示候选通过指定预测检查，不是当前状态的概率置信度。独立统计命题针对冻结完整控制器在预声明分布上的闭环行为。A 表示回合至少实际采用一次有效学习命令；H 表示任一车辆成本或登记完成时间增加，或执行失败。主要目标是 r=P(H=1 | A=1)，并同时限制采用率和团队收益，防止用几乎始终回退来取得表面低风险。')
    eq('r = P(H∩A)/P(A) = q/a， a > 0')
    p('平衡校准含四类执行器各 768 个新独立初态；固定分层计数不能把所有采用样本直接视为同一二项样本。对层 s 的 qₛ=P(H∩A | s) 和 aₛ=P(A | s)，分别构造上界 Uq,ₛ 与下界 La,ₛ。在该项风险的 α 内，对四层的八个单侧界各分配 α/8；已知混合权重 wₛ=1/4 时，有：')
    eq('r ≤ min{1, (Σₛ wₛUq,ₛ)/(Σₛ wₛLa,ₛ)}')
    p('上述有限样本结论要求层内独立同分布、完整冻结控制器、事先固定权重及正的分母下界。证明由八个有效覆盖事件的并集界与比值单调性得到。分母下界为零时不可认证。强滞后子分布另用层内采用条件下的精确二项检验，混合风险合格不能自动推出每个场景或任意分布外风险合格。')
    p('三个固定模型均须同时满足混合及强滞后风险不超过 5%、确认分布采用率下界至少 20%、配对团队百分比改善下界大于 0.5%，并保留零观测执行失败这一工程验收条件。确认集为 1,024 个另行 IID 混合初态；每行先计算配对百分比改善，再作均值检验，其指标不同于第 1 节的合计成本改善。')
    eq('Gₖ = 100(Jᵇteam,k − Jᵍteam,k)/Jᵇteam,k')
    p('固定全部分项同时通过是交并检验：只要某项真实不合格，整个方案错误通过必然包含该不合格项错误通过。因此可用分项 p 值的最大值检验整体命题，不要求各模型独立；这不产生所有分项区间的联合覆盖。第二次尝试沿用 α=0.05/6，不重用第一轮预算。')
    p('采用与风险项使用有限样本二项界；收益项的单侧 t 下界仅在 IID 正态配对增益下精确，一般有限方差下属于渐近论证。不能将整体包装为任意分布下的有限样本保证。基线失败的配对仍保留且阻止验收；零观测失败不等于总体失败概率为零。当前正式两阶段实验尚未完成，不据此填写通过结论。')

    doc.add_page_break()
    h('4 公平对照与工程适用范围')
    p('MAMPS 已研究耦合多智能体的学习／回退选择，DMPS 已研究任务相关的安全恢复，CSPI 和 Learn then Test 已提供策略改进检验及风险校准工具。因此，联合切换、加入扰动采样或使用二项检验本身均不能作为首次贡献。本文应由同一闭环的实际收益、采用和逐车退化取舍，以及明确的模型失效机制支撑贡献。')
    table(['预登记比较','固定范围','回答的问题'], [
        ['共享证据且保留失败','640 个 IID 初态\n5,120 次执行','同等证据构造计费下的选择器比较'],
        ['候选集合与门控 factorial','相同 640 个初态\n3,200 次执行','分离动作投影与扰动门控作用'],
        ['原 CSPI 作者检验补充','256/128/128/128\n拟合／调节／安全／评估','保留原单臂标签和独立评估的信息边界'],
    ], [5.2,4.2,7.0])
    p('三组正式结果均待完整执行与审计。当前全网格中的 MPS／MAMPS 为共享物理判据适配器，不能充当原方法完整复现；已有 DMPS 固定预算任务适配没有学会完成普通测试，不能据此判定方法族无效。CSPI 与在线门控的信息、训练数据和计算方式不同，补充比较不宣称等数据或等计算。')
    p('工程方面，全网格集合门控的决策 P95 为 750.792 ms，44,039 次调用中 3,234 次超过 500 ms，不能宣称满足 500 ms 硬实时要求。另已完成的时间戳补偿开发探针复用三个初态、共 117 次真实 SUMO：严重失配条件下三种预处理均为 0/13 完成。停止仿真不等于实际安全停车，更准的状态中点也未修复任务。')
    p('扩展 OOD 计划为 22,464 次执行、192 个独立初态，完整结果仍待审计。证明适用范围必须保留碰撞、几何违规、中止和超时；不能删除失败后仅比较完成子集。当前支持的结论是：有限路径性能门控在已完成开发网格中呈现有价值的收益与退化取舍，同时仍受路径覆盖、执行器记忆、离散时间指标和计算代价限制。')
    p('参考文献与复核来源', 'Heading 2')
    for text in [
        '[1] Zhang W, Bastani O, Kumar V. MAMPS: Safe Multi-Agent Reinforcement Learning via Model Predictive Shielding. 2019. arXiv:1910.12639.',
        '[2] Banerjee A, Rahmani K, Biswas J, Dillig I. Dynamic Model Predictive Shielding for Provably Safe Reinforcement Learning. 2024. arXiv:2405.13863.',
        '[3] Cho B M, et al. CSPI-MT: Calibrated Safe Policy Improvement with Multiple Testing for Threshold Policies. 2024. arXiv:2408.12004.',
        '[4] Angelopoulos A N, et al. Learn then Test: Calibrating Predictive Algorithms to Achieve Risk Control. 2022. arXiv:2110.01052.',
        '[5] 完整原始记录、冻结协议与重放入口：github.com/Leo-7GOAT/CoopGate。',
    ]:
        paragraph = p(text)
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_after = Pt(3)
        for run in paragraph.runs:
            run.font.size = Pt(9)

    assert doc.sections[0]._sectPr.xml == section_xml
    OUT.mkdir(parents=True)
    target = OUT/NAME
    doc.save(target)
    with ZipFile(REFERENCE) as before, ZipFile(target) as after:
        preserve = ['word/styles.xml','word/theme/theme1.xml','word/numbering.xml']
        for name in preserve:
            if name in before.namelist():
                assert before.read(name) == after.read(name), name
    assert sha(REFERENCE) == REFERENCE_SHA
    record = dict(utc=datetime.now(timezone.utc).isoformat(), document_sha256=sha(target),
        reference_sha256=REFERENCE_SHA, source_sha256=sha(__file__),
        completed_inputs_sha256={REPORT:REPORT_SHA, interpretation.relative_to(ROOT).as_posix():sha(interpretation)},
        existing_manuscript_unchanged=True, active_partial_outcomes_read=False,
        new_scientific_result_generated=False, native_equations=len(list(doc.element.body.iter(qn('m:oMath')))),
        explicit_page_sections=4, final_visual_review_still_required=True)
    (OUT/'build.json').write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(dict(output=str(target), **record),ensure_ascii=False))


if __name__ == '__main__':
    build()
