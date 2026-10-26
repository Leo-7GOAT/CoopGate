"""Verify and deliver the visually reviewed four-page supplement."""
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import shutil
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
import pypdfium2 as pdfium

ROOT = Path(__file__).resolve().parents[1]
RUN = ROOT/'runs/completed_evidence_supplement_20261026/revision2'
PUBLIC = ROOT/'artifacts/paper_supplement_20261026'
NAME = 'CoopGate_论文实验与理论补充_20261026.docx'


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def verify():
    build = read(RUN/'build.json')
    document = RUN/NAME
    assert sha(document) == build['document_sha256']
    assert sha(ROOT/'operations/build_completed_evidence_supplement.py') == build['source_sha256']
    for name, expected in build['completed_inputs_sha256'].items():
        assert sha(ROOT/name) == expected
    source = ROOT/'runs/manuscript_v17/candidate_1.docx'
    assert sha(source) == build['reference_sha256']
    doc = Document(document)
    assert doc.sections[0]._sectPr.xml == Document(source).sections[0]._sectPr.xml
    assert len(list(doc.element.body.iter(qn('m:oMath')))) == 6
    aggregate = read(ROOT/'artifacts/v38/interrupted_block_grid_recovery/v36/analysis/interpretation.json')['aggregate']
    for row, key in zip(doc.tables[0].rows[1:], ['baseline','nominal_performance','bank_performance'], strict=True):
        item = aggregate[key]
        assert [cell.text for cell in row.cells][1:] == [f"{item['completed']}/{item['runs']}", str(item['adopted']), str(item['harmful_and_adopted']), f"{item['all_scope_cost_gain_percent']:.4f}%"]
    mechanism_name = 'artifacts/diagnostics/v36_all_bank_regressions/report.json'
    mechanism = read(ROOT/mechanism_name)
    assert len(mechanism['cases']) == 5
    assert mechanism['physical_replays'] == 16 and mechanism['adoptions'] == 8
    assert mechanism['new_independent_initials'] == 0 and not mechanism['new_risk_certificate']
    probe_name = 'artifacts/diagnostics/timestamped_control_probe_v2/verification.json'
    probe = read(ROOT/probe_name)
    assert probe['actual_sumo_runs'] == 117 and probe['reused_independent_initials'] == 3
    assert not probe['new_risk_certificate'] and not probe['marl_training_claimed']
    pdf = pdfium.PdfDocument(RUN/'word_render.pdf')
    assert len(pdf) == 4
    texts = [page.get_textpage().get_text_range() for page in pdf]
    assert all(len(text) > 700 and '\ufffd' not in text for text in texts)
    assert 'Dynamic Model Predictive Shielding' in texts[3]
    assert 'Differentiable Model' not in texts[3]
    assert '尚未完成' in texts[2] and '不能' in texts[3]
    pages = {str(i):sha(RUN/f'word_render/page-{i}.png') for i in range(1,5)}
    with ZipFile(source) as before, ZipFile(document) as after:
        preserved = [name for name in before.namelist() if name in after.namelist()
                     and name != 'word/document.xml' and before.read(name) == after.read(name)]
        for name in ['word/styles.xml','word/theme/theme1.xml','word/numbering.xml']:
            assert name in preserved
    record = dict(utc=datetime.now(timezone.utc).isoformat(), document_sha256=sha(document),
        pages=4, all_final_pages_visually_inspected=True, inspected_pages=[1,2,3,4],
        page_png_sha256=pages, pdf_sha256=sha(RUN/'word_render.pdf'),
        verifier_sha256=sha(__file__), additional_completed_evidence_sha256={
            mechanism_name:sha(ROOT/mechanism_name), probe_name:sha(ROOT/probe_name)},
        native_equations=6, numerical_table_matches_completed_report=True,
        original_manuscript_unchanged=True, active_partial_efficacy_read=False,
        final_risk_fairness_ood_results_not_claimed=True,
        preserved_package_parts=preserved,
        renderer='Packaged renderer attempted: missing soffice.exe; separate hidden read-only Word export; bundled pypdfium2 rasterization',
        original_export_empty_word_instance_closed_after_unique_pid_creation_and_empty_document_checks=True,
        primary_references_verified_on='2026-10-26', primary_reference_urls=[
            'https://arxiv.org/abs/1910.12639','https://arxiv.org/abs/2405.13863',
            'https://arxiv.org/abs/2408.12004','https://arxiv.org/abs/2110.01052'])
    PUBLIC.mkdir(exist_ok=False)
    shutil.copyfile(document, PUBLIC/NAME)
    shutil.copyfile(RUN/'build.json', PUBLIC/'build.json')
    shutil.copyfile(RUN.parent/'artifact.md', PUBLIC/'artifact.md')
    (PUBLIC/'qa.json').write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    desktop = ROOT.parent.parent/NAME
    with desktop.open('xb') as output:
        output.write(document.read_bytes())
    assert sha(desktop) == sha(PUBLIC/NAME) == build['document_sha256']
    print(json.dumps(dict(desktop=str(desktop),public=str(PUBLIC/NAME),pages=4,sha256=sha(desktop)),ensure_ascii=False))


if __name__ == '__main__':
    verify()
