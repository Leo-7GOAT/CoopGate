"""Verify the retained Word manuscript against completed evidence and final render.

This checks content and package integrity. It does not perform visual inspection;
the separate qa.json records that manual review and the exact inspected PNG hashes.
Run with the bundled document Python runtime.
"""
import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from lxml.etree import tostring
import pypdfium2 as pdfium

ROOT = Path(__file__).resolve().parents[1]


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def read(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def canonical(element):
    return tostring(element, method="c14n", exclusive=True)


def verify(reference, document, pdf, render_dir, build_record, output):
    build = read(build_record)
    assert sha(reference) == build["reference_sha256"]
    assert sha(document) == build["document_sha256"]
    assert sha(ROOT / "tools/build_v17_manuscript.py") == build["source_sha256"]
    for path, digest in build["input_sha256"].items():
        assert sha(ROOT / path) == digest, path
    old, new = Document(reference), Document(document)
    old_math = [canonical(e) for e in old.element.body.iter(qn("m:oMath"))]
    new_math = [canonical(e) for e in new.element.body.iter(qn("m:oMath"))]
    assert len(old_math) == 23 and len(new_math) == 28
    assert new_math[:23] == old_math
    assert len(old.tables) == 20 and len(new.tables) == 22
    assert [canonical(t._tbl) for t in new.tables[:20]] == [canonical(t._tbl) for t in old.tables]
    assert len(new.sections) == len(old.sections) == 1
    assert new.sections[0]._sectPr.xml == old.sections[0]._sectPr.xml
    assert len(new.inline_shapes) == len(old.inline_shapes) + 2
    assert all(s._inline.docPr.get("descr") for s in list(new.inline_shapes)[-2:])
    with ZipFile(reference) as before, ZipFile(document) as after:
        assert set(before.namelist()) <= set(after.namelist())
        changed = [n for n in before.namelist() if before.read(n) != after.read(n)]
        added = sorted(set(after.namelist()) - set(before.namelist()))
        assert changed == build["changed_parts"] and added == build["added_parts"]
        assert set(changed) <= {"word/document.xml", "docProps/core.xml", "[Content_Types].xml", "word/_rels/document.xml.rels"}

    comparison = read(ROOT / "artifacts/v25/historical_common_inference_audit/report.json")
    table_rows = new.tables[20].rows[1:]
    assert len(table_rows) == len(comparison["rows"]) == 9
    for row, data in zip(table_rows, comparison["rows"], strict=True):
        cells = [c.text for c in row.cells]
        assert data["model"] in cells[0]
        for written, key in zip(cells[1].split(" / "), ["calibration_actual", "calibration_bundle"], strict=True):
            assert abs(float(written.rstrip("%")) - 100 * data[key]["threshold_upper"]) <= .005
        assert abs(float(cells[2].rstrip("%")) - data["final_cost_gain_percent"]) <= .0005
        adoption = data["final_actual_adoption"]
        assert cells[3] == f"{adoption['adopted']}/512"
        assert cells[4] == f"{adoption['harmful_adopted']}/{adoption['adopted']}"
    timing = read(ROOT / "artifacts/v25/buffered_phase_presentation/report.json")
    rows = new.tables[21].rows[1:]
    assert len(rows) == len(timing["group_timings"]) == 5
    for row, data in zip(rows, timing["group_timings"], strict=True):
        cells = [c.text for c in row.cells]
        assert cells[0] == data["mode"] and int(cells[1]) == data["decisions"]
        assert abs(float(cells[2]) - data["reference_select_ms"] / 1000) <= .0005
        assert abs(float(cells[3]) - data["buffered_select_ms"] / 1000) <= .0005
        assert abs(float(cells[4].rstrip("%")) - data["reduction_percent"]) <= .005

    final_pdf = pdfium.PdfDocument(pdf)
    assert len(final_pdf) == 28
    texts = [p.get_textpage().get_text_range() for p in final_pdf]
    assert all(len(t.strip()) > 200 and "\ufffd" not in t for t in texts)
    for index, title in {23: "20 共同推断规则下的公平比较", 24: "21 保持完整判据的计算改进", 25: "22 采用后的风险迁移边界", 26: "匹配迁移的采用与退化事件", 27: "23 讨论与结论", 28: "参考文献"}.items():
        assert texts[index-1].startswith(title)
    assert "仍在执行" in texts[21] and "最终确认尚未开始" in texts[21]
    assert "全量结果" in texts[23] and "不新增迁移风险证书" in texts[24]
    pages = {str(i): sha(render_dir / f"page-{i}.png") for i in range(1, 29)}
    old_qa = read(ROOT / "artifacts/v21/manuscript_v16/qa.json")
    identical = [i for i in range(2, 22) if pages[str(i)] == old_qa["page_png_sha256"][str(i)]]
    result = dict(
        utc=datetime.now(timezone.utc).isoformat(), document_sha256=sha(document),
        reference_sha256=sha(reference), build_record_sha256=sha(build_record),
        verifier_source_sha256=sha(__file__), input_files_checked=len(build["input_sha256"]),
        changed_parts=changed, added_parts=added, retained_equations=23, added_equations=5,
        original_tables_unchanged=20, added_tables_checked_against_completed_data=2,
        pages=28, page_png_sha256=pages, pdf_sha256=sha(pdf),
        visually_reviewed_reference_pages_with_identical_rasters=identical,
        all_original_styles_and_theme_parts_exact=True, original_section_exact=True,
        active_calibration_outcomes_read=False, new_risk_certificate=False,
        visual_inspection_performed_by_this_script=False,
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({k: v for k, v in result.items() if k != "page_png_sha256"}, ensure_ascii=False))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    for arg in ["reference", "document", "pdf", "render-dir", "build-record", "output"]:
        parser.add_argument("--" + arg, type=Path, required=True)
    args = parser.parse_args()
    verify(args.reference, args.document, args.pdf, args.render_dir, args.build_record, args.output)
