"""Chinese second manuscript, generated from the complete frozen v2 study."""
from pathlib import Path
import argparse,json,re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
DATA=ROOT/'artifacts/v2_final'
NAMES={'traditional':'传统规控','mappo':'MAPPO','independent':'独立门控','joint_robust':'V1联合门控',
       'paired_no_gain':'去收益筛选','paired_no_reserve':'去制动余量','paired_no_memory':'去执行器记忆','paired':'完整门控'}


def build(output):
    result=json.loads((DATA/'study_summary.json').read_text())
    timing=json.loads((DATA/'timing.json').read_text())
    summary=result['summaries'];cert=result['certification']
    doc=Document();sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=Cm(2.2);sec.left_margin=sec.right_margin=Cm(2.35)
    for border in doc.styles.element.xpath('.//w:pBdr'):border.getparent().remove(border)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Caption'):
        s=doc.styles[name];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0)
        s.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    doc.styles['Normal'].font.size=Pt(11)
    doc.styles['Normal'].paragraph_format.line_spacing=1.22
    doc.styles['Normal'].paragraph_format.space_after=Pt(6)
    doc.styles['Title'].font.size=Pt(19)
    doc.styles['Subtitle'].font.size=Pt(11);doc.styles['Subtitle'].font.italic=False
    doc.styles['Heading 1'].font.size=Pt(14);doc.styles['Heading 2'].font.size=Pt(12)
    doc.styles['Caption'].font.size=Pt(9)
    for s in ('Heading 1','Heading 2'):doc.styles[s].paragraph_format.keep_with_next=True
    doc.core_properties.title='考虑执行器状态与备用收益的多车策略切换研究'
    doc.core_properties.author='李东旭'

    def p(text,style=None):
        q=doc.add_paragraph(text,style)
        if style is None:q.paragraph_format.first_line_indent=Cm(.74)
        return q
    def h(text):doc.add_heading(text,1)
    def eq(text):
        q=doc.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER
        m=OxmlElement('m:oMath')
        def run(parent,value):
            r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=value
            t.set(qn('xml:space'),'preserve');r.append(t);parent.append(r)
        position=0
        for token in re.finditer(r'([A-Za-z]+)_(?:\{([^}]+)\}|([A-Za-z0-9ω]+))',text):
            if token.start()>position:run(m,text[position:token.start()])
            script=OxmlElement('m:sSub');base=OxmlElement('m:e');sub=OxmlElement('m:sub')
            run(base,token.group(1));run(sub,token.group(2) or token.group(3))
            script.append(base);script.append(sub);m.append(script);position=token.end()
        if position<len(text):run(m,text[position:])
        q._p.append(m)
    def table(headers,rows,widths):
        t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
        borders=OxmlElement('w:tblBorders')
        for edge in ('top','left','bottom','right','insideH','insideV'):
            el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
        t._tbl.tblPr.append(borders)
        for c,w,label in zip(t.rows[0].cells,widths,headers):c.width=Cm(w);c.text=label
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for values in rows:
            for c,w,value in zip(t.add_row().cells,widths,values):c.width=Cm(w);c.text=str(value)
        for i,row in enumerate(t.rows):
            el=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(el)
            for cell in row.cells:
                cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'E7ECF2' if i==0 else ('FFFFFF' if i%2 else 'F5F7FA'))
                cell._tc.get_or_add_tcPr().append(sh)
                for q in cell.paragraphs:
                    q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.space_after=Pt(4)
                    q.paragraph_format.space_before=Pt(4);q.paragraph_format.line_spacing=1.0
                    for r in q.runs:r.font.size=Pt(9);r.bold=i==0
        p('')

    p('考虑执行器状态与备用收益的多车策略切换研究','Title')
    p('李东旭','Subtitle')
    h('摘要')
    success='、'.join(f'{r["successes"]}/15' for r in summary['test']['paired'])
    shifted='、'.join(f'{r["successes"]}/15' for r in summary['shift']['paired'])
    passed=sum(r['certified'] for r in cert)
    p('针对多车强化学习策略与传统规划控制之间的接管问题，本文研究执行器持续状态、候选收益与安全检查之间的耦合。'
      '提出一个备用相对门控实现：显式保留上一纵向指令和正在执行的换道目标，在共享纵向MPC预测器下比较候选与全备用分支，'
      '并以决策后的连续时间制动余量筛选候选。本文给出制动包络的精确计算、成对收益的条件误差界，以及冻结闭环控制器的回合级统计检验。'
      '在SUMO连续合流路网上训练三个独立策略，完成1030次主实验和90次执行器记忆补充消融。'
      f'完整门控在同分布新场景中的成功次数分别为{success}，在更强制动场景中分别为{shifted}；'
      f'独立认证集的8个风险检验中有{passed}项通过预设5%目标。'
      '完整门控保持了接近传统规控的完成率，但学习动作使用比例较低；执行器记忆消融未显示稳定的额外收益。'
      '结果用于辨析各模块的实际作用及适用分布。局部制动模型结论不等于多车无限时域安全，有限预测收益也不等于整回合性能保证。')
    p('关键词  多智能体强化学习  混合控制  备用策略  制动余量  统计风险检验')

    h('1 引言')
    p('学习策略能够提出灵活的多车行为，但部署时需要判断其决策是否可执行。一个直接做法是根据预测分数选择学习或传统规划。'
      '困难在于，控制系统具有持续状态：已经开始的换道不能被视为每个决策周期都可以任意重置的瞬时动作，纵向控制的变化率约束也依赖上一指令。'
      '此外，单纯避免碰撞并不保证任务完成；低速等待也可能得到较大的预测安全裕度。')
    p('本研究关注三个相互关联的问题：候选检查是否与实际执行器一致，放行学习动作是否具有相对备用方案的收益依据，'
      '以及完整闭环在指定场景分布上的失败风险是否得到独立检验。本文采用三车合流任务进行受控实验，不将问题扩展为视觉感知或完整自动驾驶系统。')
    p('本文实现的贡献包括：具有换道承诺状态的监督器接口、共享纵向QP的成对预测比较、精确制动位移包络和可复现的回合级风险报告。'
      '这些组件采用既有控制与统计工具，整体新颖性仍须通过更广泛查新和强基线验证；本文不把模块组合或标准推断方法本身作为原创定理。')

    h('2 相关研究')
    p('MAMPS已经研究多智能体回退的连锁影响，通过可恢复性判断逐步调整各智能体的模式[1]。鲁棒模型预测屏障进一步处理随机动力学和备用控制[2]。'
      '因此，联合切换、备用控制和不确定性预测并非本文首先提出。')
    p('DMPS研究任务无关的保守回退如何妨碍进展，并利用恢复规划与学习价值估计改善任务表现[3]。本文相对备用方案的收益筛选与这一研究动机相关，'
      '不能宣称首次让安全回退兼顾效率。本文当前使用有限时域效用而非已验证的长期价值函数，也没有直接复现DMPS原方法。')
    p('MAPPO为共享Actor和集中Critic的协作学习提供基础[4]。本文用行为克隆初始化及辅助教师损失降低训练难度，所有门控对照共用同一冻结策略。'
      '风险报告采用已有的精确二项置信上界与有限族错误控制，思路与Learn then Test相关[5]。这与将运行时裕度分数直接解释为安全概率有本质区别。')

    h('3 系统模型与候选生成')
    p('环境包含三辆受控车与两辆SUMO背景车。受控车辆选择九个离散行为，由8、14、20 m/s三档目标速度及三个车道意图组合。'
      '规则备用规划根据前车距离调速，并在存在足够间隙时汇入。两种候选均由纵向MPC执行，横向由SUMO连续换道模型执行。')
    eq('z = (x, a_prev, lane_target, r)')
    p('x包含位置、速度与加速度，a_prev为上一加速度指令，lane_target与r为正在执行换道的目标和剩余承诺时间。'
      '当r大于零时，监督器和执行器均保持该目标。联合监督器使用车组共享状态，因此整体系统不被描述为完全分布式。'
      '背景车的隐藏行为状态没有完整建模，此增广状态也不等于完整仿真器状态。')
    eq('J_MPC = ∑ [(v_{k+1} − v_ref)² + 0.1a_k² + 0.3(a_k − a_{k−1})²]')
    p('求和范围为k等于0至N−1，初始a下标−1为上一指令。MPC时域12步，步长0.5 s，最大加速度2.5 m/s²、常规最大减速度5 m/s²、加加速度上限5 m/s³。'
      '监督器预测时调用相同的OSQP纵向控制器，并按0.1 s积分。横向和背景预测仍然近似，因此共享QP并不意味着预测完全准确。')

    h('4 备用相对门控方法')
    p('对三辆车枚举八种学习与备用模式组合。每种组合先执行一个候选周期，再按规则备用预测，共预测3 s。'
      '背景加速度假设为−2.5、0和1.5 m/s²，分别计算候选和全备用分支，使用同一背景假设形成配对。')
    eq('Δ̂(b) = min_ω [Û(z_H(b,ω)) − Û(z_H(B,ω))]')
    p('效用U采用各车min(x+1.5v, x_goal)的总和，并对未完成合流施加5点代价；车道代价在接近终点时连续衰减。'
      '这种写法避免车辆跨过终点时突然失去速度奖励。效用是有限时域代理量，并非已证明正确的备用价值函数。')
    p('候选需同时满足预测几何裕度非负、候选执行一个周期后预测制动余量非负，以及配对收益至少为0.5。'
      '通过检查的候选按配对收益减去0.08倍模式切换数排序。全备用始终是默认方案；若它本身也未通过检查，则显式记录未验证回退，'
      '不能将回退动作自动当作获得了安全证书。')
    p('计算上先评估全备用分支，再按固定次序搜索其他模式。科学实验给足枚举预算，使各次运行检查相同候选集合。'
      '工程测试使用150 ms软预算，在开始下一候选前检查时间。该实现不具备中断一次预测的能力，也没有硬实时保证。')

    h('5 理论分析')
    doc.add_heading('5.1 连续时间制动包络',2)
    p('考虑固定车道的一对前后车。前车立即制动，后车在反应时间内加速，随后制动至停止。记位移分别为sF与sR，初始净间距为d，最低允许净间距为d0。定义')
    eq('D = max{0, sup_t [s_R(t) − s_F(t)]},     R = d − d_0 − D')
    p('命题1：对指定的两条运动曲线，整个制动过程保持净间距至少d0，当且仅当R非负。证明：任意时刻净间距为d+sF(t)−sR(t)，'
      '对所有时刻施加约束并取上确界即可。由于速度分段线性，位移差分段二次，最大值只可能位于相位端点或相对速度为零的分段内部点。'
      '实现穷举这些有限候选点，避免时间采样遗漏峰值。')
    p('一个反例说明仅比较最终停车位置不够：后车20 m/s、前车10 m/s，制动强度分别9与3 m/s²且无反应时间时，'
      '最终位移差约5.56 m，而过程最大接近量约8.33 m。实现中的该反例及随机密集时间网格对照均纳入测试。')
    p('实际门控假定后车反应0.5 s、反应期加速度2.5 m/s²、随后制动5 m/s²，前车可立即制动9 m/s²。'
      '该判据只对指定纵向模型成立，不涵盖不可控后车追尾、横向侵入或任意参数误差。代码检查的是决策后制动余量，尚未构造多车不变终端集合。')
    doc.add_heading('5.2 成对收益的条件误差界',2)
    p('命题2：若真实背景对应某个被检查假设，且候选与备用分支的效用预测误差各不超过ε，则放行候选的真实成对收益至少为η−2ε。'
      '证明：真实差值不小于相应预测差值减2ε，而相应预测差值不小于假设集上的最小差值；放行条件使后者至少为η。'
      '在搜索多个候选后使用本结论时，误差界必须对可能被选择的候选同时成立。')
    p('当前并未获得这样的统一误差界，因此η等于0.5不代表某个统计安全概率。有限时域代理收益也不能推出整回合累计收益改善。'
      '若未来使用满足Bellman关系的备用代价函数，并使每次选择的即时成本与后继备用代价之和不超过当前备用代价，才能通过逐步相加建立对应累计界。')
    doc.add_heading('5.3 冻结闭环的回合级风险检验',2)
    p('对一个事先冻结的控制器，在n个独立同分布回合中观测k次失败。设需要同时检验的控制器与事件组合总数为K，'
      '整体错误概率为δ，采用单侧精确二项上界。k小于n时写作')
    eq('p_U = Beta⁻¹(1 − δ/K ; k+1, n−k)')
    p('命题3：在冻结控制器及独立同分布采样假设下，以至少1−δ的置信水平，全部K个真实风险同时不超过对应上界。'
      '证明：每个单侧精确区间的欠覆盖概率至多δ/K，对全部欠覆盖事件使用并集界即可。不同控制器可共享配对场景，此证明不要求控制器之间独立。'
      '这些属于标准统计工具的应用，而非本文提出的新推断定理。')
    p('零失败时上界为1−(δ/K)的1/n次方。本文固定传统控制器及三个paired策略，分别检验碰撞和任务未完成，共K等于8项。'
      'δ等于0.05、n等于100时，即使零失败，上界仍约4.95%。只有上界不超过预设5%目标才记为通过。'
      '该目标仅演示仿真统计检验，不是道路部署标准；也不适用于分布变化或调参后复用同一认证集。')

    h('6 实验协议与实现审计')
    p('三组训练种子为31、47、59，每组先用90回合规则示范进行20轮行为克隆，再运行30000步MAPPO，附加0.1权重的教师动作损失。'
      '使用最终检查点，不根据测试结果筛选。所有方法共享冻结策略、车辆约束、换道承诺执行器及路网。每回合最长30 s，三车都到达410 m位置才算成功。')
    p('主对照有七种方法：传统规控、无门控MAPPO、独立门控、V1联合多假设门控、去收益筛选、去制动余量及完整门控。'
      '每模型每方法测试15个同分布新场景和15个更强制动场景，共630次执行。另对四个冻结控制器运行100个独立认证场景，共400次。'
      '执行器记忆消融另有90次，属于补充探索性实验。重复场景或同一场景中的车辆、帧都不能当作独立统计试验。')
    p('同分布场景来自正常、较密和前车制动模板的均匀混合。分布变化场景使用密集间距，并让背景前车制动至停车且持续更久。'
      '受控车关闭SUMO隐式安全干预，背景车保留仿真行为模型；不引入通信延迟实验。')
    p('开发审计发现，关闭内部连接会在节点附近产生约8 m位置跳变。该轮中间实验已中止并归档，正式实验改用连续路网。'
      '在五个开发场景、150个决策步上，共享MPC预测器的位置和速度RMSE分别约0.060 m和5.75×10⁻⁷ m/s；'
      '同一路网上旧纵向近似对应约0.073 m和0.148 m/s。此结果只刻画这些开发轨迹的单步误差。')
    p('效用函数终点不连续的问题也在代码审查中修正，并换用全新评估场景重跑。最终配置、场景种子和模型校验值随代码发布，'
      '中止的开发运行不与正式数据合并。')

    h('7 实验结果')
    for split,label in (('test','同分布新场景'),('shift','更强制动分布变化')):
        doc.add_heading('7.1 '+label if split=='test' else '7.2 '+label,2)
        p('每个单元格为碰撞次数与成功次数，以15回合为分母。失败罚时将所有未完成回合记为30 s，再对三组策略取均值。')
        data=[]
        for method,name in NAMES.items():
            group=summary[split][method]
            data.append([name]+[f'{r["collisions"]} / {r["successes"]}' for r in group]
                        +[f'{sum(r["penalized_time_s"] for r in group)/3:.2f}'])
        table(['方法','种子31','种子47','种子59','罚时 s'],data,[4,3,3,3,3])
        if split=='test':
            p('完整门控与传统、独立门控在三组策略上都完成15个场景。完整门控平均用时21.94 s，传统为22.43 s，独立门控为20.69 s。'
              '完整门控并非本分布中效率最高的方法。去除制动余量后，三组策略各发生一次碰撞；完整门控的平均学习动作比例仅4.7%。')
        else:
            p('完整门控三组均完成13/15回合，碰撞分别为0、0、1次；传统同样完成13/15回合，并在同一个场景发生碰撞。'
              '无门控MAPPO三组均未完成任何回合，碰撞分别为15、15、14次。去除收益筛选后，种子31与59的完成数降为5，种子47仍为13；'
              '去除制动余量后，碰撞分别升为8、7、5次。这些结果支持在本测试集上同时约束安全余量和备用相对收益。')
            p('去除执行器记忆后的完成数和碰撞数与完整门控完全一致，失败罚时均值也一致；同分布用时差仅约0.04 s。'
              '因此，当前实验不能支持执行器记忆带来稳定性能提升。完整门控在分布变化中的学习动作比例约2.1%，主要仍依赖传统备用，'
              '应定位为退化防护的初步证据，不能据此声称已经充分发挥多智能体学习的协同优势。')
    doc.add_heading('7.3 独立风险检验',2)
    table(['控制器','事件','失败数','风险上界','通过'],
          [[('传统' if r['method']=='traditional' else '门控'+str(r['training_seed'])),
            '碰撞' if r['endpoint']=='collision' else '未完成',r['failures'],f'{r["risk_upper"]:.2%}',
            '是' if r['certified'] else '否'] for r in cert],[3.5,3,2.5,3.5,2.5])
    p('八项检验均通过预设目标，传统基线也通过。传统基线的重复执行不作为额外独立证据。'
      '对分布变化的任何改善或退化，只使用该测试集本身的数据描述，不能沿用同分布认证结论。')
    doc.add_picture(str(DATA/'results.png'),width=Cm(16))
    doc.paragraphs[-1].paragraph_format.keep_with_next=True
    p('图1 各方法的任务完成率与失败罚时 黑点为三个训练种子而非置信区间','Caption')
    doc.add_heading('7.4 计算开销',2)
    table(['方法','循环均值 ms','循环P95 ms','循环最大 ms','超500 ms'],
          [[NAMES[m],f'{r["mean_ms"]:.2f}',f'{r["p95_ms"]:.2f}',f'{r["max_ms"]:.2f}',r['over_500ms']]
           for m,r in timing['loop'].items()],[3.2,3.2,3.2,3.2,3.2])
    p('上述循环时间包含策略推理、监督器、MPC与TraCI步进，在主实验完成后单进程测量。'
      '该样本测量不能证明最坏执行时间界，也不能保证其他硬件或更大车组满足相同周期。三车完整枚举不能直接扩展到任意数量车辆。'
      '150 ms软预算可能提前截断候选枚举，形成不同的闭环策略；它不自动继承完整枚举控制器的风险检验结果。')

    h('8 讨论与后续研究')
    p('本工作的工程意义在于把接管判断与实际执行状态连接起来，并将未验证回退、任务失败和计算开销显式记录。'
      '目前完整门控的完成率与传统策略相同，效率改善有限且学习动作比例偏低，现有证据主要支持学习策略退化防护。'
      '收益筛选与制动余量消融给出了值得继续检验的信号；执行器记忆没有表现出稳定收益，不适合独立作为论文的主要贡献。')
    p('理论方面，制动包络是指定纵向模型的精确结论，配对收益是依赖统一误差界的条件结论，统计检验是指定回合分布的概率结论。'
      '三者不能互相替代。目前仍缺少覆盖持续换道与不可控车辆的联合可恢复集合、统一预测误差界和完整强基线复现。')
    p('投稿层面的下一步应由消融与失败轨迹决定：验证执行器记忆在更复杂交互中的作用，构造可实际检验的收益误差界，'
      '并在同等约束和计算预算下与MAMPS、DMPS等方法比较。随后再增加道路结构和CARLA连续车辆动力学验证。'
      '在这些证据完成前，不声称首次、最先进性能或真实道路安全保证。')
    h('9 结论')
    p('本文实现并评估了具有执行器状态、备用相对收益和制动余量的多车策略切换系统，提供了条件理论分析、独立回合风险检验与可复现实验。'
      '本次推进使研究问题从启发式分数切换具体化为执行一致性、备用相对效用和分布限定风险的检验，但是否形成有竞争力的算法创新仍需更强证据。')
    h('参考文献')
    refs=['[1] Zhang W, Bastani O, Kumar V. MAMPS Safe Multi-Agent Reinforcement Learning via Model Predictive Shielding. 2019. https://arxiv.org/abs/1910.12639',
          '[2] Li S, Bastani O. Robust Model Predictive Shielding for Safe Reinforcement Learning with Stochastic Dynamics. 2019. https://arxiv.org/abs/1910.10885',
          '[3] Banerjee A, Rahmani K, Biswas J, Dillig I. Dynamic Model Predictive Shielding for Provably Safe Reinforcement Learning. NeurIPS, 2024. https://arxiv.org/abs/2405.13863',
          '[4] Yu C et al. The Surprising Effectiveness of PPO in Cooperative Multi-Agent Games. NeurIPS, 2022. https://arxiv.org/abs/2103.01955',
          '[5] Angelopoulos A N et al. Learn then Test Calibrating Predictive Algorithms to Achieve Risk Control. 2021. https://arxiv.org/abs/2110.01052',
          '[6] Eclipse SUMO Documentation. https://eclipse.dev/sumo/docs/']
    for ref in refs:
        q=p(ref);q.paragraph_format.first_line_indent=Cm(0)
        for run in q.runs:run.font.size=Pt(9)
    q=sec.footer.paragraphs[0];q.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');q._p.append(field)
    doc.save(output);print(output)


if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('--out',required=True);a=p.parse_args();build(a.out)
