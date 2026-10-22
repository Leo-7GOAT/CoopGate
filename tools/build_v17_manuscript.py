"""Integrate completed V25 evidence into a new retained-style Word revision."""
import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table
from lxml.etree import tostring

ROOT=Path(__file__).resolve().parents[1]
REFERENCE_SHA='201c52c5be21a00cc4f7ff65758891ed47456c24fa708ef71f90134277c2f6fd'


def digest(path):return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def build(reference,out,evidence_out):
    if reference.resolve()==out.resolve() or digest(reference)!=REFERENCE_SHA:
        raise ValueError('Retain the exact sixteenth manuscript')
    inputs={}
    def read(name):
        inputs[name]=digest(ROOT/name)
        return json.loads((ROOT/name).read_text(encoding='utf-8'))
    comparison=read('artifacts/v25/historical_common_inference_audit/report.json')
    timing=read('artifacts/v25/buffered_phase_presentation/report.json')
    paired=read('artifacts/v25/buffered_phase_sumo_audit/results/report.json')
    events=read('artifacts/v25/matched_event_transfer/report.json')
    event_check=read('artifacts/v25/matched_event_transfer_checks/verification.json')
    event_plot=read('artifacts/v25/matched_event_transfer_presentation/summary.json')
    scope=read('artifacts/v25/buffered_complete_scope/plan.json')
    assert comparison['scope']=='historical' and len(comparison['rows'])==9
    assert comparison['inference_rule_identical_for_all_nine']
    assert not comparison['new_risk_certificate'] and not comparison['unfinished_prospective_outcomes_read']
    assert timing['actual_sumo_runs']==paired['actual_sumo_runs']==24
    assert timing['source_report_sha256']==inputs['artifacts/v25/buffered_phase_sumo_audit/results/report.json']
    assert timing['all_non_timing_fields_exact'] and paired['all_24_trajectories_match_reviewed_archive']
    assert paired['independent_review_counts']['physical_traces']==24
    assert timing['execution_observation']['exit_code']==0
    assert len(events['cells'])==270 and event_check['source_target_summary_comparisons']==540
    assert event_check['report_sha256']==event_plot['source_sha256']==inputs['artifacts/v25/matched_event_transfer/report.json']
    assert not events['risk_certified'] and not events['active_calibration_outcomes_used']
    assert scope['actual_sumo_runs']==6912 and scope['paired_cases']==3456
    doc=Document(reference)
    original_section=doc.sections[0]._sectPr.xml
    original_math=[tostring(e,method='c14n',exclusive=True) for e in doc.element.body.iter(qn('m:oMath'))]
    assert len(original_math)==23
    prototype=deepcopy(doc.tables[0]._tbl)
    doc.core_properties.version='17'

    def replace(prefix,text):
        matches=[p for p in doc.paragraphs if p.text.startswith(prefix)]
        if len(matches)!=1:raise ValueError('Ambiguous source slot: '+prefix)
        matches[0].text=text
        return matches[0]
    replace('单车前置验证和三车模型式协同训练完成后，',
        '单车前置验证和三车模型式协同训练完成后，原17路径门控仅有一份固定模型通过全部预定风险与收益要求。'
        '新增128个独立初态的6,656次开发执行中，50 ms时序门控保留0.76%–1.03%的成本改善和28.1%–30.5%的采用率，'
        '密集交互收益高于普通场景；另12,160次匹配迁移执行暴露低速让行退化和强惯性失败。'
        '保持行为一致的缓冲实现，在一个旧初态的24次配对SUMO检查中减少75.4%完整门控耗时，仍有8/277次决策超过500 ms。'
        '本文进一步区分统计检验口径，并推导采用率对条件风险迁移界的影响。新一轮独立风险校准和完整范围时延验证仍在执行。')
    replace('本文考察完整联合组合、动作调整和时序检查对实际采用、逐车性能和计算代价的影响，',
        '本文考察完整联合组合、动作调整和时序检查对实际采用、逐车性能和计算代价的影响。'
        '第11至15节保留原17路径研究；第16至19节给出时序门控、完整开发与匹配迁移结果，以及独立风险检验。'
        '第20至22节补充共同推断口径、保持行为的计算改进和采用后风险的迁移边界。原研究与新门控的证据分别报告。')
    replace('完整执行器已在一个旧初态上通过90次实际SUMO集成，',
        '完整执行器已在一个旧初态上通过90次实际SUMO集成，含40次主执行、10次影子预测和40次部署回放，'
        '18个运行时选择均与离线接口一致。该集成不增加新独立样本。正式校准的原执行发生中断，原退出状态未知；'
        '已保留并核验原始前缀，显式恢复通过边界回放后继续补齐缺失角色。原模型、阈值、预算和统计规则保持不变。'
        '截至2026年10月22日，校准仍在执行，最终确认尚未开始。实际进程退出、完整记录封存和独立复核完成后才允许分析。')
    discussion=replace('20 讨论与结论','23 讨论与结论')
    replace('原研究的保守共形适配全部回退，',
        '原研究的保守共形适配全部回退，LTT只在一份模型采用严格门控，均值改进适配保留收益但不保证采用后条件风险。'
        '共同推断口径下只有严格门控1812通过两个旧风险目标，所有完整模型组均未通过。'
        '不可拆分的整组认证与允许选候选的多重检验回答不同命题，不能将检验负担差异当作算法收益。')
    replace('匹配外部噪声的迁移研究显示，',
        '匹配外部噪声的迁移研究显示，时序门控在名义、增益阻力和轻度惯性下均完成本组运行，低速让行仍有负收益；'
        '强惯性超出误差界并产生未完成。事件分布变化通过采用率分母影响条件风险界，物理误差界不能直接替代性能迁移保证。'
        '缓冲实现已在旧初态配对检查中明显降低计算耗时，完整开发与迁移范围的6,912次实现验证仍在进行。'
        '现有证据尚不能支持硬实时部署，独立风险校准与最终确认也仍需完成。')
    replace('代码、冻结模型、实验协议、逐回合记录和独立复核入口见公开仓库',
        '代码、冻结模型、实验协议、逐回合记录和独立复核入口见公开仓库[11,12,14,16]。'
        '当前尚无新增门控的跨模型风险认证、部署分布迁移保证或硬实时保证。')
    anchor=discussion._p.getprevious()
    def place(element):anchor.addprevious(element)
    def p(text,caption=False):
        x=doc.add_paragraph(text)
        x.paragraph_format.first_line_indent=Cm(0 if caption else .74)
        if caption:x.paragraph_format.keep_with_next=True;x.paragraph_format.space_after=Pt(5)
        place(x._p);return x
    def page(title,level=1):
        place(doc.add_page_break()._p);place(doc.add_heading(title,level)._p)
    def table(headers,rows,widths):
        element=deepcopy(prototype);place(element);t=Table(element,doc._body)
        while len(t.rows)>1:t._tbl.remove(t.rows[-1]._tr)
        while len(t.columns)<len(headers):t.add_column(Cm(2.5))
        assert len(t.columns)==len(headers) and abs(sum(widths)-16.4)<1e-9
        t.autofit=False
        for c,v in zip(t.rows[0].cells,headers):c.text=v
        fill=deepcopy(t.rows[0].cells[0]._tc.get_or_add_tcPr().find(qn('w:shd')))
        for c in t.rows[0].cells:
            props=c._tc.get_or_add_tcPr();old=props.find(qn('w:shd'))
            if old is not None:props.remove(old)
            if fill is not None:props.append(deepcopy(fill))
        for values in rows:
            for c,v in zip(t.add_row().cells,values,strict=True):c.text=str(v)
        for c,w in zip(t.columns,widths):c.width=Cm(w)
        for i,row in enumerate(t.rows):
            props=row._tr.get_or_add_trPr();props.append(OxmlElement('w:cantSplit'))
            if i==0 and props.find(qn('w:tblHeader')) is None:props.append(OxmlElement('w:tblHeader'))
            for c,w in zip(row.cells,widths):
                c.width=Cm(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in c.paragraphs:
                    paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.first_line_indent=Cm(0)
                    paragraph.paragraph_format.space_before=paragraph.paragraph_format.space_after=Pt(4)
                    paragraph.paragraph_format.line_spacing=1.05
                    paragraph.paragraph_format.keep_with_next=i==0
                    for run in paragraph.runs:run.font.size=Pt(10);run.bold=i==0
        place(doc.add_paragraph()._p)
    def figure(path,caption):
        inputs[path]=digest(ROOT/path)
        x=p('',True);x.alignment=WD_ALIGN_PARAGRAPH.CENTER
        shape=x.add_run().add_picture(str(ROOT/path),width=Cm(16.4))
        shape._inline.docPr.set('descr',caption)
        y=p(caption);y.paragraph_format.first_line_indent=Cm(0)
        for run in y.runs:run.font.size=Pt(10)
    def mr(text):
        node=OxmlElement('m:r');value=OxmlElement('m:t');value.set(qn('xml:space'),'preserve');value.text=text;node.append(value);return node
    def sub(base,value):
        node=OxmlElement('m:sSub')
        for tag,text in [('m:e',base),('m:sub',value)]:
            part=OxmlElement(tag);part.append(mr(text));node.append(part)
        return node
    def frac(num,den):
        node=OxmlElement('m:f')
        for tag,text in [('m:num',num),('m:den',den)]:
            part=OxmlElement(tag);part.append(mr(text));node.append(part)
        return node
    def summation(expression):
        node=OxmlElement('m:nary');props=OxmlElement('m:naryPr');ch=OxmlElement('m:chr');ch.set(qn('m:val'),'∑');props.append(ch);node.append(props)
        for tag,text in [('m:sub','y = 0'),('m:sup','2'),('m:e',expression)]:
            part=OxmlElement(tag);part.append(mr(text));node.append(part)
        return node
    def equation(*parts):
        x=p('');x.alignment=WD_ALIGN_PARAGRAPH.CENTER;x.paragraph_format.first_line_indent=Cm(0)
        node=OxmlElement('m:oMath')
        for part in parts:node.append(mr(part) if isinstance(part,str) else part)
        x._p.append(node)

    page('20 共同推断规则下的公平比较')
    p('相同候选和数据预算不足以保证统计命题可比。第19节要求三模型的六个目标全部通过，不能从中挑选模型；'
      'LTT部署选择器可以从九个候选中选取，因此原实现对每个风险目标下的九候选作Bonferroni校正。'
      '首轮整组检验水平为0.025；共同LTT规则的每候选每目标水平为0.025/9。检验更容易通过不能直接解释为控制器风险更低。')
    p('为分开机制效果与推断规则，下表重现已完成原17路径研究的全部九候选：校准为1,024个独立初态，'
      '最终测试为另512个初态。全部使用原共同LTT水平，不重算更宽松证书，不按最终结果挑选种子。'
      '这张历史表不包含正在校准的50 ms门控。')
    p('风险上界依次为实际模式与四模式组合目标；最终有害数以最终实际采用数为分母。',True)
    labels={'raw':'原始学习','nominal':'名义门控','portfolio0':'17路径门控'}
    rows=[]
    for r in comparison['rows']:
        a=r['final_actual_adoption']
        rows.append([labels[r['controller']]+'\n'+r['model'],
            f"{100*r['calibration_actual']['threshold_upper']:.2f}% / {100*r['calibration_bundle']['threshold_upper']:.2f}%",
            f"{r['final_cost_gain_percent']:.3f}%",f"{a['adopted']}/512",f"{a['harmful_adopted']}/{a['adopted']}"])
    table(['方法与模型','校准风险上界','最终成本改善','最终采用','最终有害采用'],rows,[3.3,4.2,2.6,2.8,3.5])
    p('在共同规则下，只有17路径门控1812通过两个风险目标，三个完整模型组均未通过。'
      '名义门控保留更多收益与采用，但风险上界较高；严格门控降低经验退化，代价是收益和采用率下降。'
      '原始学习1802的最终平均收益区间还跨越零，不能预设所有学习模型在本分布都显著优于基线。')
    p('共形适配的边际覆盖涉及校准集与新样本共同随机性，不能自动推出固定校准集下采用后的条件风险。'
      'Mean SPI的全信息配对均值检验也不保证每车不退化或采用后风险低于5%。这些是已有工具在本接口的适配[6,7,9]，'
      '并非各原文完整系统复现。新门控的整组主命题与共同口径对照将分别报告。')

    page('21 保持完整判据的计算改进')
    p('反馈基线续跑占所分析慢调用的80.6%–92.5%，跨调用缓存仅降低约0.52%完整耗时。'
      '本文复用区间证书的四个3×2状态矩阵，保留原候选、排序、33条预测路径、完整时域、准入和回退，'
      '未启用fastmath或截止时间截断。')
    p('两实现从相同上下界开始，按相同运算次序生成下一上下界，并执行相同转换检查。'
      '缓冲交换只改变存储角色，逐步归纳可得证书值一致。下游辅助函数与门控方法体逐项语法树一致。'
      '这是实现等价论证；有限回放不替代全状态机器证明。')
    p('在一个已复核旧初态、三固定模型与四种扰动上，预先冻结并交替执行两实现，共24次SUMO。'
      '全部277次配对决策、9次学习采用及所有非计时字段完全一致；独立复核重建24条物理轨迹、554个候选集合和7个不同的已准入完整预测。')
    p('下表为完整select调用耗时之和，包括分配、生成、预测、物理检查与证据整理；初始化另存。',True)
    table(['扰动模式','配对决策数','参考耗时秒','缓冲耗时秒','耗时减少'],
        [[r['mode'],r['decisions'],f"{r['reference_select_ms']/1000:.3f}",f"{r['buffered_select_ms']/1000:.3f}",f"{r['reduction_percent']:.2f}%"]
         for r in timing['group_timings']],[3.0,2.8,3.5,3.5,3.6])
    p('本组完整耗时减少75.423%。样本P95由1,400.758降至335.154 ms，P99由2,631.581降至591.055 ms；'
      '最大值仍为838.140 ms，超过500 ms的调用由34/277降至8/277。该组仅一个旧初态，且存在并发桌面负载，'
      '不能替代完整开发集时延分布或硬实时保证。')
    figure('artifacts/v25/buffered_phase_presentation/paired_gate_timing.png',
        '图1 单旧初态上的配对计算耗时与经验时延分布 三模型和四模式不增加独立有效性样本')
    p('完整验证覆盖1,536个开发与1,920个迁移案例，两实现共计划6,912次SUMO。'
      '所有原失败均保留；相同包装器记录成功与终止异常调用耗时，补全原成功调用计时。'
      '截至本稿该验证仍在执行，实时性结论等待全量结果。正式风险校准继续使用原冻结实现。')

    page('22 采用后的风险迁移边界')
    p('物理误差界约束可恢复性，性能风险还取决于采用对象和逐车结果。定义三状态事件Y：'
      '未采用、采用且无定义退化、采用且发生退化。未采用时的失败保留在总体失败指标中。'
      '令P与Q为源条件和目标条件下Y的分布，a=P(A)>0，源风险为rP=P(H|A)，并假设Q(A)>0。'
      '如果事件分布的总变差不超过ρ，则标准总变差工具[15]给出以下敏感性关系：')
    equation(sub('d','TV'),'(',sub('P','Y'),', ',sub('Q','Y'),') = ',frac('1','2'),
        summation('|P(Y = y) − Q(Y = y)|'),' ≤ ρ')
    equation(sub('r','Q'),' ≤ min{1, ',sub('r','P'),' + ',frac('ρ','a'),'}， Q(A) ≥ a − ρ')
    p('为证明风险界，记B={Y=2}、C={Y=1}，构造值域长度为1的函数：')
    equation('f(Y) = (1 − ',sub('r','Q'),') ',sub('1','B'),' − ',sub('r','Q'),sub('1','C'))
    equation(sub('E','Q'),'f = 0， ',sub('E','P'),'f = (',sub('r','P'),' − ',sub('r','Q'),')a， |',
        sub('E','P'),'f − ',sub('E','Q'),'f| ≤ ρ')
    p('整理即得上述风险界；采用率下界来自同一总变差定义。这里ρ描述三状态事件的概率变化，'
      '不是加速度残差，不能用±0.5 m/s²直接替代。较低采用率会放大通用风险上界对事件分布变化的敏感性。')
    p('该比例项可以达到等号。合成源概率为(0.80,0.19,0.01)，目标为(0.80,0.18,0.02)时，'
      '采用率都为20%，事件TV为1%，而采用后风险从5%变为10%。该例只说明数学边界，不是本项目实测风险。'
      '若在相同外生随机性下耦合两套系统，其事件不一致概率为ε，则ρ≤ε；全域实现等价时ε为零。有限回放零差异不能证明总体ε为零。')
    p('若要高置信迁移，还需同时有效的源风险上界U、源采用率下界L和偏移上界R。三者同时覆盖且L>R时：')
    equation(sub('r','Q'),' ≤ min{1, U + ',frac('R','L'),'}， Q(A) ≥ L − R > 0')
    p('三项错误概率可由联合界相加，但各界必须在实际数据使用和选择规则下有效。当前20%采用率验收线是经验条件，'
      '不能直接当作真实采用率下界L。正式风险校准也未提供目标分布偏移的有效R，因此本推导不新增迁移风险证书。')

    page('匹配迁移的采用与退化事件',2)
    p('从已完整复核的12,160次匹配迁移运行中，重建六类非基线方法、三个模型、五类场景和三种目标动力学相对名义动力学的270个单元。'
      '每单元32个旧初态，均保留3×3事件转移表；540次源与目标计数核验同原汇总完全一致。'
      '这里的源条件是匹配迁移实验的名义动力学与±0.1外部创新，不是正式风险校准的四模式混合总体。')
    figure('artifacts/v25/matched_event_transfer_presentation/event_changes.png',
        '图2 50 ms完整模型组的45个配对事件变化单元 星号表示目标动力学的物理误差假设越界')
    p('在50 ms组中，增益阻力的15/15个场景模型单元出现配对事件变化，轻度惯性为14/15，强惯性为15/15。'
      '前两类动力学下本组运行均完成且误差未越界，目标采用后也未记录定义的有害事件；这不证明零总体风险。'
      '变化可包含旧退化消失及采用对象改变，并不等同于风险必然增加。强惯性下三个模型各160次运行均越界，未完成数仍为8、8、9。')
    p('低速让行的名义条件下，三个模型仅采用2/32、2/32、3/32次，其中各有一次有害。'
      '基于经验事件分布计算的通用界在45个单元中的11个达到100%上限。小采用分母限制了通用迁移界的解释力；'
      '采用和退化事件的经验TV也不是未知总体TV的有效上界。全部270个单元与失败计数均保留[16]，本分析不增加独立样本或风险认证。')

    doc.add_paragraph('[15] Levin D A, Peres Y. Markov Chains and Mixing Times. Second edition. '
        'American Mathematical Society. Chapter 4. https://darkwing.uoregon.edu/~dlevin/MARKOV/mcmt2e.pdf')
    doc.add_paragraph('[16] CoopGate. Common-inference audit, exact buffered implementation and matched event-transfer evidence. '
        '2026. https://github.com/Leo-7GOAT/CoopGate/tree/main/artifacts/v25')
    math=[tostring(e,method='c14n',exclusive=True) for e in doc.element.body.iter(qn('m:oMath'))]
    assert math[:23]==original_math and len(math)==28
    assert doc.sections[0]._sectPr.xml==original_section
    out.parent.mkdir(parents=True,exist_ok=True);doc.save(out)
    assert digest(reference)==REFERENCE_SHA
    with ZipFile(reference) as a,ZipFile(out) as b:
        assert set(a.namelist())<=set(b.namelist())
        changed=[name for name in a.namelist() if a.read(name)!=b.read(name)]
        added=sorted(set(b.namelist())-set(a.namelist()))
        assert set(changed)<={'word/document.xml','docProps/core.xml','[Content_Types].xml','word/_rels/document.xml.rels'},changed
        assert len(added)==2 and all(name.startswith('word/media/') for name in added),added
    record=dict(document_name=out.name,document_sha256=digest(out),reference_sha256=REFERENCE_SHA,
        source_sha256=digest(__file__),input_sha256=inputs,changed_parts=changed,added_parts=added,
        retained_native_equations=23,added_native_equations=5,total_tables=len(doc.tables),
        completed_paired_sumo_runs=24,paired_decisions=277,matched_event_cells=270,
        whole_scope_sumo_runs_planned=6912,whole_scope_results_read=False,
        active_calibration_outcomes_read=False,new_risk_certificate=False,hard_real_time_certified=False,
        render_required=True)
    evidence_out.parent.mkdir(parents=True,exist_ok=True)
    evidence_out.write_text(json.dumps(record,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(record,ensure_ascii=False))


if __name__=='__main__':
    parser=argparse.ArgumentParser()
    parser.add_argument('--reference',type=Path,required=True)
    parser.add_argument('--out',type=Path,required=True)
    parser.add_argument('--evidence-out',type=Path,required=True)
    args=parser.parse_args();build(args.reference,args.out,args.evidence_out)
