"""Update the retained Word manuscript from completed V21 evidence only.

Run with the bundled document runtime. Never reads active calibration records.
The reference and all study sources are preserved; the output is a new DOCX.
"""
import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.table import Table
from lxml.etree import tostring

ROOT = Path(__file__).resolve().parents[1]
REFERENCE_SHA = '7d30395f64dcf669845a3783fd2f8f6115489fc2d35bb9a7f25316a166eb50c7'
EXECUTION_SHA = '277b6bf4d8dda9540b4a9aea7d396887817be797e3e64bd4cefc41e5730457b2'
SEEDS = ('1802', '1812', '1822')


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def build(reference, out, evidence_out):
    if reference.resolve() == out.resolve() or digest(reference) != REFERENCE_SHA:
        raise ValueError('Use the unchanged retained fifteenth manuscript')
    sources = {}

    def read(relative):
        sources[relative] = digest(ROOT / relative)
        return json.loads((ROOT / relative).read_text(encoding='utf-8'))

    completed = 'artifacts/v21/interrupted_recovery_20261022'
    phase = read(completed + '/results/phase/summary.json')
    transfer = read(completed + '/results/transfer/summary.json')
    completion = read(completed + '/execution/supervisor_exit.json')
    selection = read('artifacts/v21/validation_attempts/attempt_1/selection.json')
    plan = read('artifacts/v21/prospective_execution_attempt_1_native/frozen/plan.json')
    preflight = read('artifacts/v21/prospective_execution_preflight_native_20261022/preflight/analysis/acceptance.json')
    preflight_exit = read('artifacts/v21/prospective_execution_preflight_native_20261022/preflight/execution/supervisor_exit.json')
    assert completion['exit_code'] == preflight_exit['exit_code'] == 0
    assert preflight['integration_only'] and preflight['new_independent_initials'] == 0
    assert preflight['all_runtime_offline_profile_choices_exact'] == 18
    assert phase['actual_sumo_runs'] == 6656 and phase['independent_initials'] == 128
    assert transfer['actual_sumo_runs'] == 12160 and transfer['independent_initials'] == 160
    assert len(transfer['cells']) == 380
    assert not phase['risk_certified'] and not transfer['risk_certified']
    assert selection['selected_family'] == 'phase05' and not selection['risk_certified']
    assert selection['development_summary_sha256'] == sources[completed + '/results/phase/summary.json']
    assert sources['artifacts/v21/prospective_execution_attempt_1_native/frozen/plan.json'] == EXECUTION_SHA
    assert plan['allocated_error'] == .025 and plan['prospective_total_sumo_runs'] == 296960
    for name, observed in phase['independently_recounted_outcomes'].items():
        for key, value in observed.items():
            assert phase['results'][name][key] == value, (name, key)

    doc = Document(reference)
    original_section = doc.sections[0]._sectPr.xml
    original_math = [tostring(e, method='c14n', exclusive=True)
                     for e in doc.element.body.iter(qn('m:oMath'))]
    assert len(original_math) == 17
    prototype = deepcopy(doc.tables[0]._tbl)
    doc.core_properties.version = '16'

    def replace(prefix, content):
        matches = [p for p in doc.paragraphs if p.text.startswith(prefix)]
        if len(matches) != 1:
            raise ValueError('Ambiguous edit slot: ' + prefix)
        matches[0].text = content
        return matches[0]

    replace('本文研究协同车辆在学习决策与传统规划之间的逐周期选择。',
            '本文研究协同车辆在学习决策与传统规划之间的逐周期选择。多车分别认可的动作同时执行后可能相互干扰，'
            '团队均值也可能掩盖个别车辆延误。本文对实际执行的完整联合动作检查逐车成本与到达时间，'
            '并引入有限扰动路径、动作调整和插值到达时序检查。基线剩余价值的望远镜求和给出精确模型和库内真实扰动下的性能关系；'
            '物理可恢复性由独立过滤器约束，采用后的统计退化风险使用独立数据检验。')
    replace('单车前置验证与三车模型式协同训练完成后，',
            '单车前置验证和三车模型式协同训练完成后，原17路径门控仅有一份固定模型通过全部预定风险与收益要求。'
            '新增128个独立初态的6,656次开发执行中，50 ms时序门控保留0.76%–1.03%的成本改善和28.1%–30.5%的实际采用率，'
            '密集交互收益高于普通场景；相对原动作调整门控，收益与采用率均有损失。另12,160次匹配外部噪声的迁移执行暴露低速让行退化及强惯性失败。'
            '完整决策P95为1.55–1.69秒，仍不满足500 ms周期。新一轮4,096初态校准正在执行，跨模型风险保证尚未获得。')
    replace('本文考察完整联合组合对实际采用、逐车性能和计算代价的影响，',
            '本文考察完整联合组合、动作调整和时序检查对实际采用、逐车性能和计算代价的影响，并检验有限扰动保证及动力学迁移边界。'
            '第11至15节保留原17路径门控的已完成研究；第16至19节给出新增机制、完整开发与匹配迁移结果，以及独立风险检验设计。'
            '原研究仅一份模型通过全部要求，新门控仍须独立校准，两个阶段的证据不能互相替代。')
    replace('共形预测与Learn then Test为不确定性和风险检验提供了工具',
            '共形预测与Learn then Test为不确定性和风险检验提供工具[6,7]；MA-COPP研究多智能体策略改变后的联合轨迹分布[8]。'
            '阈值型安全策略改进与深度模型式改进已有相关研究[9,10]。第12节报告原17路径门控的匹配适配结果；第19节为新门控沿用相同候选与样本预算。'
            '完整模型组检验使用已有交并检验[13]，不将统计工具本身作为方法创新。')
    replace('本文的创新性应由组合机制的实际影响、有效采用收益和复核证据评价，',
            '本研究的贡献集中于实际执行联合组合的逐车准入、动作调整与到达时序约束，以及可复核的收益、退化和计算代价。'
            '完整配对实验用于区分机制改善与简单增加保守性。现有结果支持有条件的学习采用，但尚未证明跨模型一致认证、'
            '对公开完整系统的普遍优势或实时部署能力。')
    replace('对三个actor的学习提案与同一反馈基线动作作二元组合，',
            '本节及第12至15节描述原17路径门控研究。对三个actor的学习提案与同一反馈基线动作作二元组合，至多产生七个非基线候选。'
            '每个完整组合先经过共同物理过滤器，再检查全部车辆的成本与时间。扰动库包含一条无扰动、八条恒定异号偏置和八条固定符号脉冲。'
            '候选按名义团队收益排序，全部17条预测通过才采用；第16节另行扩展候选与预测库。')
    replace('两个分母不能互换，',
            '两个分母不能互换，压力组风险不能自动解释成每种模式的采用后风险。原17路径研究固定九个候选、两个目标，风险上限为5%，'
            '总错误认证预算为5%。每个候选与目标分配1/360预算，通过精确二项检验构造Clopper–Pearson单侧上界；这是Learn then Test的应用[7]。')
    replace('冻结权重与门控参数后，在名义、密集、低速让行、速度不对称、紧密跟车五类初态中，',
            '原17路径研究冻结权重与门控参数后，在名义、密集、低速让行、速度不对称、紧密跟车五类初态中，各取32个新初态，'
            '配对四类动力学与十个控制器，完成6,400次SUMO。四类动力学为名义模型、增益与二次阻力、轻度惯性、强惯性。'
            '各控制器使用相同隐藏参数及外生创新；跨动力学的噪声幅值差异由第18节的新实验控制。')
    replace('16 讨论与结论', '20 讨论与结论')
    replace('独立最终结果支持联合检查在本任务中的实际收益。',
            '原17路径门控的独立最终结果支持联合检查在本任务中的收益，但仅模型1812通过全部预定要求。'
            '新增时序门控开发中，三模型相对强基线保留0.76%–1.03%收益和28.1%–30.5%采用率；'
            '密集交互改善约1.21%–1.58%，普通场景约0.37%–0.55%。这些观察支持继续验证，不能替代独立最终结果。')
    replace('可恢复物理安全依赖误差与初始集合假设；',
            '可恢复物理安全依赖误差和初始集合假设；逐车性能关系依赖精确模型、有限扰动路径与终端条件；采用后风险依赖独立校准和声明的总体。'
            '时序检查减少了一些采用，也降低平均收益，经验条件风险没有在所有模型上一致下降。四模式压力组仍存在退化，'
            '不能从实际模式零次观察退化推断零风险。')
    replace('保守共形适配在最终测试全部回退基线；',
            '原研究的保守共形适配全部回退，LTT只在一份模型采用严格门控，均值改进适配保留收益但不保证采用后条件风险。'
            '新研究继续使用共同候选、校准初态和实际部署回放，按各方法的保证对象解释结果。'
            '已有工具的适配不等于原论文完整系统复现，也不能用某个比较器的成功替代失败的主风险命题。')
    replace('本文将统计校准、策略改进恒等式和模型式多智能体训练作为已有工具。',
            '匹配外部噪声的迁移研究显示，时序门控在名义、增益阻力和轻度惯性下均完成本组运行；'
            '低速让行仍有负收益，强惯性超出误差界并产生未完成。完整调用P95仍超过控制周期。'
            '后续工作需完成冻结的独立校准和最终比较，并在保持候选与判据的前提下降低计算开销；'
            '若改变动力学范围，则必须重新建模执行器和可恢复集合。')
    replace('代码、冻结模型、实验协议、逐回合记录和复核入口见公开仓库',
            '代码、冻结模型、实验协议、逐回合记录和独立复核入口见公开仓库[11,12,14]。'
            '当前尚无新增门控的跨模型风险认证、硬实时保证或开放道路安全承诺。')

    discussion = next(p for p in doc.paragraphs if p.text == '20 讨论与结论')
    anchor = discussion._p.getprevious()

    def place(element):
        anchor.addprevious(element)

    def p(text, caption=False):
        paragraph = doc.add_paragraph(text)
        paragraph.paragraph_format.first_line_indent = Cm(0 if caption else .74)
        if caption:
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_after = Pt(5)
        place(paragraph._p)
        return paragraph

    def page(title, level=1):
        place(doc.add_page_break()._p)
        place(doc.add_heading(title, level)._p)

    def table(headers, rows, widths):
        element = deepcopy(prototype)
        place(element)
        t = Table(element, doc._body)
        while len(t.rows) > 1:
            t._tbl.remove(t.rows[-1]._tr)
        while len(t.columns) < len(headers):
            t.add_column(Cm(2.5))
        assert len(t.columns) == len(headers) and abs(sum(widths) - 16.4) < 1e-9
        t.autofit = False
        for cell, value in zip(t.rows[0].cells, headers):
            cell.text = value
        fill = deepcopy(t.rows[0].cells[0]._tc.get_or_add_tcPr().find(qn('w:shd')))
        for cell in t.rows[0].cells:
            props = cell._tc.get_or_add_tcPr()
            old = props.find(qn('w:shd'))
            if old is not None:
                props.remove(old)
            if fill is not None:
                props.append(deepcopy(fill))
        for values in rows:
            assert len(values) == len(headers)
            for cell, value in zip(t.add_row().cells, values):
                cell.text = str(value)
        for col, width in zip(t.columns, widths):
            col.width = Cm(width)
        for index, row in enumerate(t.rows):
            pr = row._tr.get_or_add_trPr()
            pr.append(OxmlElement('w:cantSplit'))
            if index == 0 and pr.find(qn('w:tblHeader')) is None:
                pr.append(OxmlElement('w:tblHeader'))
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for para in cell.paragraphs:
                    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    para.paragraph_format.first_line_indent = Cm(0)
                    para.paragraph_format.space_before = para.paragraph_format.space_after = Pt(4)
                    para.paragraph_format.line_spacing = 1.05
                    para.paragraph_format.keep_with_next = index == 0
                    for run in para.runs:
                        run.font.size = Pt(10)
                        run.bold = index == 0
        blank = doc.add_paragraph()
        blank.paragraph_format.space_after = Pt(2)
        place(blank._p)

    def mr(text):
        r = OxmlElement('m:r')
        t = OxmlElement('m:t')
        t.set(qn('xml:space'), 'preserve')
        t.text = text
        r.append(t)
        return r

    def sub(base, value):
        node = OxmlElement('m:sSub')
        for tag, content in [('m:e', base), ('m:sub', value)]:
            part = OxmlElement(tag)
            part.append(mr(content) if isinstance(content, str) else content)
            node.append(part)
        return node

    def tilde(base):
        node = OxmlElement('m:acc')
        props = OxmlElement('m:accPr')
        character = OxmlElement('m:chr')
        character.set(qn('m:val'), '\u0303')
        props.append(character)
        node.append(props)
        expression = OxmlElement('m:e')
        expression.append(mr(base))
        node.append(expression)
        return node

    def fraction(num, den):
        node = OxmlElement('m:f')
        for tag, content in [('m:num', num), ('m:den', den)]:
            part = OxmlElement(tag)
            part.append(mr(content))
            node.append(part)
        return node

    def equation(*parts):
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.first_line_indent = Cm(0)
        math = OxmlElement('m:oMath')
        for part in parts:
            math.append(mr(part) if isinstance(part, str) else part)
        para._p.append(math)
        place(para._p)

    page('16 动作调整与到达时序门控')
    p('原二元组合只在每辆车的学习提案与基线动作之间选择。动作调整扩展至围绕提案与基线的有限候选集合，'
      '经共同物理过滤后对完整联合命令去重、排序和检查。预测库扩展为33条固定路径，包含名义、恒定偏置及时变路径。'
      '训练权重、共同强基线和原逐车成本及登记到达时间条件保持不变。新增候选提供更多采用机会，也增加预测开销。')
    p('登记到达时间按0.5秒决策周期累计。同一登记时间内仍可能有先后差异，原判据因此不能区分部分小幅延误。'
      '本文在既有0.1秒物理步长的预测轨迹上增加终点穿越时刻的线性插值诊断。设终点为L，'
      '某物理步由x⁻移动至x⁺且x⁻<L≤x⁺，步起点时刻为t，则定义：')
    equation(tilde('T'), ' = t + Δt ', fraction('L − x⁻', 'x⁺ − x⁻'), '， Δt = 0.1 s')
    p('该量由原离散动力学的相邻位置计算，不改变仿真更新、原成本或统计终点，也不等价于精确连续物理到达时间。'
      '对每个候选a、车辆i和库中路径ω，在保留全部原准入条件的基础上，要求插值到达差满足：')
    equation(sub(tilde('T'), 'i,ω'), '(a → b) − ', sub(tilde('T'), 'i,ω'), '(b) ≤ η， η ∈ {0.05, 0.10} s')
    p('a→b表示先执行候选命令，再持续执行同一反馈基线。当前已经到达的车辆与预测时域内未到达的情况按冻结实现处理；'
      '未获得有效完整预测的候选不能准入。50 ms与100 ms表示时序容差，不表示计算延迟、状态置信度或安全概率。')
    p('额外限制缩小同一状态下的可接受候选集合，但闭环轨迹也随之变化，不能由此推导最终条件风险必然下降。'
      '第11节的望远镜关系仍只作用于满足其假设的原成本和登记时间；不能把单步插值容差直接解释为整个行程的50 ms退化上界。'
      '真正的取舍需比较相同初态、相同外生扰动下的完整闭环结果。')

    page('17 完整开发比较与收益代价')
    p('在128个新独立初态上，四类扰动与13个控制器共完成6,656次SUMO，全部物理轨迹和准入证据通过独立重建。'
      '其中普通初态67个、密集初态61个。按冻结随机流每个初态选择一种实际模式作主统计，四模式压力组另行计数。'
      '三个固定模型始终作为完整组选择，不跨模型拼接最优门控。')
    p('表中采用列分母均为128；实际退化与组合退化均使用各自的采用数作为分母。成本改善相对配对强基线，正值表示收益。', True)
    rows = []
    for family, label in [('repair', '动作调整'), ('phase10', '100 ms'), ('phase05', '50 ms')]:
        for seed in SEEDS:
            d = phase['results'][family + '_' + seed]
            rows.append([label + '\n' + seed, f"{d['actual_mean_gain_percent']:.3f}%",
                         f"{d['actual_adopted']}/128", f"{d['actual_harmful_adopted']}/{d['actual_adopted']}",
                         f"{d['bundle_harmful_adopted']}/{d['bundle_adopted']}"])
    table(['门控与模型', '成本改善', '实际采用', '实际退化', '组合退化'], rows, [3.4, 3.1, 3.3, 3.3, 3.3])
    p('三个完整组均通过冻结开发筛选。按六个风险目标中最差开发上界分数排序，选中50 ms组进入新校准。'
      '其实际采用率为28.1%–30.5%，三个模型的平均收益逐项95%配对区间均为正；这些是开发区间，不能作为选择后的最终确认。'
      '实际模式零次观察退化并不等于零风险，四模式压力组仍各有一个有害初态。')
    p('相对动作调整门控，50 ms组的平均增量收益依次为−0.02820、−0.00366、−0.00378秒，三个配对区间均为负。'
      '模型1802的有害计数有所减少；另两模型的组合有害数仍为一、采用分母却更小，经验条件风险没有一致降低。'
      '该结果不支持全面性能提升，支持的是保留收益后的进一步独立风险检验。')
    p('50 ms组在普通场景改善0.37%–0.55%，密集场景改善1.21%–1.58%；这一差异只描述当前分布。'
      '完整决策P95为1,554–1,687 ms，超过500 ms周期。采集存在并发负载，不能将不同阶段计时直接解释成严格加速比。')

    page('18 匹配外部扰动的动力学迁移')
    p('第13节的不同动力学使用不同噪声幅值，无法单独归因于动力学变化。本节为五类场景各生成32个独立初态，'
      '在四种动力学、19个控制器间保持同一外部创新序列及±0.1 m/s²幅值，共完成12,160次SUMO。'
      '控制器仍采用原±0.5 m/s²误差假设，未重新训练或校准；全部380个实验单元保留。')
    p('下表按五类场景汇总，每个模型及基线各160次；三项计数依次对应模型1802、1812、1822。', True)
    plants = [('nominal', '名义'), ('gain_drag', '增益阻力'), ('lag_mild', '轻度惯性'), ('lag_strong', '强惯性')]
    transfer_rows = []
    for plant, label in plants:
        base = [c for c in transfer['cells'] if c['plant'] == plant and c['method'] == 'baseline']
        incomplete, exceed = [], []
        for seed in SEEDS:
            cells = [c for c in transfer['cells'] if c['plant'] == plant and c['method'] == 'phase05_' + seed]
            assert len(cells) == 5
            incomplete.append(str(sum(c['independent_initials'] - c['completed'] for c in cells)))
            exceed.append(str(sum(c['exceeded_error_runs'] for c in cells)))
        transfer_rows.append([label, ' / '.join(incomplete),
                              str(sum(c['independent_initials'] - c['completed'] for c in base)), ' / '.join(exceed)])
    table(['动力学', '50 ms组未完成', '基线未完成', '50 ms组误差越界'], transfer_rows, [3.0, 4.5, 3.4, 5.5])
    p('50 ms组在前三种动力学均完成本组运行，最大实测加速度误差分别约0.1000、0.3872、0.4363 m/s²。'
      '完成与误差未越界不构成迁移风险认证。强惯性下该组三模型分别8、8、9次未完成，基线10次未完成；'
      '各模型160次运行全部出现误差越界，原物理与性能假设已不适用。')
    p('全部控制器在强惯性下共3,040次执行，其中146次未完成，含144次中止；其余三类动力学另有5次未完成，均来自原始学习，'
      '全研究合计151次。没有记录碰撞或约束事件，但中止之后未继续模拟，不能解释为车辆已安全停稳。'
      '基线失败时较大的相对改善也不能单独支持可靠性结论。')
    p('名义动力学下各场景的50 ms组成本改善如下，单位为百分比；每格只有32个独立初态。', True)
    scene_labels = [('nominal', '普通'), ('dense', '密集'), ('slow_yield', '低速让行'),
                    ('asymmetric_speed', '速度不对称'), ('tight_headway', '紧密跟车')]
    scene_rows = []
    for scene, label in scene_labels:
        row = [label]
        for seed in SEEDS:
            c = next(c for c in transfer['cells'] if c['scenario'] == scene and c['plant'] == 'nominal'
                     and c['method'] == 'phase05_' + seed)
            row.append(f"{c['mean_gain_percent']:.4f}")
        scene_rows.append(row)
    table(['场景', '模型1802', '模型1812', '模型1822'], scene_rows, [4.4, 4.0, 4.0, 4.0])
    p('低速让行三个模型的有害采用分别为1/2、1/2、1/3，前两个模型平均收益为负。该小样本反例限制了推广范围，'
      '不能将密集开发场景的收益外推为任意复杂交通都有效。全部逐格结果及配对动力学、门控差异见数据记录[14]。')

    page('19 完整模型组的独立风险检验')
    p('新一轮只对开发阶段选定的50 ms完整组提出主命题：三份固定模型的实际模式风险与四模式压力组风险全部不超过5%。'
      '六个目标使用各自采用事件A及有害事件H，定义r=P(H=1|A=1)。给定独立初态中的采用数N，'
      '有害采用数K服从二项分布。车辆、模式、模型与回放不能增加独立初态数，六个目标之间不要求独立。')
    p('对每个目标j检验风险达到或超过α=0.05的零假设，使用精确下尾概率；N为零时返回不通过。')
    equation(sub('p', 'j'), ' = P{Binomial(', sub('N', 'j'), ', α) ≤ ', sub('K', 'j'), '}， α = 0.05')
    p('采用交并检验[13]，六项全部通过才认证完整组。若组命题不成立，总有一个由真实总体确定的目标j*风险不合格；'
      '完整组通过事件包含于该单项目标通过事件，因此首轮错误认证概率不超过预定δ。')
    equation(sub('p', 'group'), ' = maxⱼ ', sub('p', 'j'), ' ≤ δ， P{错误认证整组} ≤ P{', sub('p', 'j*'), ' ≤ δ} ≤ δ')
    p('这里报告不可拆分的整组命题，不在组内挑选通过的模型，因而无需为六项另行分摊δ。'
      '另取每项置信水平1−δ的Clopper–Pearson单侧上界U，并令Umax为六项最大值。设真实最大风险为θ，'
      '选择真实最大风险目标j*，可得：')
    equation('θ = maxⱼ ', sub('r', 'j'), '， ', sub('U', 'max'), ' = maxⱼ ', sub('U', 'j'),
             '， P{', sub('U', 'max'), ' < θ} ≤ P{', sub('U', 'j*'), ' < ', sub('r', 'j*'), '} ≤ δ')
    p('因此Umax是最大风险的有效上界；各单项U仍是逐项上界，不能宣称它们全部同时覆盖各自真实风险。'
      '代码同时要求最大p值和最大上界通过阈值，数值边界只能使结果更保守。该统计结构不降低控制器真实风险，'
      '也不认证收益、采用率、任意状态安全、新模型或迁移分布。')
    p('为防止反复更换控制器并校准直至偶然通过，第k轮预分配可求和的错误预算：')
    equation(sub('δ', 'k'), ' = ', fraction('0.05', 'k(k + 1)'), '， k = 1, 2, …， ∑ₖ ', sub('δ', 'k'), ' = 0.05')
    p('首轮δ为0.025。给定过去历史，每轮使用新独立样本且事先固定控制器和样本量，条件错误认证概率至多δk；'
      '再取期望并对轮次使用联合界，任意一轮错误认证的总概率不超过0.05。接触数据后失败或中止仍消耗该轮预算，'
      '不得重置轮次、筛掉模型或按效果追加样本。样本分割、交并检验与联合界均为已有工具。')

    page('公平对照与实际执行协议', 2)
    p('新校准固定4,096个独立初态，最终确认固定另1,024个初态，普通与密集总体混合方式不变。'
      '九个共同候选为三份模型各自的原始学习、名义门控及50 ms门控，另含一个共同基线。'
      '原始学习、名义门控、所提门控、LTT、Conformal和Mean SPI共六类部署规则共享这些候选与全部校准初态。')
    p('完整执行预算区分数据角色，影子预测和部署回放均不增加独立样本数。', True)
    table(['阶段', '独立初态', '主执行', '观测影子预测', '实际部署回放'],
          [['风险校准', '4,096', '163,840', '40,960', '0'],
           ['最终确认', '1,024', '40,960', '10,240', '40,960']], [3.2, 3.0, 3.2, 3.6, 3.4])
    p('每个初态先从可观测状态产生十个完整名义预测，固定内部随机流不读取真实扰动模式、种子或未来。'
      '最终阶段在实际轨迹产生前保存18个方法与模型组合的选择，随后执行共同控制器网格和独立完整部署回放。'
      '相同控制器被多个方法选中时共享同一实际回放；命令、状态和非计时指标逐一核对，不能按结果拼接轨迹。')
    p('LTT提供采用后条件退化检验，Conformal校准初态簇内最坏残差的边际覆盖，Mean SPI使用全信息配对均值检验。'
      '三者保证对象不同，只能称为当前共同接口下的适配。Conformal所需全部观测预测准备成本另行计入，'
      '无需该预测的方法不承担此成本；记录完整控制调用和初态总耗时，不统计通信延迟。')
    p('完整执行器已在一个旧初态上通过90次实际SUMO集成，含40次主执行、10次影子预测和40次部署回放，'
      '18个运行时选择均与离线接口一致。该集成不增加新独立样本。2026年10月22日冻结的正式协议正在执行新校准；'
      '最终确认尚未开始，新增风险证书尚未产生。主进程实际退出、完整记录封存和独立复核完成后才允许分析。')
    p('最终确认要求主风险认证有效，且三份模型各自实测采用率至少20%、平均成本改善至少0.5%、配对收益区间下界为正。'
      '完成率与计时另外完整报告。即使主认证失败，仍保留失败命题并完成共同方法比较，其他方法的成功不能补救主认证。'
      '原始基础设施失败记录和已提交前缀保留，禁止通过静默重跑或结果筛选改变样本。')

    doc.add_paragraph('[13] Berger R L, Hsu J C. Bioequivalence Trials, Intersection–Union Tests and Equivalence Confidence Sets. '
                      'Statistical Science, 1996, 11(4): 283–319. https://doi.org/10.1214/ss/1032280304')
    doc.add_paragraph('[14] CoopGate. Complete arrival-phase development, matched-disturbance transfer and prospective execution protocol. '
                      '2026. https://github.com/Leo-7GOAT/CoopGate/tree/main/artifacts/v21')

    all_math = [tostring(e, method='c14n', exclusive=True) for e in doc.element.body.iter(qn('m:oMath'))]
    assert all_math[:17] == original_math and len(all_math) == 23
    assert doc.sections[0]._sectPr.xml == original_section
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    assert digest(reference) == REFERENCE_SHA
    with ZipFile(reference) as old, ZipFile(out) as new:
        assert set(old.namelist()) == set(new.namelist())
        changed = [name for name in old.namelist() if old.read(name) != new.read(name)]
        assert set(changed) <= {'word/document.xml', 'docProps/core.xml'}, changed
    record = dict(document_name=out.name, document_sha256=digest(out), reference_sha256=REFERENCE_SHA,
                  changed_parts=changed, retained_native_equations=17, added_native_equations=6,
                  source_sha256=digest(__file__), input_sha256=sources, phase_runs=6656,
                  phase_independent_initials=128, matched_transfer_runs=12160, transfer_independent_initials=160,
                  transfer_cells=380, selected_complete_family='phase05',
                  calibration_results_read=False, new_risk_certified=False, hard_real_time_certified=False,
                  render_required=True, expected_pages='approximately 24; inspect every final page')
    evidence_out.parent.mkdir(parents=True, exist_ok=True)
    evidence_out.write_text(json.dumps(record, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(record, ensure_ascii=False))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--reference', type=Path, required=True)
    parser.add_argument('--out', type=Path, required=True)
    parser.add_argument('--evidence-out', type=Path, required=True)
    args = parser.parse_args()
    build(args.reference, args.out, args.evidence_out)
