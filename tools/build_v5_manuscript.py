"""Chinese paper revision using verified V5 results; earlier drafts are retained."""
import argparse
import json
from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
D = ROOT/'artifacts/v5'


def build(out):
    evidence = json.loads((D/'summary.json').read_text())
    summary = evidence['summary']; timing = evidence['timing']['full_portfolio']
    doc = Document(); section = doc.sections[0]
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    section.page_width = Cm(21); section.page_height = Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.1)
    section.left_margin = section.right_margin = Cm(2.3)
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption'):
        style = doc.styles[name]
        style.font.name = 'Times New Roman'; style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
    normal = doc.styles['Normal']
    normal.font.size = Pt(11); normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.space_after = Pt(6)
    doc.styles['Title'].font.size = Pt(18)
    doc.styles['Heading 1'].font.size = Pt(14)
    doc.styles['Subtitle'].font.size = Pt(11); doc.styles['Subtitle'].font.italic = False
    doc.core_properties.author = '李东旭'
    doc.core_properties.title = '联合可恢复证书下的合流规划与学习性能分析'

    def p(text, style=None):
        paragraph = doc.add_paragraph(style=style)
        pos = 0
        for token in re.finditer(r'([A-Za-z]+)_([A-Za-z]+)', text):
            paragraph.add_run(text[pos:token.start()])
            paragraph.add_run(token.group(1))
            paragraph.add_run(token.group(2)).font.subscript = True
            pos = token.end()
        paragraph.add_run(text[pos:])
        if style is None:
            paragraph.paragraph_format.first_line_indent = Cm(.74)
        return paragraph

    def h(text):
        doc.add_heading(text, 1)

    def eq(text):
        paragraph = doc.add_paragraph(); paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        math = OxmlElement('m:oMath')
        def run(parent, value):
            r = OxmlElement('m:r'); t = OxmlElement('m:t'); t.text = value
            t.set(qn('xml:space'), 'preserve'); r.append(t); parent.append(r)
        pos = 0
        pattern = r'([A-Za-zΣ]+)(?:_(?:\{([^}]+)\}|([A-Za-z0-9]+)))?(?:\^(?:\{([^}]+)\}|([A-Za-z0-9*]+)))?'
        for token in re.finditer(pattern, text):
            if token.start() > pos:
                run(math, text[pos:token.start()])
            low, high = token.group(2) or token.group(3), token.group(4) or token.group(5)
            if low is None and high is None:
                run(math, token.group(1))
            else:
                script = OxmlElement('m:sSubSup' if low and high else ('m:sSub' if low else 'm:sSup'))
                base = OxmlElement('m:e'); run(base, token.group(1)); script.append(base)
                for tag, value in [('m:sub', low), ('m:sup', high)]:
                    if value:
                        index = OxmlElement(tag); run(index, value); script.append(index)
                math.append(script)
            pos = token.end()
        if pos < len(text):
            run(math, text[pos:])
        paragraph._p.append(math)

    def table(headers, rows, widths):
        table = doc.add_table(rows=1, cols=len(headers)); table.autofit = False
        for column, width in zip(table.columns, widths):
            column.width = Cm(width)
        for cell, width, value in zip(table.rows[0].cells, widths, headers):
            cell.width = Cm(width); cell.text = value
        table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for values in rows:
            for cell, width, value in zip(table.add_row().cells, widths, values):
                cell.width = Cm(width); cell.text = str(value)
        borders = OxmlElement('w:tblBorders')
        for name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
            edge = OxmlElement('w:'+name); edge.set(qn('w:val'), 'single')
            edge.set(qn('w:sz'), '4'); edge.set(qn('w:color'), 'D9D9D9'); borders.append(edge)
        table._tbl.tblPr.append(borders)
        for i, row in enumerate(table.rows):
            row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E7ECF2' if i == 0 else 'FFFFFF')
                cell._tc.get_or_add_tcPr().append(shade)
                for paragraph in cell.paragraphs:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    paragraph.paragraph_format.space_after = paragraph.paragraph_format.space_before = Pt(5)
                    paragraph.paragraph_format.line_spacing = 1.05
                    paragraph.paragraph_format.keep_with_next = i < len(table.rows)-1
                    for run in paragraph.runs:
                        run.font.size = Pt(10); run.bold = i == 0
        doc.add_paragraph().paragraph_format.space_after = Pt(2)

    p('联合可恢复证书下的\n合流规划与学习性能分析', 'Title')
    p('李东旭', 'Subtitle')
    h('摘要')
    p('针对强化学习与传统规控切换中安全收益和学习收益难以区分的问题，本文研究三辆完全受控车的双路合流。'
      '在包含上一加速度指令的状态空间中，采用候选首步与联合制动尾部构成的可恢复证书，并检查子步内部的冲突区占用。'
      '在该固定安全合同下，构建包含传统基线、优先权和提前滑行方案的37项候选库，以完整到达时间选择可执行轨迹，运行时逐步重检并在计划失配时回退。')
    p('对80个独立新场景的五种方法共完成400次真实SUMO评估。完整规划使常规及高密度场景平均到达时间分别降低7.292%和7.290%，'
      '80个场景中64个改善、16个持平，所有执行均完成且未记录碰撞或路线合同违规。冻结的强化学习策略仍与传统基线持平。'
      '结果将收益定位于提前让行与完整目标比较，为改进学习器提供可达到的性能参照；结论限于固定路线、精确模型和全受控交通。')
    p('关键词  合流协调  多智能体强化学习  可恢复性  控制器切换  轨迹规划')
    h('1 引言与相关研究')
    p('学习控制器的动作可行性与任务效率属于不同问题。安全门控可以避免不满足合同的动作，却可能让策略只学会持续加速，由门控承担实际让行。'
      '前期960次合流执行中，三个训练种子的联合门控策略均未超过恒定加速加同一门控。由此产生本文问题：当前安全合同下是否还存在可实现、可学习的性能改善？')
    p('多智能体rollout已有相对基策略的成本改进性质[1]；DMPS通过恢复规划结合短期进展与长期奖励，处理安全屏障妨碍任务推进的问题[2]。'
      '轨迹驱动在线策略迭代也已研究固定初态和一致性条件下的单调改进[3]。本文不将这些一般思想作为首次创新，重点检验同一安全合同下的实际性能来源。')
    p('本文提供可运行的有限候选规划器、条件性安全与基线不劣论证，以及完整时域和提前让行的配对消融。'
      '本轮新增收益来自确定性规划，尚不构成新的强化学习效率结果。')

    doc.add_page_break()
    h('2 动力学与合流安全合同')
    p('系统包含三辆车，来路标签为(0,1,0)，同来路车辆0始终在车辆2之前。每辆车的状态包含路线进度s、速度v及上一指令u⁻。'
      '汇入口对应s=0；决策周期Δ=0.5 s，仿真子步h=0.1 s，速度范围0至20 m/s，指令范围−5至2.5 m/s²，指令变化绝对值不超过2.5 m/s²。')
    eq('v_i⁺ = clip(v_i + h u_i, 0, 20),     s_i⁺ = s_i + h v_i⁺')
    p('同来路车辆保持至少6.8 m车头间距。不同来路车辆不能同时占据冲突区Z=[−25,6.8]；双方均进入s≥−25的范围后，还需保持至少6.8 m路线坐标间距。'
      '子步内位置按仿射变化，解sᵢ(λ)∈Z得到占用时间区间，对不同来路要求区间不相交。相对距离在区间端点或过零点检查，从而覆盖采样点之间的冲突。')
    eq('s_i(λ) = s_{i,0} + λ(s_{i,1} − s_{i,0}),     λ ∈ [0,1]')
    p('仿真采用真实Eclipse SUMO和TraCI逐子步速度接口，读取累计行驶距离构造路线坐标。所有车辆受同一控制系统管理，固定路线并关闭内部路口连接。'
      '证明对象是路线合同；连续空间车身几何、防感知误差与混合交通安全不在该合同内。SUMO碰撞事件单独记录。')
    h('3 带指令记忆的联合可恢复证书')
    eq('B_i(u_i⁻) = max(−5, u_i⁻ − 5Δ)')
    p('备用B逐步降低指令，直至三车停止。终端集合要求速度为零、上一指令非正且路线合同成立。没有外部交通时，继续执行B保持终端集合。'
      '可恢复集合K_B由执行B可以安全停止的状态组成。候选动作先检查幅值及指令变化约束，再验证首周期和后续联合制动尾部。实现检查24个备用周期。')
    p('性质1：若初态属于K_B，模型精确、路线固定且车辆服从指令，则逐步选取通过认证的联合模式可保持路线合同和可恢复性。'
      '证明：全备用是可行候选；任意被认证首步都带有安全停止尾部，其后继状态属于K_B。由于备用保留上一指令，原证书尾部与下一周期递推一致。逐步归纳即得结论。')
    p('独立许可不能直接拼接。例如三车(s,v)分别为(−40,10)、(−40,10)、(−90,10)，上一指令均为零。'
      '分别检查“本车加速、其余制动”可获得许可，但把三个许可拼接为同时加速会使两个来路的后续占区区间重叠。联合检查避免该组合错误。')
    p('可恢复性也不保证任务完成。一车停在冲突区、其余车停在区外可满足静止终端合同，持续备用却不会完成合流。'
      '因此本文在同一证书下独立评估到达代价，而不把零碰撞直接视为学习有效。')

    doc.add_page_break()
    h('4 完整轨迹比较与运行时执行')
    p('候选0保留原始恒定加速门控的全部逻辑。其余36项由三个次序偏好(0,1,2)、(0,2,1)、(1,0,2)，三个权重比r∈{2,8,32}，'
      '以及四个滑行时长q∈{0,2,4,6}个决策周期组合。按偏好排列的干预权重为(r²,r,1)。初始q个周期内非首选车辆提案为零加速度，之后恢复最大加速度。')
    p('每个方案仍在提案与备用组成的八种联合模式中认证，以加权平方干预最小者执行；指令先投影到变化约束。'
      '次序只是偏好，证书可否决对应提案。完成车辆统一继续巡航，避免目标线停车对其他车造成偏置。')
    eq('u_k = arg min [ Σ_i w_i(u_i − p_i)² − 10⁻⁸ Σ_i u_i ],   u ∈ F(z_k)')
    p('从实测初态对37项方案分别进行模型内闭环展开，直至全部完成或25 s截止。以首次到达s≥60的时刻定义逐车代价，未完成计25 s，选择平均代价最小者。'
      '相同代价优先保留原始基线。选择在真实执行前完成，不通过回溯真实SUMO测试来挑选结果。')
    eq('J_j = (1/3) Σ_i min(T_i^j, 25),     j^* = arg min J_j')
    p('性质2：初态、模型、生命周期和目标一致且所选计划完整执行时，实际代价等于所选方案代价且不高于基线J₀。'
      '证明：候选0属于比较集合，取最小值不会更差；精确执行具有相同到达时刻。该性质来自保留基线，既不是全局最优性，也不保证每辆车单独更早到达。')
    p('运行时缓存期望状态、上一指令和动作，逐步检查状态匹配并重新认证当前动作。偏差超过10⁻⁶或证书失败时永久退出缓存，改用原反馈门控。'
      '退出后原性能不劣论证不再适用；若扰动使实测状态离开K_B，备用也不具有本文的条件保证。')
    h('5 冻结协议与消融设计')
    p('开发使用种子500至502在常规和高密度分布上的6个场景。算法、候选库、源文件哈希及80个独立测试场景在运行前冻结。'
      '两类分布各40场景，每场景执行五种方法，共400次SUMO评估；不将配对的控制器执行当作400个独立场景。')
    p('五种方法为恒定加速门控、已冻结的RL127门控、按2 s截断进展选择同一候选库、按完整代价选择但不允许提前滑行，以及完整方案。'
      '短期消融选定方案后执行其完整轨迹，不是滚动MPC。原RL检查点在新测试前指定，本轮没有重新训练或挑选学习种子。')
    p('主指标为平均截断到达时间；碰撞回合所有车均计25 s。同时记录完成率、合同违规、模型残差、指令jerk与计划回退。'
      '描述性95%区间由每类40个配对场景进行10000次bootstrap获得，不解释为总体道路风险证书。')

    doc.add_page_break()
    h('6 实验结果与收益来源')
    p('表1给出新测试集结果。冻结RL策略在80个场景上均与恒定加速基线到达代价相同，表明原策略尚未利用更长时域的协调机会。')
    names = [('constant', '恒定加速加门控'), ('frozen_rl127', '冻结RL127加门控'), ('short_portfolio', '2秒目标选择'),
             ('no_coast', '完整目标但无提前滑行'), ('full_portfolio', '完整候选库规划')]
    p('表1  平均截断到达时间  单位为秒', 'Caption')
    table(['方法', '常规 40场景', '高密度 40场景'],
          [[name, f"{summary['nominal'][method]['mean_arrival_s']:.3f}", f"{summary['dense'][method]['mean_arrival_s']:.3f}"] for method, name in names], [8.2, 4.1, 4.1])
    for split, name in [('nominal', '常规'), ('dense', '高密度')]:
        result = summary[split]['full_portfolio']; ci = result['saved_ci95_s']
        p(f"{name}分布平均节省{result['mean_saved_s']:.3f} s，降低{result['reduction_percent']:.3f}%；配对95%区间为[{ci[0]:.3f}, {ci[1]:.3f}] s。"
          f"40个场景中{result['wins_ties_losses'][0]}个改善、{result['wins_ties_losses'][1]}个持平，没有变差。")
    picture = doc.add_paragraph(); picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.add_run().add_picture(str(D/'results.png'), width=Cm(16.3))
    p('图1  左为到达时间对比  右为完整方案逐场景节省时间的累计分布', 'Caption')
    p('仅改变干预权重且不提前滑行时，常规和高密度平均仅改善0.121 s和0.183 s；采用2 s进展排序则无稳定收益。'
      '完整方案的0.800 s和0.754 s收益说明，提前让行需要结合完整到达目标才能在本候选库中被有效利用。')
    p('400次执行全部完成，未记录碰撞、路线合同违规、不可认证动作或计划回退。全部规划预测到达代价与实际值一致，'
      '最大单步模型残差为1.42×10⁻¹³。完整方案常规场景平均指令jerk由1.107降至0.725，高密度由1.218降至1.058；该量不是乘坐舒适性的完整评价。')

    doc.add_page_break()
    h('7 计算成本与工程边界')
    p(f"完整方案在开发电脑上首次规划平均{timing['planning_mean_ms']:.2f} ms，P95为{timing['planning_p95_ms']:.2f} ms。"
      f"将规划与首周期执行相加，首次响应P95为{timing['first_response_p95_ms']:.2f} ms，最大{timing['first_response_max_ms']:.2f} ms；"
      f"执行周期P95为{timing['cycle_p95_ms']:.2f} ms。80次首次响应均低于500 ms决策周期。")
    p('计算采用单工作进程。周期用时含策略推断或当前证书重检、TraCI步进和指标更新，首次规划单列而不摊薄。'
      '仿真在计算期间不推进，因此该结果不是计算延迟下的闭环安全证明，也不是最坏执行时间保证。本文不统计通信延迟。')
    p('工程实现将新规划模块与冻结的历史代码分开，保存全部400条回合记录、80组候选代价、预先指定场景的轨迹和文件哈希。'
      '批量证书逐项对照旧标量实现；实测状态与缓存轨迹不匹配时的反馈回退、SUMO模型残差及重放一致性都有自动检查。')
    p('当前规划器在初始状态做一次完整选择，适用于本基准的确定性短合流过程。连续车辆到达、背景车辆不配合、传感器误差或轮胎动力学变化会破坏其假设。'
      '推广时需要重新建立受扰可恢复集合和滚动更新机制，而不能直接使用现有条件证明宣称实车安全。')
    h('8 讨论与后续学习目标')
    p('本轮实验确定了原学习器没有利用的性能空间：在不放宽安全合同的条件下，主动让行和完整到达时间比较可取得约7.29%的均值改善。'
      '相对于原有每步加速度模仿，下一阶段可以学习次序和让行时机，并考核相对完整规划教师的代价差、计算成本及回退频率。'
      '本次测试场景应封存，新训练与调参使用独立数据。')
    p('轨迹可恢复性的布尔证书与学习性能置信度需要分别定义。现有结果尚未建立统计校准的置信度、未知模型下的鲁棒保证或超过已发表方法的公平对照。'
      '因此本稿提供机制和性能机会证据，尚不足以认定新的MARL算法贡献或顶级期刊会议投稿成熟度。')
    h('参考文献')
    references = [
        '[1] Bertsekas D. Multiagent Rollout Algorithms and Reinforcement Learning. arXiv:1910.00120. https://arxiv.org/abs/1910.00120',
        '[2] Banerjee A, Rahmani K, Biswas J, Dillig I. Dynamic Model Predictive Shielding for Provably Safe Reinforcement Learning. NeurIPS, 2024. https://arxiv.org/abs/2405.13863',
        '[3] Li Y, Chen F, Li Y, Fan C, Bertsekas D. On-Line Policy Iteration with Trajectory-Driven Policy Generation. arXiv:2604.15004, 2026. https://arxiv.org/abs/2604.15004',
        '[4] CoopGate source code and frozen V5 evidence. https://github.com/Leo-7GOAT/CoopGate']
    for reference in references:
        paragraph = p(reference)
        paragraph.paragraph_format.first_line_indent = Cm(0)
        for run in paragraph.runs:
            run.font.size = Pt(9)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out)
    print(out)


if __name__ == '__main__':
    parser = argparse.ArgumentParser(); parser.add_argument('--out', type=Path, required=True)
    build(parser.parse_args().out)
