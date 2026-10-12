"""Chinese research manuscript generated from the frozen V3 measurements."""
from pathlib import Path
import argparse,json,re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
NAMES={'raw_mpc':'原始规则MPC','raw_rl':'无门控RL','joint_mpc':'联合门控MPC','individual':'逐车拼接',
       'independent_front':'独立前车最坏制动','no_memory':'忽略指令记忆','joint_bc':'门控行为克隆','joint_constant':'门控恒定加速','joint_rl':'联合门控RL'}


def build(output):
    data=json.loads((ROOT/'artifacts/v3/final/summary.json').read_text());s=data['summaries']
    extra=json.loads((ROOT/'artifacts/v3/supplement/results.json').read_text())
    doc=Document();section=doc.sections[0]
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=section.bottom_margin=Cm(2.2);section.left_margin=section.right_margin=Cm(2.35)
    for border in doc.styles.element.xpath('.//w:pBdr'):border.getparent().remove(border)
    for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Caption'):
        style=doc.styles[name];style.font.name='Times New Roman';style.font.color.rgb=RGBColor(0,0,0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
    doc.styles['Normal'].font.size=Pt(11);doc.styles['Normal'].paragraph_format.line_spacing=1.2
    doc.styles['Normal'].paragraph_format.space_after=Pt(5)
    doc.styles['Title'].font.size=Pt(19);doc.styles['Subtitle'].font.size=Pt(11)
    doc.styles['Subtitle'].font.italic=False;doc.styles['Heading 1'].font.size=Pt(14);doc.styles['Heading 2'].font.size=Pt(12)
    doc.styles['Caption'].font.size=Pt(9)
    doc.core_properties.author='李东旭';doc.core_properties.title='基于联合制动轨迹认证的多智能体策略切换研究'
    def p(text,style=None):
        q=doc.add_paragraph(text,style)
        if style is None:q.paragraph_format.first_line_indent=Cm(.74)
        return q
    def h(text):doc.add_heading(text,1)
    def eq(text):
        q=doc.add_paragraph();q.alignment=WD_ALIGN_PARAGRAPH.CENTER;m=OxmlElement('m:oMath')
        def run(parent,value):
            r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=value;t.set(qn('xml:space'),'preserve');r.append(t);parent.append(r)
        pos=0
        for token in re.finditer(r'([A-Za-z]+)_(?:\{([^}]+)\}|([A-Za-z0-9]+))',text):
            if token.start()>pos:run(m,text[pos:token.start()])
            sub=OxmlElement('m:sSub');base=OxmlElement('m:e');index=OxmlElement('m:sub')
            run(base,token.group(1));run(index,token.group(2) or token.group(3));sub.append(base);sub.append(index);m.append(sub);pos=token.end()
        if pos<len(text):run(m,text[pos:])
        q._p.append(m)
    def table(headers,rows,widths):
        t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
        border=OxmlElement('w:tblBorders')
        for edge in ('top','bottom','left','right','insideH','insideV'):
            e=OxmlElement('w:'+edge)
            for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]:e.set(qn('w:'+key),value)
            border.append(e)
        t._tbl.tblPr.append(border)
        for col,width in zip(t.columns,widths):col.width=Cm(width)
        for c,w,text in zip(t.rows[0].cells,widths,headers):c.width=Cm(w);c.text=text
        repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            for c,w,text in zip(t.add_row().cells,widths,row):c.width=Cm(w);c.text=str(text)
        for i,row in enumerate(t.rows):
            no=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no)
            for c in row.cells:
                c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7ECF2' if i==0 else ('F5F7FA' if i%2==0 else 'FFFFFF'));c._tc.get_or_add_tcPr().append(shade)
                for q in c.paragraphs:
                    q.alignment=WD_ALIGN_PARAGRAPH.CENTER;q.paragraph_format.space_after=q.paragraph_format.space_before=Pt(4);q.paragraph_format.line_spacing=1
                    for r in q.runs:r.font.size=Pt(9);r.bold=i==0
        gap=doc.add_paragraph();gap.paragraph_format.line_spacing=Pt(4)
        gap.paragraph_format.space_after=Pt(4)
    def mean(split,method,key):return sum(r[key] for r in s[split][method])/3
    p('基于联合制动轨迹认证的多智能体策略切换研究','Title');p('李东旭','Subtitle')
    h('摘要')
    p('针对多车学习控制中的备用策略接管问题，本文研究具有加速度指令变化率约束的固定车道纵向子系统。'
      '构造依赖上一指令的联合制动轨迹，利用候选首周期与备用尾部定义可恢复集合，并在明确动力学及外部领车假设下证明集合不变性与切换闭环安全。'
      '对固定车序和相邻对约束，进一步将完整模式搜索精确分解为链式动态规划。真实SUMO中完成三个独立策略的门控内训练及2020次固定协议评估。'
      '联合门控RL在同分布及双次急刹测试中均未发生碰撞或2米间距违规，而忽略指令记忆及逐车直接拼接均出现失效。'
      f'同分布平均行驶进展为{mean("test","joint_rl","progress_m"):.2f}米，低于同一门控下MPC的{mean("test","joint_mpc","progress_m"):.2f}米与恒定加速策略的{mean("test","joint_constant","progress_m"):.2f}米。'
      '结果支持联合备用检查的作用，但尚不支持学习策略超过强传统基线。本文结论限于声明的纵向离散模型，不覆盖合流、感知误差或真实车辆部署。')
    p('关键词  多智能体强化学习  可恢复性  预测安全过滤  联合制动  动态规划')
    h('1 问题动机与研究定位')
    p('根据置信分数在学习策略与传统控制之间切换时，容易隐含两个未经验证的前提：传统备用始终安全，以及某辆车被接管后不会影响其他车辆的安全判断。'
      '前期合流原型的开发审计中，9次运行包含2772个候选，其中2691个未通过制动余量检查，只有18个同时通过所有条件；396个决策周期中有387个周期的备用方案也未通过制动余量检查。'
      '这表明采用保守独立前车假设的余量判据，与实际依赖规则备用的闭环之间存在明显脱节。')
    p('本阶段将理论对象限定为固定车序下的三车纵向协同。外部领车可以急刹，三辆受控车则共同遵守带指令记忆的备用策略。'
      '目标是构造真正可执行的备用轨迹，并在最终联合动作层面检查约束。本阶段任务与原合流任务不同，因此两阶段指标不作直接提升比例比较。')
    p('本文提供三个可审查的结果：明确模型下的可恢复性证明、独立切换导致证书前提失效的数值反例，以及利用相邻依赖结构的精确模式求解实现。'
      '同时通过行为克隆、MPC及恒定动作对照检验学习是否提供额外收益，避免将安全过滤器本身的作用归因于强化学习。')
    h('2 相关研究与贡献边界')
    p('模型预测安全认证已经通过通往安全终端集合的备用轨迹处理学习控制安全[1,2]。MAMPS考虑多智能体接管的连锁影响[3]，DMPS研究任务收益导向的恢复规划与学习[4]。'
      'AAMAS 2025的组合式屏障工作进一步使用假设保证推理分解多智能体安全义务，并包含车队案例[5]。因此，备用轨迹、联合检查、可组合安全和屏障内学习本身均不能作为本文首次提出的思想。')
    p('本文当前的具体实现特点是：将上一加速度指令纳入备用递推、以共同制动承诺替代受控前车任意独立急刹假设，并利用本问题的链结构精确求解模式选择。'
      '这些差异仍需与符合原论文条件的强基线比较；数学推导是已有备用轨迹与动态规划工具在具有指令变化率约束的纵向系统中的具体化。')
    h('3 动力学与可执行备用策略')
    p('系统含三辆受控车与一辆外部领车，车辆顺序固定，无横向切入和不受控后车。外部领车加速度不低于−9 m/s²且速度非负。'
      '控制周期为0.5秒，仿真子步为0.1秒，速度上限20 m/s。每条受控指令保持五个子步，状态包括车辆位置、速度和上一加速度指令。')
    eq('v_i⁺ = clip(v_i + h u_i, 0, 20),     p_i⁺ = p_i + h v_i⁺')
    p('指令范围为−5至2.5 m/s²，相邻指令变化不超过2.5 m/s²，即指令jerk上限5 m/s³。速度截断时实际加速度变化不必满足相同上限，故不把指令jerk等同于实车舒适性保证。'
      'SUMO使用显式加速度适配器：每个子步按上式计算目标速度并通过TraCI setSpeed执行。关闭隐式跟驰防撞与换道干预，只保留初始化的几何碰撞检查。')
    p('每辆受控车的备用指令由上一指令逐步降至制动下限，三辆车同时执行。外部领车以最强允许制动预测。')
    eq('B_i(u_{i,prev}) = max(−5, u_{i,prev} − 5Δ)')
    p('备用轨迹持续至三辆受控车全部停止，上一指令非正且相邻净间距不低于2米，构成终端集合T。由于领车不倒车，继续执行备用可以保持T。'
      '实现最多检查24个备用周期；未达到终端则不签发证书。初态按备用轨迹计算间距后随机增加余量，并显式检查其属于可恢复集合，不声称能救回任意初始危险状态。')
    h('4 联合证书与理论结论')
    doc.add_heading('4.1 可恢复集合与切换闭环',2)
    p('定义K_B为执行联合备用B时始终满足相邻间距约束且有限时间到达T的状态集合。对任何候选联合指令，先推演一个周期，再执行B至T；仅当整条轨迹通过检查时放行。'
      '所有候选失效时执行已经具有备用轨迹的B。数值实现增加0.05米缓冲，该缓冲不是已证明的实车模型误差界。')
    p('命题1：在精确离散模型、固定车序及领车界成立时，K_B在备用B下不变。证明：执行B一步后，后续备用序列恰好是原轨迹尾部，因为备用递推保留了上一指令。'
      '尾部仍安全并到达T。实际领车的位置和速度不低于最坏制动轨迹，且离散更新对领车位置和速度单调，所以领车与首车的间距不会比最坏预测更小。')
    p('命题2：若初态属于K_B，且每步执行经联合证书验证的候选或备用B，则闭环持续满足相邻间距约束并保留可恢复性。'
      '证明：候选首周期由证书保证安全，其末态拥有备用尾部，故属于K_B；备用本身由命题1保证可行。逐步归纳即得。'
      '每个子步内位置按下一速度线性推进，检查相邻子步端点覆盖该执行模型的线性插值；该说法不直接适用于其他连续车辆模型。')
    doc.add_heading('4.2 独立拼接反例',2)
    p('开发种子0的初始备用裕度约6.85米。三辆车都提出2.5 m/s²时，第一、第二、第三车的局部裕度分别约−13.47、9.30、6.85米。'
      '如果只按局部结果让第一辆车改为−2.5 m/s²，其他两车继续加速，则最终联合备用裕度降至约−1.15米。'
      '后车的检查原本依赖前车继续加速，切换改变了前提。联合检查拒绝该组合。负裕度表示2米间距约束不能得到证书，并不等于已经发生物理碰撞。')
    doc.add_heading('4.3 精确链式模式求解',2)
    p('每辆车仅有首选与备用两种首步模式，之后均执行只依赖自身指令的备用尾部。因此可以独立计算每辆车的两条位置轨迹，'
      '再构建相邻模式兼容表。记Cᵢ(a,b)表示前车模式a与本车模式b的轨迹相容，cᵢ(b)为本车干预代价，则')
    eq('D_i(b) = c_i(b) + min{D_{i−1}(a) : C_i(a,b)=1}')
    p('命题3：相邻约束、可独立生成的备用轨迹及可加干预代价条件成立时，上述二状态动态规划与完整模式穷举等价。'
      '证明：给定末车模式b，可行前缀仅通过前一模式a与末车耦合；同一a下更昂贵的前缀可被最优前缀替代而不改变末端可行性，故Bellman递推成立。'
      '固定停止时域H时，轨迹及兼容表计算为O(nH)，模式求解为O(n)，仅需2+4(n−1)个兼容表项。实现采用回溯指针，避免复制增长的路径。')
    p('该分解适用于本问题的固定链结构。换道、超车、非相邻约束或依赖其他车辆状态的备用控制可能破坏分解条件。'
      '动态规划本身为标准算法，本文不据此主张通用多智能体问题具有线性复杂度。')
    h('5 学习与实验协议')
    p('使用参数共享Actor及集中式Critic，五种加速度提案为−5、−2.5、0、1.25、2.5 m/s²，并投影到指令变化率范围。'
      '三个训练种子83、97、109各使用30回合联合门控MPC示范、15轮行为克隆，然后运行12000步门控内PPO。'
      'PPO概率比对应策略提出的动作，门控替换属于环境转移；奖励使用执行后的进展，并惩罚指令变化和干预。选用最终检查点，没有按测试表现选模型。')
    p('九种方法包括原始规则MPC、无门控RL、联合门控MPC、逐车拼接、独立前车最坏制动、忽略指令记忆、门控行为克隆、门控恒定最大加速及联合门控RL。'
      '每模型每方法测试30个同分布场景与30个未训练的双次急刹时序，共1620次执行；另对三个联合门控RL和一个联合门控MPC各运行100个独立认证回合，共2020次主实验。')
    p('同分布混合巡航、普通制动与急刹，双次急刹改变扰动时序但仍满足领车制动界。所有回合固定25秒，主指标为三辆受控车的平均行驶距离；碰撞回合的进展记0。'
      '相邻净间距低于2米即记录违规，不以碰撞作为唯一安全指标。不同方法和训练种子共享场景，重复基线不能作为额外独立样本。')
    p('对固定控制器分别检验碰撞与间距违规，共8项，采用单侧精确二项上界及Bonferroni校正，整体置信水平95%。'
      '描述性配对bootstrap按每个训练种子内的30个场景重采样，不作多重比较显著性声明。统计检验采用标准方法，不作为理论创新。')
    h('6 实验结果与消融')
    for split,label in (('test','同分布场景'),('shift','双次急刹场景')):
        doc.add_heading('6.1 '+label if split=='test' else '6.2 '+label,2)
        table(['方法','种子83','种子97','种子109','进展 m','jerk'],
            [[name]+[f'{r["collisions"]}/{r["gap_violation_episodes"]}' for r in s[split][method]]+
             [f'{mean(split,method,"progress_m"):.2f}',f'{mean(split,method,"jerk"):.2f}'] for method,name in NAMES.items()],
             [4.1,2.2,2.2,2.2,2.7,2.6])
        p('种子列为碰撞回合数与间距违规回合数，每项分母30；进展与指令jerk取三组策略均值。jerk单位为m/s³，提前碰撞回合的jerk不与无碰撞方法作舒适性优势比较。')
    p('联合门控RL在两类测试的三个训练种子中均为零碰撞、零间距违规，且日志中没有离开可恢复集或执行未认证指令的周期。'
      '忽略指令记忆在同分布中产生5、10、14次碰撞，在双次急刹中产生11、20、22次；逐车拼接在两类测试中也产生违规。'
      '这与理论前提及数值反例一致，但有限测试本身不能替代模型条件。')
    p(f'联合门控RL的提案保留率在两类场景中分别为{mean("test","joint_rl","retained_fraction"):.1%}与{mean("shift","joint_rl","retained_fraction"):.1%}。'
      '门控内强化学习在三个种子中都提高了相对于行为克隆初始化的平均进展，但未稳定超过联合门控MPC，且三个种子在两类场景中均未超过门控恒定加速的平均进展。'
      '恒定加速更快但指令变化更大；当前数据尚未证明学习带来相对于调优传统控制的Pareto优势。')
    p('独立前车最坏制动同样保持约束，但同分布下其进展低于相同策略的联合门控。这里放宽的是受控车辆之间的独立急刹假设；'
      '外部领车的制动界保持一致。对双次急刹，收益并非所有种子都一致，不能宣称普遍性能提升。')
    doc.add_picture(str(ROOT/'artifacts/v3/final/results.png'),width=Cm(16));doc.paragraphs[-1].paragraph_format.keep_with_next=True
    p('图1 固定车道测试的行驶进展与间距违规 黑点为三个训练种子','Caption')
    doc.add_heading('6.3 独立风险检验',2)
    p('三个联合门控RL策略与联合门控MPC在各自100个独立回合中均没有碰撞或间距违规。8项同时检验的风险上界均为4.95%，通过预设5%目标。'
      '该目标用于展示有限样本统计推断，不是道路部署标准；该结果也没有证明学习优于MPC。')
    h('7 工程验证与复杂度')
    p('链式求解属于固定协议之外的算法等价性与工程补充，不作为独立调参后的效果证据。'
      '随机状态测试将链式输出与完整轨迹穷举比较；三个发布策略还分别重放一个预先确定的SUMO测试场景，合计150个决策周期的指令和状态一致。')
    table(['实现','循环均值 ms','P95 ms','最大 ms','超500 ms'],
        [[name,f'{extra["timing"][key]["mean_ms"]:.2f}',f'{extra["timing"][key]["p95_ms"]:.2f}',f'{extra["timing"][key]["max_ms"]:.2f}',extra['timing'][key]['over_500ms']]
         for key,name in [('joint_mpc','联合门控MPC'),('joint_rl','穷举联合门控RL'),('chain_rl','链式联合门控RL')]],[4,3,3,3,3])
    p('完整循环含策略推理、门控与SUMO步进，在主实验结束后单仿真进程测量，每方法5个场景。样本P95与最大值不能作为操作系统最坏执行时间保证。')
    table(['受控车数','兼容表项','完整模式数','链式总耗时均值 ms'],
        [[r['vehicles'],r['pair_entries'],str(r['joint_modes']),f'{r["mean_total_ms"]:.3f}'] for r in extra['scaling']], [3,3,6,4])
    p('扩展性表来自合成固定车序状态，包含轨迹生成、兼容表构建与动态规划。它不是64车SUMO闭环实验。'
      '理论复杂度来自固定时域和链式依赖条件；更复杂交互需要不同分解。')
    h('8 局限与下一步')
    p('当前的主要推进是将经验备用改为具有明确前提的可执行证书，并实现同一安全选择问题的精确结构化求解。'
      '但车辆模型没有执行误差，车序固定，没有横向切入和不可控后车；理论不能覆盖原合流场景，更不能直接覆盖实车。')
    p('学习部分仍是主要短板。简单纵向任务中，恒定加速加安全门控已经形成很强的进展基线，说明仅增加学习训练并不能自然产生协同决策贡献。'
      '后续应将可恢复证书接回需要让行、次序选择或换道承诺的场景，并检验学习是否在相同安全和计算预算下改善效率与指令平顺性。'
      '还需开展原方法条件相符的强基线比较、模型误差下的鲁棒证书，以及链式分解失效时的算法设计。')
    h('9 结论')
    p('本文在明确限定的纵向子系统中实现了联合备用轨迹、递归可恢复性检查与精确链式模式求解，并通过真实SUMO、多种消融及独立回合检验公开其表现。'
      '现有证据支持指令记忆和最终联合动作检查的必要性；对于学习策略是否提供超越强传统基线的额外收益，实验仍给出否定或不足的证据。'
      '后续研究需在具有协同决策需求的任务与动力学不确定性下，进一步检验这一方法的适用性。')
    h('参考文献')
    references=[
      '[1] Wabersich K P, Zeilinger M N. Linear model predictive safety certification for learning-based control. 2018. https://arxiv.org/abs/1803.08552',
      '[2] Wabersich K P, Zeilinger M N. A predictive safety filter for learning-based control of constrained nonlinear dynamical systems. Automatica, 2021, 129:109597. https://doi.org/10.1016/j.automatica.2021.109597',
      '[3] Zhang W, Bastani O, Kumar V. MAMPS Safe Multi-Agent Reinforcement Learning via Model Predictive Shielding. 2019. https://arxiv.org/abs/1910.12639',
      '[4] Banerjee A, Rahmani K, Biswas J, Dillig I. Dynamic Model Predictive Shielding for Provably Safe Reinforcement Learning. NeurIPS, 2024. https://arxiv.org/abs/2405.13863',
      '[5] Brorholt A H, Larsen K G, Schilling C. Compositional Shielding and Reinforcement Learning for Multi-Agent Systems. AAMAS, 2025:399–407. https://www.ifaamas.org/Proceedings/aamas2025/pdfs/p399.pdf',
      '[6] Eclipse SUMO Documentation. https://eclipse.dev/sumo/docs/']
    for ref in references:
        q=p(ref);q.paragraph_format.first_line_indent=Cm(0)
        q.paragraph_format.space_after=Pt(3)
        for r in q.runs:r.font.size=Pt(9)
    footer=section.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    doc.save(output);print(output)


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--out',required=True);build(parser.parse_args().out)
